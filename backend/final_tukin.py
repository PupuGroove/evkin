"""Optional SKP finalization. All amounts are immutable integer-Rupiah snapshots."""
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import importlib
import io
import json
import re
import uuid
from itertools import islice
from zipfile import BadZipFile

import duckdb
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from backend.skp_assessment import (
    VALID_PREDICATES as CANONICAL_SKP_RATES,
    annotate_skp_tukin_provenance,
    canonical_skp_reduction_rate,
    resolve_skp_for_tukin,
    skp_ready_for_tukin,
)
from backend.tukin_combined import (
    calculation_scope,
    calculate_tukin_amounts,
    integration_readiness,
)

ensure_integration_schema = importlib.import_module(
    'backend.migrations.048_create_canonical_tukin_integration'
).ensure_canonical_tukin_integration_schema

MODES = {'ABSENSI_ONLY', 'ABSENSI_SKP'}
SOURCES = {'ASSESSMENT', 'TEMPLATE_IMPORT', 'EXTERNAL_IMPORT'}
SKP_UNAVAILABLE_LABELS = {'-', 'TIDAK BISA DINILAI', 'TIDAK MEMBUAT SKP'}
SKP_NORMAL_RATES = dict(CANONICAL_SKP_RATES)


def resolve_skp_rate(skp_result):
    if not skp_result:
        return None

    explicit_present = 'potongan_skp_percent' in skp_result
    explicit_value = skp_result.get('potongan_skp_percent')
    predicate = str(skp_result.get('predikat_skp') or '').strip().upper()
    report_status = str(skp_result.get('report_status') or '').upper()
    consecutive_missing = int(skp_result.get('consecutive_missing') or 0)

    if explicit_present:
        explicit_rate = decimal(explicit_value)
        if explicit_rate is None or not 0 <= explicit_rate <= 100:
            return None
        return float(explicit_rate)

    canonical_status = report_status
    if predicate in {'-', ''} or predicate in SKP_UNAVAILABLE_LABELS:
        canonical_status = 'NOT_SUBMITTED_VERIFIED'
    return canonical_skp_reduction_rate(predicate, canonical_status, consecutive_missing)


def signature(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str, ensure_ascii=False).encode()).hexdigest()


def decimal(value):
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def rupiah(value):
    return int(value.quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def period(value):
    value = str(value or '').strip()
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', value):
        raise ValueError('Periode wajib YYYY-MM')
    return value


def calculate_final_tukin(nip, periode, mode, attendance_result, employee_master, skp_result=None):
    period(periode)
    if mode not in MODES:
        raise ValueError('Mode tidak valid')
    person = employee_master or {}
    absence = attendance_result or {}
    row = dict(nip=str(nip), periode=periode, tahun=int(periode[:4]), bulan=int(periode[5:]), mode=mode,
               **{k: person.get(k) for k in ('nama', 'provinsi', 'kabupaten_kota', 'kecamatan', 'jabatan')})
    row.update({k: None for k in ('tukin_bruto', 'komponen_absen_30', 'potongan_absen_percent',
        'potongan_absen_rp', 'predikat_skp', 'komponen_skp_70', 'potongan_skp_percent',
        'potongan_skp_rp', 'total_potongan', 'tukin_dibayar', 'skp_result_id')})
    row.update(status='READY', reason='', skp_source='TIDAK DIGUNAKAN', status_skp='TIDAK DIGUNAKAN',
               skp_source_status='TIDAK DIGUNAKAN', skp_calculation_mode='TIDAK DIGUNAKAN',
               skp_fallback_policy=None, skp_description='',
               attendance_result_id=absence.get('id') or signature([nip, periode, absence]))
    issues = []
    gross = decimal(person.get('tukin_gross'))
    rate = decimal(absence.get('total'))
    if gross is None or gross < 0:
        issues.append('MISSING_TUKIN')
    if not attendance_result:
        issues.append('ATTENDANCE_NOT_FOUND')
    elif absence.get('_duplicate'):
        issues.append('DUPLICATE_ATTENDANCE')
    elif rate is None or not 0 <= rate <= 100:
        issues.append('INVALID_PERCENTAGE')
    if absence.get('_unverified'):
        issues.append('ATTENDANCE_NOT_VERIFIED')
    if person.get('eligible') is False:
        issues.append('MASTER_NOT_ELIGIBLE')
    skp = skp_result or {}
    skp_rate = None
    if mode == 'ABSENSI_SKP':
        row.update(skp_source=skp.get('source_type'), status_skp=skp.get('status_skp'),
                   predikat_skp=skp.get('predikat_skp'), skp_result_id=skp.get('id'))
        if not skp_result:
            semantic = resolve_skp_for_tukin(None, attendance_participant=bool(attendance_result))
            row.update(skp_source='TIDAK TERSEDIA', status_skp=semantic['source_status'],
                       skp_source_status=semantic['source_status'],
                       skp_calculation_mode=semantic['ready_mode'],
                       skp_fallback_policy=semantic['fallback_policy'],
                       skp_description=semantic['description'], predikat_skp='SKP TIDAK TERSEDIA')
            if semantic['ready']:
                skp_rate = semantic['fallback_rate']
            else:
                issues.append('SKP_NOT_FOUND')
        elif skp.get('_duplicate'):
            issues.append('DUPLICATE_SKP')
        else:
            if skp.get('source_type') in {'ASSESSMENT', 'DIRECT_FINAL_IMPORT'}:
                semantic = skp_ready_for_tukin(
                    skp.get('assessment') or skp,
                    skp.get('source_type') == 'DIRECT_FINAL_IMPORT',
                )
                if not semantic['ready']:
                    issues.append('SKP_NOT_FINAL')
                row.update(skp_source_status=semantic['source_status'],
                           skp_calculation_mode=semantic['ready_mode'],
                           skp_fallback_policy=semantic['fallback_policy'],
                           skp_description=semantic['description'])
            elif skp.get('status_skp') != 'IMPORTED_FINAL':
                issues.append('SKP_NOT_FINAL')
            else:
                row.update(skp_source_status='FINAL_IMPORTED', skp_calculation_mode='NORMAL',
                           skp_fallback_policy=None, skp_description='SKP final tersedia')
            skp_rate = resolve_skp_rate(skp)
            predicate = str(skp.get('predikat_skp') or '').strip().upper()
            if skp_rate is None and predicate in SKP_UNAVAILABLE_LABELS:
                issues.append('SKP_UNAVAILABLE')
            elif skp_rate is None or not 0 <= skp_rate <= 100:
                issues.append('INVALID_PERCENTAGE')
            if skp.get('payment_eligible') is False:
                issues.append('SKP_PAYMENT_HOLD')
        if skp_result and (skp.get('nip') not in (None, str(nip)) or skp.get('periode') not in (None, periode)):
            issues.append('SKP_NOT_FOUND')
    if gross is not None and gross >= 0:
        row['tukin_bruto'] = rupiah(gross)
        row['komponen_absen_30'] = rupiah(Decimal(row['tukin_bruto']) * Decimal('.30'))
        if mode == 'ABSENSI_SKP':
            # Complementary rounded components ensure even 100% + 100% reconciles.
            row['komponen_skp_70'] = row['tukin_bruto'] - row['komponen_absen_30']
    if rate is not None and 0 <= rate <= 100:
        row['potongan_absen_percent'] = str(rate)
    if skp_rate is not None and 0 <= skp_rate <= 100:
        row['potongan_skp_percent'] = str(skp_rate)
    if issues:
        row.update(status='HOLD', reason=';'.join(dict.fromkeys(issues)))
        return row
    amounts = calculate_tukin_amounts(gross, rate, skp_rate if mode == 'ABSENSI_SKP' else 0)
    row['potongan_absen_rp'] = amounts['attendance_deduction']
    if mode == 'ABSENSI_SKP':
        row['potongan_skp_rp'] = amounts['skp_deduction']
    row['total_potongan'] = row['potongan_absen_rp'] + (row['potongan_skp_rp'] or 0)
    row['tukin_dibayar'] = amounts['calculated_entitlement']
    assert row['tukin_dibayar'] >= 0
    assert row['tukin_bruto'] == row['tukin_dibayar'] + row['potongan_absen_rp'] + (row['potongan_skp_rp'] or 0)
    return row


def migrate(db_path):
    with duckdb.connect(str(db_path)) as c:
        c.execute('BEGIN')
        c.execute('CREATE TABLE IF NOT EXISTS final_tukin_schema(version INTEGER PRIMARY KEY, applied_at TIMESTAMP DEFAULT now())')
        c.execute('''CREATE TABLE IF NOT EXISTS master_skp_rules(
            id VARCHAR PRIMARY KEY, predikat VARCHAR, potongan_percent DECIMAL(9,4),
            is_active BOOLEAN DEFAULT true, effective_from DATE, effective_to DATE, notes VARCHAR)''')
        c.execute('''CREATE TABLE IF NOT EXISTS skp_import_batches(
            id VARCHAR PRIMARY KEY, periode VARCHAR, source_type VARCHAR, source_file VARCHAR,
            source_sheet VARCHAR, source_hash VARCHAR, row_count INTEGER, created_by VARCHAR,
            imported_at TIMESTAMP DEFAULT now())''')
        c.execute('ALTER TABLE skp_import_batches ADD COLUMN IF NOT EXISTS excluded_count INTEGER DEFAULT 0')
        c.execute('ALTER TABLE skp_import_batches ADD COLUMN IF NOT EXISTS validation_report VARCHAR')
        c.execute('''CREATE TABLE IF NOT EXISTS skp_results(
            id VARCHAR PRIMARY KEY, nip VARCHAR NOT NULL, periode VARCHAR NOT NULL, tahun INTEGER, bulan INTEGER,
            nama VARCHAR, predikat_skp VARCHAR, potongan_skp_percent DECIMAL(9,4), status_skp VARCHAR,
            source_type VARCHAR, source_file VARCHAR, source_sheet VARCHAR, assessment_id VARCHAR,
            reviewer VARCHAR, approver VARCHAR, evidence VARCHAR, notes VARCHAR, import_batch_id VARCHAR NOT NULL,
            skp_percentage_source VARCHAR, payload VARCHAR, created_at TIMESTAMP DEFAULT now(),
            updated_at TIMESTAMP DEFAULT now(), UNIQUE(nip,periode,source_type,import_batch_id))''')
        c.execute('''CREATE TABLE IF NOT EXISTS final_tukin_batches(
            id VARCHAR PRIMARY KEY, periode VARCHAR, mode VARCHAR, version INTEGER, skp_source VARCHAR,
            skp_batch_id VARCHAR, jumlah_sdm INTEGER, total_tukin_bruto BIGINT, total_potongan_absen BIGINT,
            total_potongan_skp BIGINT, total_tukin_dibayar BIGINT, source_signature VARCHAR,
            created_by VARCHAR, created_at TIMESTAMP DEFAULT now(), archived_at TIMESTAMP,
            UNIQUE(periode,mode,version))''')
        c.execute('''CREATE TABLE IF NOT EXISTS final_tukin_results(
            id VARCHAR PRIMARY KEY, calculation_batch_id VARCHAR, nip VARCHAR, periode VARCHAR,
            tahun INTEGER, bulan INTEGER, mode VARCHAR, version INTEGER,
            tukin_bruto BIGINT, komponen_absen_30 BIGINT, potongan_absen_percent DECIMAL(9,4), potongan_absen_rp BIGINT,
            predikat_skp VARCHAR, komponen_skp_70 BIGINT, potongan_skp_percent DECIMAL(9,4), potongan_skp_rp BIGINT,
            total_potongan BIGINT, tukin_dibayar BIGINT, skp_source VARCHAR, skp_result_id VARCHAR,
            attendance_result_id VARCHAR, status VARCHAR, reason VARCHAR, source_signature VARCHAR, snapshot VARCHAR,
            calculated_at TIMESTAMP DEFAULT now(), created_at TIMESTAMP DEFAULT now(), updated_at TIMESTAMP DEFAULT now(),
            UNIQUE(nip,periode,mode,version))''')
        if not c.execute('SELECT count(*) FROM final_tukin_schema WHERE version=1').fetchone()[0]:
            for name, rate in [('SANGAT BAIK',0),('BAIK',0),('BUTUH PERBAIKAN',5),('KURANG',10),('SANGAT KURANG',15)]:
                c.execute('''INSERT INTO master_skp_rules
                    (id,predikat,potongan_percent,is_active,effective_from,effective_to,notes)
                    VALUES(?,?,?,true,?,NULL,?)''',
                          [str(uuid.uuid4()), name, rate, '1900-01-01', 'Konfigurasi awal; tinjau sesuai kebijakan instansi'])
            c.execute('INSERT INTO final_tukin_schema(version) VALUES(1)')
        c.execute('INSERT INTO final_tukin_schema(version) VALUES(2) ON CONFLICT DO NOTHING')
        ensure_integration_schema(c)
        c.execute('COMMIT')


def records(c, table, filters=None):
    if not c.execute('SELECT count(*) FROM information_schema.tables WHERE table_name=?', [table]).fetchone()[0]:
        return []
    filters=filters or {}
    if filters:
        columns={r[0] for r in c.execute(f'DESCRIBE "{table}"').fetchall()}
        if not set(filters)<=columns: return []
    query=f'SELECT * FROM "{table}"'
    if filters: query+=' WHERE '+' AND '.join('"'+k+'"=?' for k in filters)
    cursor = c.execute(query,list(filters.values()))
    names = [d[0] for d in cursor.description]
    return [dict(zip(names, values)) for values in cursor.fetchall()]


def insert(c, table, row):
    keys = list(row)
    c.execute(f'INSERT INTO "{table}" (' + ','.join('"'+k+'"' for k in keys) + ') VALUES (' + ','.join('?' for _ in keys) + ')', list(row.values()))


def insert_many(c, table, rows):
    """Bounded parameter batches avoid per-row DuckDB round trips for national BNBA."""
    if not rows:
        return
    keys=list(rows[0])
    prefix=f'INSERT INTO "{table}" ('+','.join('"'+k+'"' for k in keys)+') VALUES '
    placeholder='('+','.join('?' for _ in keys)+')'
    for start in range(0,len(rows),500):
        chunk=rows[start:start+500]
        c.execute(prefix+','.join(placeholder for _ in chunk),[r.get(k) for r in chunk for k in keys])


ALIASES = {
    'nip': ['NIP','Nomor Induk Pegawai','NIP PPPK'],
    'nama': ['Nama','Nama Pegawai','Nama SDM'],
    'periode': ['Periode','Periode Penilaian','Assessment Period','Bulan'],
    'tahun': ['Tahun','Year'],
    'predikat_skp': ['Predikat Kinerja Periodik','Predikat SKP','Hasil SKP','Predikat'],
    'potongan_skp_percent': ['Potongan SKP (%)','Persentase Potongan SKP','Potongan SKP'],
    'reviewer': ['Penilai','Reviewer'], 'status_laporan': ['Status Laporan SKP','Status SKP'],
    'notes': ['Catatan','Catatan Pokja','Alasan'], 'evidence': ['Evidence','Referensi Eviden','Eviden'],
    'target': ['Target'], 'realisasi': ['Realisasi'], 'rekomendasi': ['Rekomendasi'],
}


def header(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower().replace('%','percent'))


def auto_mapping(values):
    result = {}
    for key, aliases in ALIASES.items():
        for alias in aliases:
            hits = [i for i,v in enumerate(values) if header(v) == header(alias)]
            if hits:
                if len(hits)==1: result[key]=hits[0]
                break
    return result


def import_period(value, year, fallback):
    if value in (None,''): return fallback
    if isinstance(value,datetime): return value.strftime('%Y-%m')
    text=str(value).strip()
    if re.fullmatch(r'\d{4}-\d{2}',text): return period(text)
    months=['januari','februari','maret','april','mei','juni','juli','agustus','september','oktober','november','desember']
    month=months.index(text.lower())+1 if text.lower() in months else int(text) if text.isdigit() else 0
    if 1<=month<=12:
        return period(f'{int(year or fallback[:4]):04d}-{month:02d}')
    raise ValueError('Periode import tidak dikenali')


class FinalTukinService:
    def __init__(self, db_path):
        self.db_path = str(db_path)
        migrate(self.db_path)

    def rule_rate(self, c, predicate, periode):
        rows = c.execute('''SELECT potongan_percent FROM master_skp_rules WHERE is_active
            AND upper(predikat)=? AND effective_from<=CAST(? AS DATE)
            AND (effective_to IS NULL OR effective_to>=CAST(? AS DATE))''',
            [predicate, periode+'-01', periode+'-01']).fetchall()
        return rows[0][0] if len(rows) == 1 else None

    def preview_import(self, content, filename, periode, source_type, sheet=None, mapping=None, header_row=None, percent_unit='POINTS'):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        period(periode)
        if source_type not in SOURCES - {'ASSESSMENT'}:
            raise ValueError('Pilih TEMPLATE_IMPORT atau EXTERNAL_IMPORT')
        if percent_unit not in {'POINTS','FRACTION'}:
            raise ValueError('Satuan persentase tidak valid')
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile,InvalidFileException,KeyError) as exc:
            raise ValueError('Workbook tidak valid. Gunakan file Excel .xlsx yang dapat dibuka.') from exc
        try:
            if not sheet:
                def score(name):
                    return max((len(auto_mapping(row)) for row in islice(wb[name].values,15)),default=0)
                sheet=max(wb.sheetnames,key=score)
            if sheet not in wb.sheetnames:
                raise ValueError('Sheet tidak ditemukan')
            cells = list(wb[sheet].iter_rows())
            if not cells:
                raise ValueError('Sheet kosong')
            if header_row is None:
                header_row = max(range(min(50,len(cells))), key=lambda i: len(auto_mapping([c.value for c in cells[i]])))
            if not 0 <= header_row < len(cells):
                raise ValueError('Baris header tidak valid')
            headers = [str(c.value or '') for c in cells[header_row]]
            mapping = auto_mapping(headers) if mapping is None else mapping
            if any(k not in ALIASES or not isinstance(v,int) or not 0 <= v < len(headers) for k,v in mapping.items()):
                raise ValueError('Mapping tidak valid')
            output = []
            rule_cache={}
            with duckdb.connect(self.db_path) as c:
                for index, cells_row in enumerate(cells[header_row+1:], header_row+2):
                    if all(cell.value in (None,'') for cell in cells_row):
                        continue
                    def get(key):
                        pos = mapping.get(key)
                        return cells_row[pos].value if pos is not None and pos < len(cells_row) else None
                    nip = str(get('nip') or '').strip().lstrip("'")
                    raw_predicate=get('predikat_skp')
                    predicate = str(raw_predicate if raw_predicate is not None else '').strip().upper()
                    errors = []
                    try: p=import_period(get('periode'),get('tahun'),periode)
                    except (ValueError,TypeError): p=str(get('periode') or '');errors.append('PERIOD_INVALID')
                    if not re.fullmatch(r'\d{18}', nip):
                        errors.append('NIP_INVALID' if nip else 'NIP_EMPTY')
                    if p != periode:
                        errors.append('PERIOD_MISMATCH')
                    raw_rate = get('potongan_skp_percent')
                    explicit = raw_rate not in (None,'')
                    rate = None
                    if explicit:
                        text = str(raw_rate).strip()
                        rate = decimal(text.replace('%','').replace(',','.'))
                        cell = cells_row[mapping['potongan_skp_percent']]
                        if rate is not None and '%' not in text and ('%' in cell.number_format or percent_unit == 'FRACTION'):
                            rate *= 100
                    else:
                        if predicate not in rule_cache: rule_cache[predicate]=self.rule_rate(c,predicate,periode)
                        rate = rule_cache[predicate]
                    unavailable = not explicit and predicate in SKP_UNAVAILABLE_LABELS
                    if not unavailable and (rate is None or not 0 <= rate <= 100 or rate != rate.quantize(Decimal('.0001'))):
                        errors.append('INVALID_PERCENTAGE')
                    row = dict(nip=nip, periode=p, nama=str(get('nama') or ''), predikat_skp=predicate or None,
                               potongan_skp_percent=str(rate) if rate is not None else None,
                               skp_percentage_source='IMPORTED' if explicit else 'RULE_MAPPING',
                               reviewer=str(get('reviewer') or ''), evidence=str(get('evidence') or ''),
                               notes=str(get('notes') or ''), source_row=index, errors=errors,
                               raw={f'{i+1}:{h}': cell.value for i,(h,cell) in enumerate(zip(headers,cells_row))})
                    output.append(row)
            counts = Counter((r['nip'],r['periode']) for r in output)
            for r in output:
                if counts[(r['nip'],r['periode'])] > 1:
                    r['errors'].append('DUPLICATE_SKP')
            stats = dict(jumlah_baris=len(output), nip_valid=sum(bool(re.fullmatch(r'\d{18}',r['nip'])) for r in output),
                         nip_kosong=sum(not r['nip'] for r in output), duplikat=sum('DUPLICATE_SKP' in r['errors'] for r in output),
                         predikat_kosong=sum(not r['predikat_skp'] for r in output), potongan_kosong=sum(r['potongan_skp_percent'] is None for r in output),
                         skp_hold=sum(str(r['predikat_skp'] or '').upper() in SKP_UNAVAILABLE_LABELS and r['potongan_skp_percent'] is None for r in output),
                         invalid=sum(bool(r['errors']) for r in output))
            response = dict(file=filename,sheet=sheet,sheets=wb.sheetnames,header_row=header_row,headers=headers,mapping=mapping,
                            periode=periode,source_type=source_type,source_hash=hashlib.sha256(content).hexdigest(),stats=stats,rows=output)
            response['preview_signature'] = signature(response)
            return response
        finally:
            wb.close()

    def import_preview(self, preview, expected_signature, actor='User', valid_only=False):
        if preview['preview_signature'] != expected_signature:
            raise ValueError('Preview berubah; lakukan preview kembali')
        if (preview['stats']['invalid'] and not valid_only) or not preview['rows']:
            raise ValueError('Import ditolak: data kosong atau masih memiliki konflik')
        selected=[r for r in preview['rows'] if not r['errors']]
        excluded=[{k:r.get(k) for k in ('source_row','nip','periode','predikat_skp','errors')} for r in preview['rows'] if r['errors']]
        if not selected: raise ValueError('Tidak ada baris valid untuk diimport')
        batch = str(uuid.uuid4())
        with duckdb.connect(self.db_path) as c:
            c.execute('BEGIN')
            if c.execute('SELECT count(*) FROM skp_import_batches WHERE periode=? AND source_type=? AND source_hash=? AND source_sheet=?',
                         [preview['periode'],preview['source_type'],preview['source_hash'],preview['sheet']]).fetchone()[0]:
                raise ValueError('File/sheet ini sudah diimport untuk periode dan sumber yang sama')
            insert(c,'skp_import_batches',dict(id=batch,periode=preview['periode'],source_type=preview['source_type'],
                source_file=preview['file'],source_sheet=preview['sheet'],source_hash=preview['source_hash'],row_count=len(selected),created_by=actor,
                excluded_count=len(excluded),validation_report=json.dumps(excluded,ensure_ascii=False)))
            imported_rows=[]
            for row in selected:
                imported_rows.append(dict(id=str(uuid.uuid4()),nip=row['nip'],periode=row['periode'],tahun=int(row['periode'][:4]),bulan=int(row['periode'][5:]),
                    **{k:row.get(k) for k in ('nama','predikat_skp','potongan_skp_percent','reviewer','evidence','notes','skp_percentage_source')},
                    status_skp='IMPORTED_FINAL',source_type=preview['source_type'],source_file=preview['file'],source_sheet=preview['sheet'],
                    import_batch_id=batch,payload=json.dumps(row['raw'],default=str,ensure_ascii=False)))
            insert_many(c,'skp_results',imported_rows)
            c.execute('COMMIT')
        return dict(import_batch_id=batch, imported=len(selected),excluded=len(excluded))

    def sources(self, periode, mode, source_type=None, skp_batch_id=None):
        from backend.shared_master import SharedMasterService
        period(periode)
        if mode not in MODES:
            raise ValueError('Mode tidak valid')
        master_service=SharedMasterService(self.db_path)
        master = master_service.snapshot_map(periode)
        missing_profiles={nip for nip,row in master.items() if row.get('tukin_gross') is None}
        if missing_profiles:
            latest=master_service.snapshot_map(periode,fallback_latest_profile=True)
            for nip in missing_profiles:
                if latest.get(nip,{}).get('tukin_gross') is not None:
                    master[nip]=latest[nip]
                    master[nip]['_period_profile_fallback']=True
        with duckdb.connect(self.db_path) as c:
            attendance = records(c,'rekap_sdm',{'bulan':periode})
            flags = records(c,'workflow_flags')
            verified = any(r.get('k') == 'semifinal:'+periode and r.get('verified') == 1 for r in flags)
            for config in records(c,'workflow_config'):
                workflow=config.get('_workflow')
                if isinstance(workflow,str):
                    try: workflow=json.loads(workflow)
                    except ValueError: workflow={}
                if isinstance(workflow,dict) and workflow.get('month')==periode and workflow.get('semifinalVerified') is False:
                    verified=False
            for row in attendance:
                row['_unverified'] = not verified
            skp = []
            if mode == 'ABSENSI_SKP':
                if source_type not in SOURCES:
                    raise ValueError('Pilih sumber SKP')
                if source_type == 'ASSESSMENT':
                    history = {r['nip']:r for r in records(c,'tukin_skp_monthly',{'assessment_period':periode})}
                    conditions = records(c,'skp_condition_decision',{'assessment_period':periode,'status':'APPROVED'})
                    rules={}
                    assessment_rows = annotate_skp_tukin_provenance(
                        c, periode, records(c,'skp_assessment_monthly',{'assessment_period':periode})
                    )
                    for r in assessment_rows:
                        predicate = str(r.get('final_predicate') or '').upper()
                        if predicate not in rules: rules[predicate]=self.rule_rate(c,predicate,periode)
                        rate = rules[predicate]
                        previous = history.get(r['nip'],{})
                        decisions=[d for d in conditions if d.get('nip')==r['nip']]
                        payment_eligible=False if any(d.get('payment_effect') in {'HOLD','UNPAID'} for d in decisions) else previous.get('payment_eligible')
                        if r.get('report_status') != 'SUBMITTED':
                            rate = previous.get('skp_reduction_rate')
                        direct = bool(r.get('_direct_final_import'))
                        skp.append(dict(id=str(uuid.uuid5(uuid.NAMESPACE_URL,signature([r,previous,rate,decisions]))),nip=r['nip'],periode=periode,
                            predikat_skp=predicate,potongan_skp_percent=rate,status_skp=r['review_status'],
                            source_type='DIRECT_FINAL_IMPORT' if direct else 'ASSESSMENT',
                            payment_eligible=payment_eligible,assessment=r, history=previous,decisions=decisions))
                else:
                    if not skp_batch_id:
                        raise ValueError('Pilih batch import SKP tersimpan')
                    skp = records(c,'skp_results',{'periode':periode,'source_type':source_type,'import_batch_id':skp_batch_id})
        return master, attendance, skp

    def preview(self, periode, mode='ABSENSI_ONLY', source_type=None, skp_batch_id=None):
        master, attendance, skp = self.sources(periode,mode,source_type,skp_batch_id)
        ac = Counter(str(r['nip']) for r in attendance)
        sc = Counter(str(r['nip']) for r in skp)
        am = {str(r['nip']):r for r in attendance}
        sm = {str(r['nip']):r for r in skp}
        rows = []
        for nip in sorted(set(am)|set(sm)):
            a = am.get(nip)
            s = sm.get(nip)
            if a and ac[nip]>1: a = {**a,'_duplicate':True}
            if s and sc[nip]>1: s = {**s,'_duplicate':True}
            row = calculate_final_tukin(nip,periode,mode,a,master.get(nip),s)
            row['source_signature'] = signature([master.get(nip),a,s])
            row['tukin_bruto_source']='MASTER SDM TERKINI' if master.get(nip,{}).get('_period_profile_fallback') else 'MASTER SDM PERIODE'
            warnings=[]
            if master.get(nip,{}).get('_period_profile_fallback'):
                warnings.append('Histori bruto periode belum tersedia; memakai Master SDM terkini sesuai workflow lama')
            if s and s.get('nama') and s['nama']!=row['nama']: warnings.append('Nama SKP berbeda dari Master SDM')
            row['warning']='; '.join(warnings)
            rows.append(row)
        stats = dict(rekap_sdm=len(attendance),data_skp=len(skp),match=len(set(am)&set(sm)),
                     skp_tidak_ada=len(set(am)-set(sm)) if mode=='ABSENSI_SKP' else 0,
                     absensi_tidak_ada=len(set(sm)-set(am)),duplikat=sum(v>1 for v in sc.values()),
                     ready=sum(r['status']=='READY' for r in rows),hold=sum(r['status']=='HOLD' for r in rows))
        return dict(periode=periode,mode=mode,source_type=source_type if mode=='ABSENSI_SKP' else None,
                    skp_batch_id=skp_batch_id if mode=='ABSENSI_SKP' else None,rows=rows,stats=stats,
                    preview_signature=signature(rows))

    def finalize(self, periode, mode, preview_signature, source_type=None, skp_batch_id=None, created_by='User'):
        preview = self.preview(periode,mode,source_type,skp_batch_id)
        if preview_signature != preview['preview_signature']:
            raise ValueError('Data sumber berubah; preview kembali sebelum Final')
        if not preview['rows'] or not preview['stats']['ready']:
            raise ValueError('Belum ada SDM siap Final')
        batch = str(uuid.uuid4())
        assessment_rows = self.sources(periode,mode,source_type,skp_batch_id)[2] if mode=='ABSENSI_SKP' and source_type=='ASSESSMENT' else []
        with duckdb.connect(self.db_path) as c:
            c.execute('BEGIN')
            version = c.execute('SELECT coalesce(max(version),0)+1 FROM final_tukin_batches WHERE periode=? AND mode=?',[periode,mode]).fetchone()[0]
            ready = [r for r in preview['rows'] if r['status']=='READY']
            insert(c,'final_tukin_batches',dict(id=batch,periode=periode,mode=mode,version=version,skp_source=preview['source_type'],
                skp_batch_id=preview['skp_batch_id'],jumlah_sdm=len(preview['rows']),source_signature=preview_signature,created_by=created_by,
                total_tukin_bruto=sum(r['tukin_bruto'] for r in ready),total_potongan_absen=sum(r['potongan_absen_rp'] for r in ready),
                total_potongan_skp=sum(r['potongan_skp_rp'] or 0 for r in ready),total_tukin_dibayar=sum(r['tukin_dibayar'] for r in ready)))
            columns = {r[0] for r in c.execute('DESCRIBE final_tukin_results').fetchall()}
            final_rows=[]
            for row in preview['rows']:
                row.update(id=str(uuid.uuid4()),calculation_batch_id=batch,version=version,
                           status='FINALIZED' if row['status']=='READY' else 'HOLD')
                saved = {k:v for k,v in row.items() if k in columns}
                saved['snapshot'] = json.dumps(row,default=str,ensure_ascii=False)
                final_rows.append(saved)
            insert_many(c,'final_tukin_results',final_rows)
            # Assessment normalization is persisted with this batch, preserving source provenance.
            if mode=='ABSENSI_SKP' and source_type=='ASSESSMENT':
                for r in assessment_rows:
                    if c.execute('SELECT count(*) FROM skp_results WHERE id=?',[r['id']]).fetchone()[0]:
                        continue
                    assessment = r['assessment']
                    insert(c,'skp_results',dict(id=r['id'],nip=r['nip'],periode=periode,tahun=int(periode[:4]),bulan=int(periode[5:]),
                        predikat_skp=r['predikat_skp'],potongan_skp_percent=r['potongan_skp_percent'],status_skp=r['status_skp'],source_type='ASSESSMENT',
                        assessment_id=periode+':'+r['nip'],import_batch_id=batch,skp_percentage_source='RULE_MAPPING',payload=json.dumps(r,default=str),
                        source_file=assessment.get('source_file'),reviewer=assessment.get('reviewed_by'),approver=assessment.get('approved_by'),
                        evidence=assessment.get('evidence_ref'),notes=assessment.get('reviewer_notes')))
            c.execute('COMMIT')
        return dict(calculation_batch_id=batch,version=version,stats=preview['stats'])

    def batches(self):
        with duckdb.connect(self.db_path) as c:
            return sorted([r for r in records(c,'final_tukin_batches') if r['archived_at'] is None],key=lambda r:r['created_at'],reverse=True)

    def results(self, batch_id):
        with duckdb.connect(self.db_path) as c:
            batches = [r for r in records(c,'final_tukin_batches',{'id':batch_id}) if r['archived_at'] is None]
            if not batches: raise ValueError('Batch tidak ditemukan atau diarsipkan')
            batch = batches[0]
            saved = records(c,'final_tukin_results',{'calculation_batch_id':batch_id})
        current = self.preview(batch['periode'],batch['mode'],batch['skp_source'],batch['skp_batch_id'])
        current_map = {r['nip']:r for r in current['rows']}
        output = []
        outdated_ids=[]
        with duckdb.connect(self.db_path) as c:
            for saved_row in saved:
                row = json.loads(saved_row['snapshot'])
                if saved_row['status']=='OUTDATED' or current_map.get(row['nip'],{}).get('source_signature') != saved_row['source_signature']:
                    row.update(status='OUTDATED',reason='SOURCE_CHANGED',status_label='PERLU FINALISASI ULANG')
                    if saved_row['status']!='OUTDATED': outdated_ids.append(row['id'])
                output.append(row)
            for start in range(0,len(outdated_ids),500):
                ids=outdated_ids[start:start+500]
                c.execute("UPDATE final_tukin_results SET status='OUTDATED',reason='SOURCE_CHANGED',updated_at=now() WHERE id IN ("+','.join('?' for _ in ids)+')',ids)
        return dict(batch=batch,rows=output,source_changed=current['preview_signature']!=batch['source_signature'])

    def calculation_results(self, calculation_id):
        """Read a canonical calculation snapshot without recalculating it."""
        with duckdb.connect(self.db_path) as c:
            ensure_integration_schema(c)
            batch = c.execute(
                """SELECT id,assessment_period,version,calculation_version,status,source_fingerprint,
                          mismatch_count,ready_count,calculated_by,calculated_at
                   FROM tukin_calculation_batch WHERE id=?""", [calculation_id]
            ).fetchone()
            if not batch:
                raise ValueError('Calculation snapshot tidak ditemukan')
            rows = c.execute(
                """SELECT snapshot,status,tukin_gross,attendance_deduction,skp_deduction,
                          adjustment,total_deduction,tukin_paid,attendance_source_id,
                          attendance_batch_id,skp_source_id,skp_source_version,skp_batch_id,
                          master_version,calculated_at
                   FROM tukin_calculation_result WHERE calculation_id=? ORDER BY nip""",
                [calculation_id],
            ).fetchall()
        keys = ['calculation_id','assessment_period','version','calculation_version','status',
                'source_fingerprint','mismatch_count','ready_count','calculated_by','calculated_at']
        row_keys = ['snapshot_status','tukin_gross','attendance_deduction','skp_deduction',
                    'adjustment','total_deduction','tukin_paid','attendance_source_id',
                    'attendance_batch_id','skp_source_id','skp_source_version','skp_batch_id',
                    'master_version','calculated_at']
        output = []
        for stored in rows:
            row = json.loads(stored[0])
            row.update({key: value for key, value in zip(row_keys, stored[1:])})
            # Stable canonical aliases consumed by Final UI/export. No arithmetic occurs here.
            row.update(
                periode=batch[1],
                tukin_bruto=stored[2],
                komponen_absensi=row.get('attendance_component'),
                potongan_absensi=stored[3],
                komponen_skp=row.get('skp_component'),
                potongan_skp=stored[4],
                adjustment=stored[5],
                total_potongan=stored[6],
                tukin_dibayar=stored[7],
            )
            output.append(row)
        batch_data = dict(zip(keys, batch))
        scope = output[0].get('scope', 'all') if output else 'all'
        batch_data['scope'] = scope
        current = integration_readiness(self.db_path, batch_data['assessment_period'], scope)
        batch_data['source_status'] = (
            'CURRENT' if current['source_fingerprint'] == batch_data['source_fingerprint'] else 'STALE'
        )
        return {'batch': batch_data, 'rows': output}

    def calculations(self, assessment_period=''):
        if assessment_period:
            period(assessment_period)
        with duckdb.connect(self.db_path) as c:
            ensure_integration_schema(c)
            query = """SELECT id calculation_id,assessment_period,version,calculation_version,status,
                              source_fingerprint,mismatch_count,ready_count,calculated_by,calculated_at
                       FROM tukin_calculation_batch"""
            params = []
            if assessment_period:
                query += ' WHERE assessment_period=?'
                params.append(assessment_period)
            query += ' ORDER BY assessment_period DESC,version DESC'
            cursor = c.execute(query, params)
            names = [item[0] for item in cursor.description]
            rows = [dict(zip(names, values)) for values in cursor.fetchall()]
            for row in rows:
                row['scope'] = calculation_scope(c, row['calculation_id'])
        readiness = {}
        for row in rows:
            p = row['assessment_period']
            key = (p, row['scope'])
            if key not in readiness:
                readiness[key] = integration_readiness(self.db_path, p, row['scope'])
            row['source_status'] = (
                'CURRENT' if row['source_fingerprint'] == readiness[key]['source_fingerprint'] else 'STALE'
            )
        return rows

    def approve_calculation(self, calculation_id, approved_by, reason):
        """Lock Final Tukin to one immutable calculation snapshot."""
        approved_by = str(approved_by or '').strip()
        reason = str(reason or '').strip()
        if not approved_by or not reason:
            raise ValueError('Approver dan alasan persetujuan wajib diisi')
        snapshot = self.calculation_results(calculation_id)
        batch = snapshot['batch']
        if batch['status'] != 'CALCULATED':
            raise ValueError('Hanya calculation berstatus CALCULATED yang dapat disetujui')
        if not batch['ready_count']:
            raise ValueError('Calculation snapshot tidak memiliki hasil siap Final')
        if int(batch.get('mismatch_count') or 0) > 0:
            raise ValueError(
                f"Finalisasi diblokir: {batch['mismatch_count']} NIP/data sumber masih perlu diperbaiki"
            )
        if batch['source_status'] == 'STALE':
            raise ValueError('Sumber berubah setelah perhitungan; buat calculation version baru')
        decision_id = str(uuid.uuid4())
        with duckdb.connect(self.db_path) as c:
            ensure_integration_schema(c)
            c.execute('BEGIN TRANSACTION')
            try:
                existing = c.execute(
                    """SELECT calculation_id,status FROM tukin_final_decision
                       WHERE assessment_period=? ORDER BY created_at DESC LIMIT 1""",
                    [batch['assessment_period']],
                ).fetchone()
                if existing and existing[0] == calculation_id:
                    raise ValueError('Calculation ini sudah memiliki Final Tukin')
                if existing and existing[1] == 'PAID':
                    raise ValueError('Periode sudah dibayar; Final baru tidak dapat dibuat')
                if existing:
                    previous = c.execute(
                        'SELECT version FROM tukin_calculation_batch WHERE id=?', [existing[0]]
                    ).fetchone()
                    if previous and int(previous[0]) >= int(batch['version']):
                        raise ValueError('Final baru harus menunjuk calculation version yang lebih baru')
                c.execute(
                    """INSERT INTO tukin_final_decision
                       (id,assessment_period,calculation_id,calculation_version,status,
                        approved_by,approved_at,locked_at,reason,created_at,updated_at)
                       VALUES (?,?,?,?, 'FINAL',?,now(),now(),?,now(),now())""",
                    [decision_id, batch['assessment_period'], calculation_id,
                     int(batch['version']), approved_by, reason],
                )
                c.execute('COMMIT')
            except Exception:
                c.execute('ROLLBACK')
                raise
        return {'status': 'FINAL', 'final_decision_id': decision_id,
                'calculation_id': calculation_id, 'calculation_version': batch['version'],
                'assessment_period': batch['assessment_period']}

    def final_decisions(self, assessment_period=''):
        with duckdb.connect(self.db_path) as c:
            ensure_integration_schema(c)
            query = """SELECT d.*,b.status calculation_status,b.calculation_version engine_version,
                              b.ready_count,b.mismatch_count,b.source_fingerprint,b.calculated_at
                       FROM tukin_final_decision d JOIN tukin_calculation_batch b ON b.id=d.calculation_id"""
            params = []
            if assessment_period:
                period(assessment_period)
                query += ' WHERE d.assessment_period=?'
                params.append(assessment_period)
            query += ' ORDER BY d.created_at DESC,d.calculation_version DESC'
            cursor = c.execute(query, params)
            names = [item[0] for item in cursor.description]
            rows = [dict(zip(names, values)) for values in cursor.fetchall()]
            for row in rows:
                row['scope'] = calculation_scope(c, row['calculation_id'])
        by_period = {}
        readiness = {}
        for row in rows:
            p = row['assessment_period']
            by_period.setdefault(p, row['id'])
            key = (p, row['scope'])
            if key not in readiness:
                readiness[key] = integration_readiness(self.db_path, p, row['scope'])
            row['is_current_decision'] = row['id'] == by_period[p]
            row['source_status'] = (
                'CURRENT' if row['source_fingerprint'] == readiness[key]['source_fingerprint'] else 'STALE'
            )
            row['display_status'] = (
                row['status'] if row['is_current_decision'] else f"{row['status']} (SUPERSEDED)"
            )
        return rows

    def final_decision(self, decision_id):
        decisions = [row for row in self.final_decisions() if row['id'] == decision_id]
        if not decisions:
            raise ValueError('Final Tukin tidak ditemukan')
        decision = decisions[0]
        snapshot = self.calculation_results(decision['calculation_id'])
        return {'decision': decision, 'batch': snapshot['batch'], 'rows': snapshot['rows']}

    def mark_paid(self, decision_id, paid_by, payment_reference):
        paid_by = str(paid_by or '').strip()
        payment_reference = str(payment_reference or '').strip()
        if not paid_by or not payment_reference:
            raise ValueError('Petugas dan referensi pembayaran wajib diisi')
        with duckdb.connect(self.db_path) as c:
            ensure_integration_schema(c)
            c.execute('BEGIN TRANSACTION')
            try:
                decision = c.execute(
                    'SELECT status,assessment_period FROM tukin_final_decision WHERE id=?', [decision_id]
                ).fetchone()
                if not decision:
                    raise ValueError('Final Tukin tidak ditemukan')
                if decision[0] == 'PAID':
                    raise ValueError('Final Tukin sudah dibayar')
                if decision[0] != 'FINAL':
                    raise ValueError('Hanya Final Tukin terkunci yang dapat dibayar')
                latest = c.execute(
                    """SELECT id FROM tukin_final_decision WHERE assessment_period=?
                       ORDER BY created_at DESC,calculation_version DESC LIMIT 1""", [decision[1]]
                ).fetchone()
                if not latest or latest[0] != decision_id:
                    raise ValueError('Final Tukin ini sudah digantikan calculation version yang lebih baru')
                c.execute(
                    """UPDATE tukin_final_decision SET status='PAID',paid_by=?,paid_at=now(),
                       payment_reference=?,updated_at=now() WHERE id=? AND status='FINAL'""",
                    [paid_by, payment_reference, decision_id],
                )
                c.execute('COMMIT')
            except Exception:
                c.execute('ROLLBACK')
                raise
        return {'status': 'PAID', 'final_decision_id': decision_id}

    def export_final_decision(self, decision_id):
        """Export only a snapshot explicitly locked as FINAL or PAID."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        result = self.final_decision(decision_id)
        decision = result['decision']
        if decision['status'] not in {'FINAL', 'PAID'}:
            raise ValueError('Hanya Final Tukin terkunci/dibayar yang dapat diekspor untuk pembayaran')
        wb = Workbook()
        metadata = wb.active
        metadata.title = 'FINAL DECISION'
        metadata.append(['Field', 'Value'])
        for key in ('id','assessment_period','calculation_id','calculation_version','status',
                    'approved_by','approved_at','locked_at','paid_by','paid_at',
                    'payment_reference','reason','source_status','calculated_at'):
            value = decision.get(key)
            if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                value = "'" + value
            metadata.append([key, value])
        rows_sheet = wb.create_sheet('FINAL TUKIN')
        columns = [
            ('NIP','nip'),('Nama','nama'),('Jabatan','jabatan'),('Unit/Wilayah','unit'),
            ('Flag 3T/Non-3T','wilayah_3t'),('Tukin Bruto','tukin_bruto'),
            ('Predikat SKP','predicate_skp'),('Potongan SKP %','skp_rate'),
            ('Komponen SKP 70%','komponen_skp'),('Potongan SKP Rp','potongan_skp'),
            ('Potongan Absensi %','attendance_rate'),('Komponen Absensi 30%','komponen_absensi'),
            ('Potongan Absensi Rp','potongan_absensi'),('Total Potongan Rp','total_potongan'),
            ('Hak Tukin Final','tukin_dibayar'),('Status','payment_status'),('Keterangan','issues'),
            ('tukin_dibayar','tukin_dibayar'),('calculation_id','calculation_id'),
            ('version','version'),('snapshot_status','snapshot_status'),('calculated_at','calculated_at'),
            ('attendance_source_id','attendance_source_id'),('attendance_batch_id','attendance_batch_id'),
            ('skp_source_id','skp_source_id'),('skp_source_version','skp_source_version'),
            ('skp_batch_id','skp_batch_id'),('master_version','master_version'),
            # Appended audit columns keep the established Final export stable.
            ('Sumber SKP','skp_source_label'),('Status Sumber SKP','skp_source_status'),
            ('Mode Perhitungan SKP','skp_calculation_mode'),
            ('Kebijakan Fallback SKP','skp_fallback_policy'),('Keterangan SKP','skp_description'),
            ('Status Workflow SKP','skp_workflow_status'),('SKP Normal Siap Tukin','skp_ready_for_tukin'),
        ]
        rows_sheet.append([label for label, _ in columns])
        for row in result['rows']:
            rows_sheet.append([row.get(key) for _, key in columns])
        for sheet in (metadata, rows_sheet):
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17365D')
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    def legacy_preview_adapter(self, periode):
        """Deprecated read-only adapter: latest canonical snapshot, never recalculated."""
        calculations = self.calculations(periode)
        if not calculations:
            raise ValueError('Belum ada calculation snapshot canonical untuk periode ini')
        result = self.calculation_results(calculations[0]['calculation_id'])
        result['deprecated'] = True
        result['replacement'] = '/api/final-tukin/calculation/{calculation_id}'
        result['preview_signature'] = calculations[0]['calculation_id']
        return result

    def legacy_finalize_adapter(self, payload):
        """Deprecated approval adapter. The signature is a canonical calculation id."""
        calculation_id = payload.get('calculation_id') or payload.get('preview_signature')
        result = self.approve_calculation(
            calculation_id, payload.get('approved_by') or payload.get('created_by'),
            payload.get('reason'),
        )
        return {**result, 'calculation_batch_id': calculation_id, 'deprecated': True,
                'replacement': '/api/final-tukin/approve-calculation'}

    def legacy_export_adapter(self, identifier):
        """Resolve a legacy export identifier to its canonical Final Decision."""
        decisions = self.final_decisions()
        decision = next((row for row in decisions if row['id'] == identifier), None)
        if decision is None:
            decision = next((row for row in decisions if row['calculation_id'] == identifier), None)
        if decision is None:
            raise ValueError('Export legacy hanya tersedia untuk calculation yang sudah menjadi Final canonical')
        return self.export_final_decision(decision['id'])

    def archive(self,batch_id):
        with duckdb.connect(self.db_path) as c:
            c.execute('UPDATE final_tukin_batches SET archived_at=now() WHERE id=? AND archived_at IS NULL',[batch_id])
        return {'archived':batch_id}

    def export(self,batch_id):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        result = self.results(batch_id)
        rows = result['rows']
        wb = Workbook(); wb.remove(wb.active)
        def sheet(name, data, columns=None):
            ws = wb.create_sheet(name)
            keys = columns or (list(data[0]) if data else ['Informasi'])
            ws.append(keys)
            for row in data:
                values = []
                for key in keys:
                    value = row.get(key)
                    if key.endswith('_percent') and value is not None: value=float(Decimal(str(value)))
                    if isinstance(value,datetime): value=value.isoformat()
                    if isinstance(value,Decimal): value=str(value)
                    if isinstance(value,str) and value.startswith(('=','+','-','@')): value="'"+value
                    values.append(value)
                ws.append(values)
            ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='17365D')
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width=min(38,max(16,len(str(column[0].value))+2))
                for cell in column[1:]:
                    if isinstance(cell.value,int): cell.number_format='#,##0'
        cols=['nip','nama','provinsi','kabupaten_kota','kecamatan','jabatan','periode','tukin_bruto','komponen_absen_30',
              'potongan_absen_percent','potongan_absen_rp','predikat_skp','komponen_skp_70','potongan_skp_percent','potongan_skp_rp',
              'total_potongan','tukin_dibayar','mode','skp_source','status','reason','version','tukin_bruto_source','warning']
        sheet('BNBA FINAL',rows,cols)
        metrics=['tukin_bruto','potongan_absen_rp','potongan_skp_rp','total_potongan','tukin_dibayar']
        summaries=[]
        for group in ['SEMUA','provinsi','kabupaten_kota','predikat_skp','mode']:
            groups={}
            for row in rows:
                key='SEMUA' if group=='SEMUA' else str(row.get(group) or 'TIDAK DIGUNAKAN')
                if group=='kabupaten_kota': key=str(row.get('provinsi') or '—')+' / '+key
                groups.setdefault(key,[]).append(row)
            for key,items in groups.items():
                payable=[r for r in items if r['tukin_dibayar'] is not None]
                summaries.append(dict(pengelompokan=group,nilai=key,jumlah_sdm=len(items),hold=sum(r['status']=='HOLD' for r in items),
                                      **{m:sum(r.get(m) or 0 for r in payable) for m in metrics}))
        sheet('REKAP',summaries)
        sheet('DATA BERMASALAH',[r for r in rows if r['status']!='FINALIZED' or r.get('warning')],cols)
        sheet('METADATA PERHITUNGAN',[dict(field=k,value=v) for k,v in {**result['batch'],
            'sumber_absensi':'rekap_sdm (verifikasi Semi Final)','rounding':'ROUND_HALF_UP; 70% = bruto - komponen 30%; potongan dari komponen integer',
            'source_changed':result['source_changed']}.items()])
        stream=io.BytesIO(); wb.save(stream); stream.seek(0)
        return stream


def create_router(db_path_getter):
    router=APIRouter(prefix='/api/final-tukin',tags=['final-tukin'])
    def svc(): return FinalTukinService(db_path_getter())
    def run(action):
        try: return action()
        except ValueError as exc: raise HTTPException(400,str(exc)) from exc

    @router.post('/skp/preview')
    async def import_preview(file: UploadFile=File(...), periode: str=Form(...), source_type: str=Form(...),
                             sheet: str=Form(''), mapping: str=Form(''), header_row: int=Form(-1), percent_unit: str=Form('POINTS')):
        content=await file.read()
        result=run(lambda:svc().preview_import(content,file.filename,periode,source_type,sheet or None,
                    json.loads(mapping) if mapping else None,header_row if header_row>=0 else None,percent_unit))
        return {**result,'rows':result['rows'][:100], 'issues':[
            {k:r.get(k) for k in ('source_row','nip','periode','predikat_skp','potongan_skp_percent','errors')}
            for r in result['rows'] if r['errors']]}

    @router.post('/skp/import')
    async def import_commit(file: UploadFile=File(...), periode: str=Form(...), source_type: str=Form(...),
                            sheet: str=Form(...), mapping: str=Form(...), header_row: int=Form(...),
                            preview_signature: str=Form(...), percent_unit: str=Form('POINTS'), valid_only: bool=Form(False)):
        content=await file.read()
        service=svc()
        return run(lambda:service.import_preview(service.preview_import(content,file.filename,periode,source_type,sheet,
                    json.loads(mapping),header_row,percent_unit),preview_signature,valid_only=valid_only))

    @router.get('/skp/batches')
    def import_batches():
        service=svc()
        with duckdb.connect(service.db_path) as c: return records(c,'skp_import_batches')

    @router.post('/preview', deprecated=True)
    def preview(payload: dict):
        return run(lambda:svc().legacy_preview_adapter(str(payload.get('periode') or '')))

    @router.post('/finalize', deprecated=True)
    def finalize(payload: dict): return run(lambda:svc().legacy_finalize_adapter(payload))

    @router.get('/batches', deprecated=True)
    def batches():
        return {'rows': svc().calculations(), 'deprecated': True,
                'replacement': '/api/tukin/combined/calculations?assessment_period=YYYY-MM'}

    @router.get('/results/{batch_id}', deprecated=True)
    def results(batch_id: str):
        result = run(lambda:svc().calculation_results(batch_id))
        return {**result, 'deprecated': True,
                'replacement': '/api/final-tukin/calculation/{calculation_id}'}

    @router.get('/calculations')
    def calculations(assessment_period: str = ''):
        return {'rows': run(lambda:svc().calculations(assessment_period))}

    @router.get('/calculation/{calculation_id}')
    def calculation(calculation_id: str):
        return run(lambda:svc().calculation_results(calculation_id))

    @router.post('/approve-calculation')
    def approve_calculation(payload: dict):
        return run(lambda:svc().approve_calculation(
            payload.get('calculation_id'), payload.get('approved_by'), payload.get('reason')
        ))

    @router.get('/decisions')
    def decisions(assessment_period: str = ''):
        return {'rows': run(lambda:svc().final_decisions(assessment_period))}

    @router.get('/decisions/{decision_id}')
    def decision(decision_id: str):
        return run(lambda:svc().final_decision(decision_id))

    @router.post('/decisions/{decision_id}/paid')
    def mark_paid(decision_id: str, payload: dict):
        return run(lambda:svc().mark_paid(
            decision_id, payload.get('paid_by'), payload.get('payment_reference')
        ))

    @router.get('/decisions/{decision_id}/export')
    def export_decision(decision_id: str):
        stream = run(lambda:svc().export_final_decision(decision_id))
        return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 headers={'Content-Disposition':'attachment; filename="FINAL_TUKIN_LOCKED.xlsx"'})

    @router.post('/archive/{batch_id}', deprecated=True)
    def archive(batch_id: str,payload: dict):
        raise HTTPException(410, 'Arsip Final legacy dinonaktifkan; calculation snapshot canonical immutable')

    @router.get('/export/{batch_id}', deprecated=True)
    def export(batch_id: str):
        stream=run(lambda:svc().legacy_export_adapter(batch_id))
        return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 headers={'Content-Disposition':'attachment; filename="EVKIN_Final_Tukin_CANONICAL.xlsx"',
                                          'Deprecation':'true'})
    return router

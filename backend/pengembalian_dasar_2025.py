"""Importer khusus dasar kewajiban dan progress Tukin SP2D 2025."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import re
import uuid

import duckdb
from openpyxl import load_workbook

from backend.pengembalian_archive import ArchiveService
from backend.pengembalian_progress import ProgressService, _stage_number, normalize_bank


SOURCE_TYPE = "DASAR_PENGEMBALIAN_2025"
PERIODE_ID = "DASAR-PENGEMBALIAN-2025"
YEAR = 2025
COMPONENT = "TUKIN"
REQUIRED = {"nip", "nama", "provinsi", "kabupaten_kota", "tahun_sp2d",
            "nama_bank", "nomor_rekening", "final_oktober", "final_november",
            "final_desember", "final_total_pengembalian"}


def _key(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _text(value):
    text = str(value or "").strip()
    return text or None


def _nip(value):
    text = _text(value)
    if not text:
        return None
    text = re.sub(r"\.0$", "", text)
    return text if re.fullmatch(r"\d{8,24}", text) else None


def _rupiah(value):
    if value is None or value == "":
        return 0, None
    if isinstance(value, bool):
        return None, "NOMINAL_MALFORMED"
    try:
        if isinstance(value, (int, float, Decimal)):
            number = Decimal(str(value))
        else:
            text = re.sub(r"[^0-9,.-]", "", str(value).strip())
            if not text:
                return None, "NOMINAL_MALFORMED"
            if "," in text and "." in text:
                text = text.replace(".", "").replace(",", ".")
            elif "," in text:
                text = text.replace(",", ".")
            number = Decimal(text)
    except (InvalidOperation, ValueError):
        return None, "NOMINAL_MALFORMED"
    if number != number.to_integral_value():
        return None, "NOMINAL_BUKAN_INTEGER"
    result = int(number)
    return (None, "NOMINAL_NEGATIF") if result < 0 else (result, None)


class DasarPengembalian2025Service:
    def __init__(self, db_path):
        self.db_path = str(db_path)
        ArchiveService(self.db_path).repository.ensure_schema()
        ProgressService(self.db_path)

    @staticmethod
    def _read(content):
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            candidates = []
            for sheet in workbook.worksheets:
                for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row), values_only=True), 1):
                    headers = [_key(value) for value in values]
                    score = len(REQUIRED.intersection(headers))
                    if "nip" in headers and score >= 8:
                        candidates.append((score, -row_number, sheet, row_number, headers))
            if not candidates:
                raise ValueError("HEADER_DASAR_PENGEMBALIAN_2025_TIDAK_DITEMUKAN")
            _, _, sheet, header_row, headers = max(candidates, key=lambda item: item[:2])
            rows = []
            for source_row, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({headers[index]: values[index] if index < len(values) else None
                             for index in range(len(headers))} | {"_source_row": source_row})
            return sheet.title, headers, rows
        finally:
            workbook.close()

    @staticmethod
    def _stage_columns(headers):
        stages = {}
        invalid = []
        for index, header in enumerate(headers):
            if "tahap" not in header:
                continue
            match = re.search(r"tahap_([ivxlcdm]+|\d+)", header)
            stage = _stage_number(match.group(1)) if match else None
            if not stage:
                invalid.append(header)
                continue
            field = None
            if header.startswith("send_"): field = "send"
            elif header.startswith("debit_"): field = "debit"
            elif header.startswith("ntpn_"): field = "ntpn"
            elif header.startswith("gagal_"): field = "gagal"
            elif header.startswith("keterangan_gagal_"): field = "alasan"
            if field:
                stages.setdefault(stage, {})[field] = index
        return stages, invalid

    def preview(self, content, filename="dasar-2025.xlsx", file_hash=None):
        digest = file_hash or sha256(content).hexdigest()
        sheet, headers, source_rows = self._read(content)
        stages, invalid_stage_headers = self._stage_columns(headers)
        normalized, errors, warnings = [], [], []
        seen = Counter()
        stage_counts = Counter()
        totals = Counter()
        for source in source_rows:
            if not _text(source.get("nip")) and not any(_text(source.get(field)) for field in
                ("nama", "nama_bank", "tahun_sp2d", "nomor_rekening")):
                continue  # workbook summary/control row, not an SDM source row
            row_number = source["_source_row"]
            row_errors, row_warnings = [], []
            nip = _nip(source.get("nip"))
            if not nip: row_errors.append("NIP_KOSONG_ATAU_INVALID")
            year_value, year_error = _rupiah(source.get("tahun_sp2d"))
            if year_error or year_value != YEAR: row_errors.append("TAHUN_SP2D_BUKAN_2025")
            amounts = {}
            for field in ("final_oktober", "final_november", "final_desember", "final_total_pengembalian"):
                amounts[field], error = _rupiah(source.get(field))
                if error: row_errors.append(f"{field.upper()}_{error}")
            monthly_total = sum(amounts.get(field) or 0 for field in ("final_oktober", "final_november", "final_desember"))
            if amounts.get("final_total_pengembalian") is not None and monthly_total != amounts["final_total_pengembalian"]:
                row_errors.append("FINAL_TOTAL_TIDAK_COCOK")
            if nip: seen[nip] += 1
            progress = []
            for stage, mapping in sorted(stages.items()):
                values = {}
                for field in ("send", "debit", "gagal"):
                    raw = source.get(headers[mapping[field]]) if field in mapping else 0
                    values[field], error = _rupiah(raw)
                    if error: row_errors.append(f"TAHAP_{stage}_{field.upper()}_{error}")
                values["ntpn"] = _text(source.get(headers[mapping["ntpn"]])) if "ntpn" in mapping else None
                values["alasan"] = _text(source.get(headers[mapping["alasan"]])) if "alasan" in mapping else None
                send, debit, gagal = values.get("send") or 0, values.get("debit") or 0, values.get("gagal") or 0
                if debit + gagal > send and send:
                    row_warnings.append(f"TAHAP_{stage}_REALISASI_DAN_GAGAL_MELEBIHI_SEND")
                if send or debit or gagal:
                    stage_counts[stage] += 1
                    progress.append({"tahap": stage, **values, "pending": max(send - debit - gagal, 0)})
            reference_progress, ref_error = _rupiah(source.get("total_progress_2025"))
            reference_remaining, remaining_error = _rupiah(source.get("sisa_kewajiban"))
            calculated_progress = sum(item["debit"] or 0 for item in progress)
            calculated_remaining = max(monthly_total - calculated_progress, 0)
            if not ref_error and reference_progress != calculated_progress: row_warnings.append("TOTAL_PROGRESS_REFERENSI_TIDAK_COCOK")
            if not remaining_error and reference_remaining != calculated_remaining: row_warnings.append("SISA_REFERENSI_TIDAK_COCOK")
            normalized.append({"source_row": row_number, "nip": nip, "nama": _text(source.get("nama")),
                "provinsi": _text(source.get("provinsi")), "kabupaten_kota": _text(source.get("kabupaten_kota")),
                "tahun_sp2d": year_value, "nama_bank": _text(source.get("nama_bank")),
                "nomor_rekening": _text(source.get("nomor_rekening")), "status_master": _text(source.get("status_master")),
                "batch_bpk": _text(source.get("batch_bpk")), "kode_anak_satker": _text(source.get("kode_anak_satker")),
                **amounts, "calculated_total": monthly_total, "progress": progress,
                "reference": {"total_progress": reference_progress, "sisa": reference_remaining,
                              "status": _text(source.get("status_progress"))},
                "raw": {key: value for key, value in source.items() if key != "_source_row"},
                "errors": row_errors, "warnings": row_warnings})
            for field, value in amounts.items(): totals[field] += value or 0
        duplicates = {nip for nip, count in seen.items() if count > 1}
        for row in normalized:
            if row["nip"] in duplicates: row["errors"].append("DUPLICATE_NIP")
            if row["errors"]: errors.append({"row": row["source_row"], "nip": row["nip"], "messages": row["errors"]})
            if row["warnings"]: warnings.append({"row": row["source_row"], "nip": row["nip"], "messages": row["warnings"]})
        if invalid_stage_headers:
            errors.append({"row": None, "messages": ["TAHAP_HEADER_INVALID: " + ", ".join(invalid_stage_headers)]})
        conn = duckdb.connect(self.db_path, read_only=True)
        try:
            duplicate = conn.execute("""SELECT import_batch_id,imported_at FROM pengembalian_progress_import_batch
                WHERE source_file_hash=? AND status IN ('COMMITTED','PARTIAL') LIMIT 1""", [digest]).fetchone()
        finally:
            conn.close()
        states = Counter()
        total_realisasi = 0
        for row in normalized:
            paid = sum(item["debit"] or 0 for item in row["progress"])
            total_realisasi += paid
            obligation = row["calculated_total"]
            states["overpayment" if paid > obligation else "exact" if paid == obligation else "receivable"] += 1
        summary = {"jumlah_baris": len(normalized), "jumlah_nip_unik": len(seen), "tahun_sp2d": [YEAR],
            "komponen": COMPONENT, "total_oktober": totals["final_oktober"], "total_november": totals["final_november"],
            "total_desember": totals["final_desember"], "total_kewajiban": totals["final_total_pengembalian"],
            "progress_per_tahap": {str(key): value for key, value in sorted(stage_counts.items())},
            "total_realisasi": total_realisasi, "jumlah_exact_payment": states["exact"],
            "jumlah_masih_piutang": states["receivable"], "jumlah_overpayment": states["overpayment"],
            "jumlah_invalid": len(errors), "jumlah_duplicate": len(duplicates)}
        return {"status": "ok", "valid": not errors and not duplicate, "source_type": SOURCE_TYPE,
            "filename": filename, "file_hash": digest, "sheet": sheet, "summary": summary,
            "duplicate_file": {"import_id": duplicate[0], "imported_at": str(duplicate[1])} if duplicate else None,
            "errors": errors[:50], "warnings": warnings[:50], "rows": normalized}

    def commit(self, content, filename, imported_by="system", file_hash=None):
        preview = self.preview(content, filename, file_hash)
        if not preview["valid"]:
            raise ValueError("IMPORT_TIDAK_VALID_ATAU_SUDAH_PERNAH_COMMITTED")
        import_id = "DASAR2025_" + uuid.uuid4().hex.upper()
        rows = preview["rows"]
        conn = duckdb.connect(self.db_path)
        try:
            conn.execute("BEGIN")
            duplicate = conn.execute("""SELECT import_batch_id FROM pengembalian_progress_import_batch
                WHERE source_file_hash=? AND status IN ('COMMITTED','PARTIAL') LIMIT 1""", [preview["file_hash"]]).fetchone()
            if duplicate: raise ValueError(f"FILE_ALREADY_IMPORTED: {duplicate[0]}")
            metadata = {"source_type": SOURCE_TYPE, "tahun_sp2d": YEAR, "komponen": COMPONENT,
                        "periode_id": PERIODE_ID, "filename": filename}
            conn.execute("""INSERT INTO pengembalian_archive_import
                (import_id,source_type,filename,file_hash,metadata,periode_id,imported_at,imported_by,row_count,
                 valid_count,invalid_count,status,dataset_type,source_file_name,archive_version,is_active)
                VALUES (?, ?, ?, ?, ?, ?, now(), ?, ?, ?, 0, 'ACTIVE', 'DATA_AWAL', ?, 'DASAR_2025', TRUE)""",
                [import_id, SOURCE_TYPE, filename, preview["file_hash"], json.dumps(metadata), PERIODE_ID, imported_by,
                 len(rows), len(rows), filename])
            archive_ids = {}
            for row in rows:
                row_id = str(uuid.uuid4()); archive_ids[row["nip"]] = row_id
                conn.execute("""INSERT INTO pengembalian_archive_data_awal
                    (row_id,import_id,source_type,source_row,nip,nama,periode_id,tahun_sp2d,tahun,periode,komponen,
                     tukin_pengembalian_oktober,tukin_pengembalian_november,tukin_pengembalian_desember,
                     total_potongan_tukin,source_file,source_sheet,provinsi,kabupaten_kota,kdanaksatker,
                     nomor_rekening,kode_bank,keterangan,validation_flags,raw_data)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    [row_id, import_id, SOURCE_TYPE, row["source_row"], row["nip"], row["nama"], PERIODE_ID,
                     YEAR, YEAR, "2025-10/12", COMPONENT,
                     row["final_oktober"], row["final_november"], row["final_desember"], row["calculated_total"],
                     filename, preview["sheet"], row["provinsi"], row["kabupaten_kota"], row["kode_anak_satker"],
                     row["nomor_rekening"], normalize_bank(row["nama_bank"]), row["status_master"], "[]",
                     json.dumps(row["raw"], default=str)])
                obligation_id = "OBL-" + sha256(f"{PERIODE_ID}|{row['nip']}|{YEAR}".encode()).hexdigest()[:24]
                conn.execute("""INSERT INTO pengembalian_master_obligation
                    (obligation_id,periode_id,nip,tahun_sp2d,nama_sdm,provinsi,final_tukin,final_uang_makan,final_total,
                     status_sp2d,tahun_sp2d_awal,tahun_sp2d_update,sumber_sp2d,flags_json,source_record_ids,
                     opening_balance,awal_tukin,awal_uang_makan,awal_total,source_import_id,created_at,updated_at)
                    VALUES (?,?,?,?,?,?,?,0,?,'NORMAL_SP2D',2025,2025,?,'[]',?,?,?,0,?,?,now(),now())
                    ON CONFLICT (periode_id,nip,tahun_sp2d) DO UPDATE SET nama_sdm=excluded.nama_sdm,
                    provinsi=excluded.provinsi,final_tukin=excluded.final_tukin,
                    final_total=excluded.final_tukin+coalesce(pengembalian_master_obligation.final_uang_makan,0),
                    sumber_sp2d=excluded.sumber_sp2d,source_record_ids=excluded.source_record_ids,
                    awal_tukin=excluded.awal_tukin,awal_total=excluded.awal_tukin+coalesce(pengembalian_master_obligation.awal_uang_makan,0),
                    source_import_id=excluded.source_import_id,updated_at=now()""",
                    [obligation_id, PERIODE_ID, row["nip"], YEAR, row["nama"], row["provinsi"], row["calculated_total"],
                     row["calculated_total"], SOURCE_TYPE, json.dumps([row_id]), row["calculated_total"],
                     row["calculated_total"], row["calculated_total"], import_id])
            transaction_rows = []
            for row in rows:
                obligation_id = "OBL-" + sha256(f"{PERIODE_ID}|{row['nip']}|{YEAR}".encode()).hexdigest()[:24]
                for item in row["progress"]:
                    events = [("BERHASIL", item["debit"] or 0, item.get("ntpn"), None),
                              ("GAGAL", item["gagal"] or 0, None, item.get("alasan")),
                              ("PENDING", item["pending"] or 0, None, None)]
                    for status, amount, ntpn, reason in events:
                        if not amount: continue
                        fingerprint = sha256(f"{preview['file_hash']}|{row['source_row']}|{item['tahap']}|{status}".encode()).hexdigest()
                        transaction_rows.append((str(uuid.uuid4()), fingerprint, import_id, obligation_id, row, item, status, amount, ntpn, reason))
            conn.execute("""INSERT INTO pengembalian_progress_import_batch
                (import_batch_id,source_filename,source_file_hash,status,total_rows,valid_rows,invalid_rows,duplicate_rows,
                 berhasil_rows,gagal_rows,total_tukin,total_uang_makan,total_debit,nominal_berhasil,nominal_gagal,
                 metadata_json,imported_by,import_mode,periode_id)
                VALUES (?,?,?,'COMMITTED',?,?,0,0,?,?,?,?,?,?,?,?,?,'DASAR_PENGEMBALIAN_2025',?)""",
                [import_id, filename, preview["file_hash"], len(transaction_rows), len(transaction_rows),
                 sum(1 for item in transaction_rows if item[6] == "BERHASIL"),
                 sum(1 for item in transaction_rows if item[6] == "GAGAL"),
                 sum(item[7] for item in transaction_rows), 0, sum(item[7] for item in transaction_rows),
                 sum(item[7] for item in transaction_rows if item[6] == "BERHASIL"),
                 sum(item[7] for item in transaction_rows if item[6] == "GAGAL"), json.dumps(metadata), imported_by, PERIODE_ID])
            for tx_id, fingerprint, batch_id, obligation_id, row, item, status, amount, ntpn, reason in transaction_rows:
                successful = amount if status == "BERHASIL" else 0
                conn.execute("""INSERT INTO pengembalian_progress_transaction
                    (transaction_id,fingerprint,import_batch_id,obligation_id,nip,nama_sdm,tahun_sp2d,provinsi,periode_id,
                     nama_bank_original,nama_bank_normalized,periode_debit,komponen_debit,nominal_debit_tukin,
                     nominal_debit_uang_makan,nominal_debit_total,status_debit,keterangan_gagal,source_filename,
                     source_row_number,validation_flags,raw_data,created_by,is_active,tahap,nomor_rekening,
                     nominal_diajukan,nominal_berhasil,ntpn,source_sheet)
                    VALUES (?,?,?,?,?,?,2025,?,?,?,? ,?,'TUKIN',?,0,?,?,?,?,?,'[]',?,?,TRUE,?,?,?,?,?,?)""",
                    [tx_id, fingerprint, batch_id, obligation_id, row["nip"], row["nama"], row["provinsi"], PERIODE_ID,
                     row["nama_bank"], normalize_bank(row["nama_bank"]), f"Tahap {item['tahap']}", amount, amount,
                     status, reason, filename, row["source_row"], json.dumps({"source_type": SOURCE_TYPE,
                     "source_row": row["source_row"], "send": item["send"], "reference": row["reference"]}, default=str),
                     imported_by, item["tahap"], row["nomor_rekening"], amount, successful, ntpn, preview["sheet"]])
            conn.execute("COMMIT")
        except Exception:
            conn.execute("ROLLBACK")
            raise
        finally:
            conn.close()
        result = ProgressService(self.db_path).progress({"periode_id": PERIODE_ID}, limit=max(len(rows), 1))
        return {"status": "ok", "import_id": import_id, "source_type": SOURCE_TYPE,
                "obligations": len(rows), "transactions": len(transaction_rows), "summary": result["summary"]}

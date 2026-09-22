"""Canonical monthly Master Pengembalian service.

21062026 MASTER UPDATE- (1).xlsx is the official opening source. Raw source
rows are kept losslessly, normalized monthly rows are versioned, and FINAL
PENGEMBALIAN is synchronized as the official Tukin obligation.
"""
from __future__ import annotations

import base64
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import importlib.util
import json
import os
from pathlib import Path
import re
import threading
import uuid

import duckdb
from backend.duckdb_maintenance import delete_where_without_secondary_indexes
import pandas as pd
from openpyxl import Workbook, load_workbook


DB = os.environ.get("EVKIN_DUCKDB_PATH") or str(
    Path(__file__).resolve().parent / "data" / "EVKIN_MASTER.duckdb"
)
_MIGRATED_DATABASES = set()
_MIGRATION_LOCK = threading.Lock()
OFFICIAL_SOURCE_FILENAME = "21062026 MASTER UPDATE- (1).xlsx"
OFFICIAL_SOURCE_CODE = "MASTER_UPDATE_21062026"
LEGACY_SOURCE_CODE = "MASTER_SEND_BPK"
REQUIRED_HEADERS = {"nip", "tahun", "bulan", "final_pengembalian"}
MONTHS = {
    "JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4,
    "MEI": 5, "JUNI": 6, "JULI": 7, "AGUSTUS": 8,
    "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12,
}

FIELDS = {
    "penilai": ("penilai",), "nip": ("nip",), "nama": ("nama",),
    "jabatan": ("jabatan",), "tahap": ("tahap",),
    "predikat_kinerja": ("predikat_kinerja_periodik",), "tahun": ("tahun",),
    "bulan": ("bulan",), "potongan_skp_persen": ("potongan_skp", "potongan_skp_persen"),
    "potongan_esdm_persen": ("potongan_absensi_esdm", "potongan_absensi_esdm_persen"),
    "potongan_eabsensi_persen": ("potongan_absensi_e_absensi", "potongan_absensi_e_absensi_persen"),
    "tukin": ("tukin",), "potongan_skp": ("potongan_skp_2",),
    "potongan_kehadiran_esdm": ("potongan_kehadiran_esdm",),
    "potongan_kehadiran_eabsensi": ("potongan_kehadiran_e_absensi",),
    "tukin_dibayar_esdm": ("tukin_dibayar_esdm",),
    "tukin_dibayar_eabsensi": ("tukin_dibayar_e_absensi",),
    "final_kehadiran": ("final_kehadiran",),
    "final_pengembalian": ("final_pengembalian",),
    "potongan_kehadiran_awal": ("potongan_kehadiran_awal",),
    "pengembalian_awal": ("pengembalian_awal",),
    "tahun_sp2d_update": ("status_sp2d_update", "tahun_sp2d_update"),
    "kdasatker_update": ("kdasatker_update",),
    "tahun_sp2d_awal": ("status_sp2d", "tahun_sp2d_awal"),
    "kdasatker_awal": ("kdasatker",), "status_send_bpk": ("status_send_bpk",),
    "status_tendik": ("status_tendik",), "zona_waktu": ("zona_waktu",),
    "provinsi": ("propinsi", "provinsi"), "kabupaten_kota": ("kabupaten", "kabupaten_kota"),
    "cuti": ("cuti",), "tanggal_mulai_cuti": ("tanggal_mulai_cuti",),
    "tanggal_selesai_cuti": ("tanggal_selesai_cuti",),
    "tendik_hitung_absen": ("tendik_hitung_absen",), "tendik_persentase": ("presentase", "persentase"),
    "absen_final_tendik": ("absen_final_tendik",),
    "pengembalian_final_tendik": ("pengembalian_final_tendik",),
    "status_pengembalian_final": ("status_pengembalian_final",),
    "status_pengembalian_awal": ("status_pengembalian_awal",),
    "cek_selisih": ("cek_selisih",), "keterangan_penyesuaian": ("s", "keterangan_penyesuaian"),
    "cek_bup": ("cek_bup",),
}

DB_COLUMNS = [
    "row_id", "import_id", "periode_id", "source_row_number", "business_key",
    "penilai", "nip", "nama", "jabatan", "tahap", "predikat_kinerja", "tahun",
    "bulan", "bulan_nama", "potongan_skp_persen", "potongan_esdm_persen",
    "potongan_eabsensi_persen", "tukin", "potongan_skp", "potongan_kehadiran_esdm",
    "potongan_kehadiran_eabsensi", "tukin_dibayar_esdm", "tukin_dibayar_eabsensi",
    "final_kehadiran", "final_pengembalian", "potongan_kehadiran_awal",
    "pengembalian_awal", "tahun_sp2d_update", "kdasatker_update", "tahun_sp2d_awal",
    "kdasatker_awal", "status_send_bpk", "status_tendik", "zona_waktu", "provinsi",
    "kabupaten_kota", "cuti", "tanggal_mulai_cuti", "tanggal_selesai_cuti",
    "tendik_hitung_absen", "tendik_persentase", "absen_final_tendik",
    "pengembalian_final_tendik", "status_pengembalian_final",
    "status_pengembalian_awal", "cek_selisih", "alasan_perubahan",
    "keterangan_penyesuaian", "cek_bup",
    "selisih_pengembalian", "pengembalian_awal_exact", "final_pengembalian_exact",
    "selisih_pengembalian_exact", "status_perubahan", "validation_status", "validation_flags",
]


def _key(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return None if not text or text.lower() in {"nan", "none", "null"} else text


def _nip(value):
    value = _text(value)
    return re.sub(r"\.0$", "", value) if value else None


def _decimal(value):
    if value is None or value == "" or str(value).strip().startswith("#"):
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value))
    text = re.sub(r"[^0-9,.-]", "", str(value).strip())
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _percent(value):
    number = _decimal(value)
    if number is None:
        return None
    if isinstance(value, str) and "%" in value:
        number /= 100
    return float(number)


def _integer(value):
    number = _decimal(value)
    return int(number) if number is not None else None


def _month(value):
    number = _integer(value)
    if number and 1 <= number <= 12:
        return number
    return MONTHS.get(str(value or "").strip().upper())


def _json(value):
    return json.dumps(value, ensure_ascii=False, default=str, separators=(",", ":"))


def _change_reason(value, status_perubahan):
    """Normalize the source `cek selisih` while preserving it separately."""
    source = _text(value)
    if source and re.fullmatch(r"0+(?:[.,]0+)?", source):
        source = None
    if not source:
        return "TIDAK ADA PERUBAHAN" if status_perubahan == "TIDAK BERUBAH" else "PERUBAHAN TANPA KETERANGAN"
    text = re.sub(r"\s+", " ", source.upper()).strip()
    has_absence = any(token in text for token in ("ABSEN", "ESDM", "KEHADIRAN"))
    has_zone = any(token in text for token in ("ZONA", "WITA", "WIT"))
    has_tendik = "TENDIK" in text
    has_leave = "CUTI" in text
    if "TAHAP 2" in text:
        return "TAHAP 2"
    if "MASIH PENGEMBALIAN" in text:
        return "MASIH PENGEMBALIAN"
    if has_leave and has_zone:
        return "PENYESUAIAN CUTI & ZONA WAKTU"
    if has_absence and has_zone:
        return "PENYESUAIAN ABSENSI & ZONA WAKTU"
    if has_tendik and has_zone:
        return "PENYESUAIAN TENDIK & ZONA WAKTU"
    if has_leave:
        return "PENYESUAIAN CUTI"
    if has_tendik:
        return "PENYESUAIAN TENDIK"
    if has_absence:
        return "PENYESUAIAN ABSENSI"
    if has_zone:
        return "PENYESUAIAN ZONA WAKTU"
    return text


class MonthlyMasterService:
    def __init__(self, db_path=None):
        self.db_path = str(db_path or os.environ.get("EVKIN_DUCKDB_PATH") or DB)
        with _MIGRATION_LOCK:
            if self.db_path not in _MIGRATED_DATABASES:
                for number in ("042_create_pengembalian_monthly_master.py", "043_add_pengembalian_change_reason.py",
                               "044_expand_pengembalian_money_precision.py",
                               "047_normalize_progress_obligation_snapshot.py"):
                    migration = Path(__file__).resolve().parent / "migrations" / number
                    spec = importlib.util.spec_from_file_location(f"monthly_master_{number}", migration)
                    module = importlib.util.module_from_spec(spec)
                    spec.loader.exec_module(module)
                    module.main(self.db_path)
                _MIGRATED_DATABASES.add(self.db_path)

    def connect(self):
        return duckdb.connect(self.db_path)

    @staticmethod
    def decode_file(payload):
        encoded = payload.get("file_base64") or payload.get("content_base64")
        if not encoded:
            raise ValueError("file_base64 wajib diisi")
        if "," in encoded and encoded.split(",", 1)[0].startswith("data:"):
            encoded = encoded.split(",", 1)[1]
        try:
            return base64.b64decode(encoded, validate=True)
        except Exception as exc:
            raise ValueError("file_base64 tidak valid") from exc

    @staticmethod
    def read_xlsx(content):
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            candidates = []
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                for header_row, raw_headers in enumerate(
                    worksheet.iter_rows(min_row=1, max_row=min(25, worksheet.max_row), values_only=True), 1
                ):
                    normalized = [_key(value) for value in raw_headers]
                    if REQUIRED_HEADERS.issubset(set(normalized)):
                        recognized = sum(1 for field in FIELDS.values() if any(alias in normalized for alias in field))
                        candidates.append((recognized, -sheet_index, -header_row, worksheet, header_row, raw_headers))
            if not candidates:
                raise ValueError(
                    "Tidak ditemukan sheet detail dengan header wajib: NIP, Tahun, Bulan, FINAL PENGEMBALIAN"
                )
            _, _, _, sheet, header_row, raw_headers = max(candidates, key=lambda item: item[:3])
            headers = []
            counts = Counter()
            for value in raw_headers:
                base = _key(value)
                counts[base] += 1
                headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
            rows = []
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1
            ):
                if not any(value not in (None, "") for value in values):
                    continue
                rows.append({"_source_row": row_number, **{
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                }})
            return sheet.title, rows
        finally:
            workbook.close()

    @staticmethod
    def normalize_rows(source_rows):
        normalized = []
        seen = Counter()
        for raw in source_rows:
            get = lambda field: next((raw.get(alias) for alias in FIELDS[field] if raw.get(alias) not in (None, "")), None)
            nip = _nip(get("nip")); year = _integer(get("tahun")); month = _month(get("bulan"))
            business_key = f"{nip or ''}|{year or ''}|{month or ''}"
            seen[business_key] += 1
            flags = []
            errors = []
            if not nip: errors.append("NIP_REQUIRED")
            elif not re.fullmatch(r"\d{10,24}", nip): flags.append("NIP_FORMAT_REVIEW")
            if year is None or year < 2000 or year > 2100: errors.append("YEAR_INVALID")
            if month is None: errors.append("MONTH_INVALID")
            final_return = _decimal(get("final_pengembalian"))
            initial_return = _decimal(get("pengembalian_awal"))
            final_presence = _decimal(get("final_kehadiran"))
            initial_presence = _decimal(get("potongan_kehadiran_awal"))
            skp = _decimal(get("potongan_skp"))
            if final_return is None: errors.append("FINAL_PENGEMBALIAN_REQUIRED")
            if any(str(get(field) or "").strip().startswith("#") for field in (
                "tendik_hitung_absen", "tendik_persentase", "absen_final_tendik", "pengembalian_final_tendik"
            )): flags.append("TENDIK_SOURCE_REF_ERROR")
            if get("penilai") in (None, ""): flags.append("PENILAI_EMPTY")
            if final_return is not None and final_presence is not None and skp is not None and abs(final_return-final_presence-skp) >= 1:
                errors.append("FINAL_FORMULA_MISMATCH")
            if initial_return is not None and initial_presence is not None and skp is not None and abs(initial_return-initial_presence-skp) >= 1:
                flags.append("INITIAL_FORMULA_MISMATCH")
            status_perubahan = "NAIK" if final_return is not None and initial_return is not None and final_return > initial_return else "TURUN" if final_return is not None and initial_return is not None and final_return < initial_return else "TIDAK BERUBAH"
            cek_selisih = _text(get("cek_selisih"))
            row = {
                "source_row_number": int(raw.get("_source_row") or 0), "business_key": business_key,
                "penilai": _text(get("penilai")), "nip": nip, "nama": _text(get("nama")),
                "jabatan": _text(get("jabatan")), "tahap": _text(get("tahap")),
                "predikat_kinerja": _text(get("predikat_kinerja")), "tahun": year, "bulan": month,
                "bulan_nama": _text(get("bulan")), "potongan_skp_persen": _percent(get("potongan_skp_persen")),
                "potongan_esdm_persen": _percent(get("potongan_esdm_persen")),
                "potongan_eabsensi_persen": _percent(get("potongan_eabsensi_persen")),
                "tukin": _decimal(get("tukin")), "potongan_skp": skp,
                "potongan_kehadiran_esdm": _decimal(get("potongan_kehadiran_esdm")),
                "potongan_kehadiran_eabsensi": _decimal(get("potongan_kehadiran_eabsensi")),
                "tukin_dibayar_esdm": _decimal(get("tukin_dibayar_esdm")),
                "tukin_dibayar_eabsensi": _decimal(get("tukin_dibayar_eabsensi")),
                "final_kehadiran": final_presence, "final_pengembalian": final_return,
                "potongan_kehadiran_awal": initial_presence, "pengembalian_awal": initial_return,
                "tahun_sp2d_update": _integer(get("tahun_sp2d_update")),
                "kdasatker_update": _text(get("kdasatker_update")),
                "tahun_sp2d_awal": _integer(get("tahun_sp2d_awal")),
                "kdasatker_awal": _text(get("kdasatker_awal")),
                "status_send_bpk": _text(get("status_send_bpk")),
                "status_tendik": (_text(get("status_tendik")) or "").upper() or None,
                "zona_waktu": (_text(get("zona_waktu")) or "").upper() or None,
                "provinsi": (_text(get("provinsi")) or "").upper() or None,
                "kabupaten_kota": _text(get("kabupaten_kota")), "cuti": _text(get("cuti")),
                "tanggal_mulai_cuti": _text(get("tanggal_mulai_cuti")),
                "tanggal_selesai_cuti": _text(get("tanggal_selesai_cuti")),
                "tendik_hitung_absen": _text(get("tendik_hitung_absen")),
                "tendik_persentase": _text(get("tendik_persentase")),
                "absen_final_tendik": _text(get("absen_final_tendik")),
                "pengembalian_final_tendik": _text(get("pengembalian_final_tendik")),
                "status_pengembalian_final": _text(get("status_pengembalian_final")),
                "status_pengembalian_awal": _text(get("status_pengembalian_awal")),
                "cek_selisih": cek_selisih,
                "alasan_perubahan": _change_reason(cek_selisih, status_perubahan),
                "keterangan_penyesuaian": _text(get("keterangan_penyesuaian")),
                "cek_bup": _text(get("cek_bup")),
                "selisih_pengembalian": final_return-initial_return if final_return is not None and initial_return is not None else None,
                "pengembalian_awal_exact": float(initial_return) if initial_return is not None else None,
                "final_pengembalian_exact": float(final_return) if final_return is not None else None,
                "selisih_pengembalian_exact": float(final_return-initial_return) if final_return is not None and initial_return is not None else None,
                "status_perubahan": status_perubahan,
                "flags": flags, "errors": errors, "raw": {k: v for k, v in raw.items() if k != "_source_row"},
            }
            normalized.append(row)
        for row in normalized:
            if seen[row["business_key"]] > 1:
                row["errors"].append("DUPLICATE_NIP_YEAR_MONTH")
            row["validation_status"] = "INVALID" if row["errors"] else "WARNING" if row["flags"] else "VALID"
            row["validation_flags"] = row["errors"] + row["flags"]
        return normalized

    @staticmethod
    def summarize(rows):
        statuses = Counter(row["validation_status"] for row in rows)
        flags = Counter(flag for row in rows for flag in row["validation_flags"])
        valid_rows = [row for row in rows if row["validation_status"] != "INVALID"]
        return {
            "total_rows": len(rows), "valid_rows": statuses["VALID"],
            "warning_rows": statuses["WARNING"], "invalid_rows": statuses["INVALID"],
            "duplicate_rows": flags["DUPLICATE_NIP_YEAR_MONTH"],
            "total_awal": float(sum((row["pengembalian_awal"] or 0) for row in valid_rows)),
            "total_final": float(sum((row["final_pengembalian"] or 0) for row in valid_rows)),
            "unique_nip": len({row["nip"] for row in valid_rows}),
            "years": sorted({row["tahun"] for row in valid_rows}),
            "months": dict(sorted(Counter(row["bulan_nama"] for row in valid_rows).items())),
            "flags": dict(flags.most_common()),
        }

    def preview_file(self, content, filename=None):
        sheet, source_rows = self.read_xlsx(content)
        rows = self.normalize_rows(source_rows)
        return {"status": "ok", "filename": filename, "sheet": sheet,
                "file_hash": sha256(content).hexdigest(), "summary": self.summarize(rows),
                "samples": [{k: (float(v) if isinstance(v, Decimal) else v) for k, v in row.items() if k not in {"raw"}}
                            for row in rows[:10]]}

    def _ensure_period(self, conn, period_id, rows, created_by):
        found = conn.execute("SELECT periode_id FROM pengembalian_period WHERE periode_id=?", [period_id]).fetchone()
        if found:
            return
        years = sorted({row["tahun"] for row in rows if row["tahun"]})
        months = sorted({row["bulan"] for row in rows if row["bulan"]})
        year = years[-1] if years else datetime.now().year
        conn.execute("""INSERT INTO pengembalian_period
            (periode_id,nama_periode,tahun,bulan_awal,bulan_akhir,status,tahun_proses,
             bulan_proses,tahun_penghasilan,bulan_penghasilan,is_active,created_by,notes)
            VALUES (?,?,?,?,?,'ACTIVE',?,?,?,?,TRUE,?,'Dibuat otomatis dari 21062026 MASTER UPDATE')""",
            [period_id, f"Master Pengembalian Q4 {year}", year, min(months or [1]), max(months or [12]),
             year, max(months or [12]), year, min(months or [1]), created_by])

    def import_file(self, content, filename, period_id, imported_by="system", metadata=None):
        if not period_id:
            raise ValueError("periode_id wajib diisi")
        sheet, source_rows = self.read_xlsx(content)
        rows = self.normalize_rows(source_rows)
        summary = self.summarize(rows)
        file_hash = sha256(content).hexdigest()
        import_id = str(uuid.uuid4())
        conn = self.connect()
        try:
            duplicate = conn.execute("""SELECT import_id FROM pengembalian_monthly_import
                WHERE periode_id=? AND source_file_hash=? AND status='ACTIVE' LIMIT 1""",
                [period_id, file_hash]).fetchone()
            if duplicate:
                return {"status": "duplicate", "import_id": duplicate[0], "summary": summary}
            conn.execute("BEGIN")
            self._ensure_period(conn, period_id, rows, imported_by)
            conn.execute("""UPDATE pengembalian_monthly_import SET status='SUPERSEDED',
                is_active=FALSE,superseded_at=now() WHERE periode_id=? AND is_active=TRUE""", [period_id])
            conn.execute("""INSERT INTO pengembalian_monthly_import
                (import_id,periode_id,source_filename,source_file_hash,source_sheet,status,is_active,
                 total_rows,valid_rows,warning_rows,invalid_rows,duplicate_rows,total_awal,total_final,
                 metadata_json,validation_json,imported_by,activated_at)
                VALUES (?,?,?,?,?,'ACTIVE',TRUE,?,?,?,?,?,?,?,?,?,?,now())""",
                [import_id, period_id, filename, file_hash, sheet, summary["total_rows"],
                 summary["valid_rows"], summary["warning_rows"], summary["invalid_rows"],
                 summary["duplicate_rows"], summary["total_awal"], summary["total_final"],
                 _json(metadata or {}), _json(summary), imported_by])
            source_values = []
            master_values = []
            for row in rows:
                source_id = str(uuid.uuid4())
                source_values.append([source_id, import_id, row["source_row_number"], row["business_key"],
                                      _json(row["raw"]), row["validation_status"], _json(row["validation_flags"])])
                if row["validation_status"] == "INVALID":
                    continue
                values = {**row, "row_id": str(uuid.uuid4()), "import_id": import_id, "periode_id": period_id,
                          "validation_flags": _json(row["validation_flags"])}
                master_values.append([values.get(column) for column in DB_COLUMNS])
            source_columns = ["source_row_id", "import_id", "source_row_number", "business_key",
                              "raw_json", "validation_status", "validation_flags"]
            source_frame = pd.DataFrame(source_values, columns=source_columns)
            master_frame = pd.DataFrame(master_values, columns=DB_COLUMNS)
            conn.register("monthly_source_stage", source_frame)
            conn.register("monthly_master_stage", master_frame)
            conn.execute(f"""INSERT INTO pengembalian_monthly_source_row ({','.join(source_columns)})
                SELECT {','.join(source_columns)} FROM monthly_source_stage""")
            conn.execute(f"""INSERT INTO pengembalian_monthly_master ({','.join(DB_COLUMNS)})
                SELECT {','.join(DB_COLUMNS)} FROM monthly_master_stage""")
            conn.unregister("monthly_source_stage")
            conn.unregister("monthly_master_stage")
            self._sync_obligations(conn, period_id)
            conn.execute("COMMIT")
        except Exception:
            conn.execute("ROLLBACK")
            raise
        finally:
            conn.close()
        return {"status": "ok", "import_id": import_id, "period_id": period_id,
                "file_hash": file_hash, "summary": summary}

    @staticmethod
    def _sync_obligations(conn, period_id):
        conn.execute("""UPDATE pengembalian_master_obligation SET final_tukin=0,
            final_total=0,flags_json='["SOURCE_SUPERSEDED"]',updated_at=now()
            WHERE periode_id=?""", [period_id])
        conn.execute("""INSERT INTO pengembalian_master_obligation
            (obligation_id,periode_id,nip,tahun_sp2d,nama_sdm,provinsi,final_tukin,
             final_uang_makan,final_total,status_sp2d,tahun_sp2d_awal,tahun_sp2d_update,
             sumber_sp2d,flags_json,source_record_ids,opening_balance,created_at,updated_at,
             awal_tukin,awal_uang_makan,awal_total,source_import_id)
            SELECT 'OBL-MONTHLY-' || md5(periode_id || '|' || nip || '|' ||
                    coalesce(tahun_sp2d_update,tahun_sp2d_awal,tahun)::VARCHAR),
                   periode_id,nip,coalesce(tahun_sp2d_update,tahun_sp2d_awal,tahun),
                   max(nama),max(provinsi),sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)),0,
                   sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)),
                   CASE WHEN max(tahun_sp2d_update) IS DISTINCT FROM max(tahun_sp2d_awal)
                        THEN 'UPDATE_SP2D' ELSE 'NORMAL_SP2D' END,
                    max(tahun_sp2d_awal),max(tahun_sp2d_update),?,
                   CASE WHEN count_if(validation_status='WARNING')>0
                        THEN '["SOURCE_WARNING"]' ELSE '[]' END,
                   to_json(list(row_id)),sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)),now(),now(),
                   sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),0,
                   sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),max(import_id)
            FROM vw_pengembalian_monthly_current WHERE periode_id=?
            GROUP BY periode_id,nip,coalesce(tahun_sp2d_update,tahun_sp2d_awal,tahun)
            ON CONFLICT (periode_id,nip,tahun_sp2d) DO UPDATE SET
                obligation_id=excluded.obligation_id,nama_sdm=excluded.nama_sdm,
                provinsi=excluded.provinsi,final_tukin=excluded.final_tukin,
                final_uang_makan=excluded.final_uang_makan,final_total=excluded.final_total,
                status_sp2d=excluded.status_sp2d,tahun_sp2d_awal=excluded.tahun_sp2d_awal,
                tahun_sp2d_update=excluded.tahun_sp2d_update,sumber_sp2d=excluded.sumber_sp2d,
                flags_json=excluded.flags_json,source_record_ids=excluded.source_record_ids,
                opening_balance=excluded.opening_balance,awal_tukin=excluded.awal_tukin,
                awal_uang_makan=excluded.awal_uang_makan,awal_total=excluded.awal_total,
                source_import_id=excluded.source_import_id,updated_at=now()""", [OFFICIAL_SOURCE_CODE, period_id])

    def dashboard(self, period_id=None):
        conn = self.connect()
        try:
            where = " WHERE periode_id=?" if period_id else ""
            params = [period_id] if period_id else []
            row = conn.execute(f"""SELECT count(*),count(DISTINCT nip),sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)),sum(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0)),
                count_if(status_perubahan='NAIK'),count_if(status_perubahan='TURUN'),
                count_if(status_perubahan='TIDAK BERUBAH'),count_if(validation_status='WARNING')
                FROM vw_pengembalian_monthly_current{where}""", params).fetchone()
            def groups(column):
                return [{column: r[0], "rows": r[1], "total_final": float(r[2] or 0)} for r in conn.execute(
                    f"SELECT {column},count(*),sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)) FROM vw_pengembalian_monthly_current{where} GROUP BY {column} ORDER BY {column}", params).fetchall()]
            status_month_cursor = conn.execute(f"""SELECT coalesce(status_send_bpk,'BELUM MASUK BPK') status_send_bpk,
                bulan,bulan_nama,count(*),sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0))
                FROM vw_pengembalian_monthly_current{where}
                GROUP BY 1,2,3 ORDER BY 1 DESC,2""", params)
            by_status_month = [{"status_send_bpk": r[0], "bulan": r[1], "bulan_nama": r[2],
                                "rows": r[3], "total_awal": float(r[4] or 0),
                                "total_final": float(r[5] or 0)} for r in status_month_cursor.fetchall()]
            sp2d_cursor = conn.execute(f"""SELECT coalesce(status_send_bpk,'BELUM MASUK BPK'),
                bulan,bulan_nama,tahun_sp2d_update,count(*),sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)) FROM vw_pengembalian_monthly_current{where}
                GROUP BY 1,2,3,4 ORDER BY 1,2,4""", params)
            by_sp2d_status_month = [{"status_send_bpk": r[0], "bulan": r[1], "bulan_nama": r[2],
                                     "tahun_sp2d_update": r[3], "rows": r[4],
                                     "total_awal": float(r[5] or 0), "total_final": float(r[6] or 0)}
                                    for r in sp2d_cursor.fetchall()]
            reason_where = where + (" AND " if where else " WHERE ") + "status_perubahan<>'TIDAK BERUBAH'"
            reason_cursor = conn.execute(f"""SELECT coalesce(alasan_perubahan,'PERUBAHAN TANPA KETERANGAN'),
                count(*),count(DISTINCT nip),sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)),
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)),sum(abs(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0))),
                count_if(status_perubahan='NAIK'),count_if(status_perubahan='TURUN')
                FROM vw_pengembalian_monthly_current{reason_where}
                GROUP BY 1 ORDER BY 6 DESC,2 DESC""", params)
            by_reason = [{"alasan_perubahan": r[0], "rows": r[1], "sdm": r[2],
                          "total_awal": float(r[3] or 0), "total_final": float(r[4] or 0),
                          "dampak_nominal": float(r[5] or 0), "naik": r[6], "turun": r[7]}
                         for r in reason_cursor.fetchall()]
            active = conn.execute("""SELECT import_id,periode_id,source_filename,imported_at,total_rows,
                valid_rows,warning_rows,invalid_rows,total_final,source_sheet FROM pengembalian_monthly_import
                WHERE is_active=TRUE ORDER BY imported_at DESC""").fetchall()
            return {"status":"ok","total_rows":row[0],"total_sdm":row[1],"total_awal":float(row[2] or 0),
                    "total_final":float(row[3] or 0),"selisih":float(row[4] or 0),"naik":row[5],"turun":row[6],
                    "tidak_berubah":row[7],"warning_rows":row[8],"by_month":groups("bulan_nama"),
                    "by_zone":groups("zona_waktu"),"by_status_month":by_status_month,
                    "by_sp2d_status_month":by_sp2d_status_month,"by_reason":by_reason,
                    "active_imports":[{"import_id":r[0],"periode_id":r[1],
                    "filename":r[2],"imported_at":str(r[3]),"total_rows":r[4],"valid_rows":r[5],
                    "warning_rows":r[6],"invalid_rows":r[7],"total_final":float(r[8] or 0),
                    "source_sheet":r[9]} for r in active]}
        finally:
            conn.close()

    def list_rows(self, filters=None, limit=100, offset=0):
        filters = dict(filters or {})
        clauses=[]; params=[]
        exact = {"periode_id":"periode_id","tahun":"tahun","bulan":"bulan","tahun_sp2d":"tahun_sp2d_update",
                 "zona_waktu":"zona_waktu","status_tendik":"status_tendik","provinsi":"provinsi",
                 "status_send_bpk":"status_send_bpk","status_perubahan":"status_perubahan",
                 "validation_status":"validation_status","alasan_perubahan":"alasan_perubahan"}
        for key,column in exact.items():
            if filters.get(key) not in (None,""):
                clauses.append(f"{column}=?"); params.append(filters[key])
        if filters.get("search"):
            clauses.append("(nip ILIKE '%'||?||'%' OR nama ILIKE '%'||?||'%' OR cek_selisih ILIKE '%'||?||'%')")
            params.extend([filters["search"],filters["search"],filters["search"]])
        where = " WHERE "+" AND ".join(clauses) if clauses else ""
        conn=self.connect()
        try:
            total=conn.execute("SELECT count(*) FROM vw_pengembalian_monthly_current"+where,params).fetchone()[0]
            columns="row_id,periode_id,nip,nama,jabatan,tahun,bulan,bulan_nama,tahap,predikat_kinerja,tukin,potongan_skp,potongan_kehadiran_esdm,potongan_kehadiran_eabsensi,final_kehadiran,potongan_kehadiran_awal,coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE)) pengembalian_awal,coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE)) final_pengembalian,coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE)) selisih_pengembalian,status_perubahan,tahun_sp2d_awal,tahun_sp2d_update,kdasatker_awal,kdasatker_update,status_send_bpk,status_tendik,zona_waktu,provinsi,kabupaten_kota,cek_selisih,alasan_perubahan,keterangan_penyesuaian,validation_status,validation_flags,source_row_number,import_id"
            cursor=conn.execute(f"SELECT {columns} FROM vw_pengembalian_monthly_current{where} ORDER BY nip,tahun,bulan LIMIT ? OFFSET ?",params+[int(limit),int(offset)])
            names=[item[0] for item in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            for row in rows:
                for key,value in list(row.items()):
                    if isinstance(value,Decimal): row[key]=float(value)
            return {"status":"ok","rows":rows,"total":total,"limit":limit,"offset":offset,"has_more":offset+len(rows)<total}
        finally: conn.close()

    def detail(self,row_id):
        conn=self.connect()
        try:
            cursor=conn.execute("SELECT * FROM vw_pengembalian_monthly_current WHERE row_id=?",[row_id])
            names=[item[0] for item in cursor.description]; row=cursor.fetchone()
            if not row: raise ValueError("Data tidak ditemukan")
            result=dict(zip(names,row))
            lineage = conn.execute("""SELECT source_filename,source_file_hash,source_sheet,imported_at,
                imported_by FROM pengembalian_monthly_import WHERE import_id=?""", [result["import_id"]]).fetchone()
            result["source_filename"], result["source_file_hash"], result["source_sheet"] = lineage[:3]
            result["source_imported_at"], result["source_imported_by"] = str(lineage[3]), lineage[4]
            result["raw_source"]=json.loads(conn.execute("""SELECT raw_json FROM pengembalian_monthly_source_row
                WHERE import_id=? AND source_row_number=?""",[result["import_id"],result["source_row_number"]]).fetchone()[0])
            for key,value in list(result.items()):
                if isinstance(value,Decimal): result[key]=float(value)
            return result
        finally: conn.close()

    def reasons(self, period_id=None, status_send_bpk=None, bulan=None):
        filters = {"periode_id": period_id, "status_send_bpk": status_send_bpk, "bulan": bulan}
        clauses=[]; params=[]
        for key,column in (("periode_id","periode_id"),("status_send_bpk","status_send_bpk"),("bulan","bulan")):
            if filters.get(key) not in (None, ""):
                clauses.append(f"{column}=?"); params.append(filters[key])
        clauses.append("status_perubahan<>'TIDAK BERUBAH'")
        where=" WHERE "+" AND ".join(clauses)
        conn=self.connect()
        try:
            cursor=conn.execute(f"""SELECT coalesce(alasan_perubahan,'PERUBAHAN TANPA KETERANGAN') alasan,
                count(*) AS "rows",count(DISTINCT nip) sdm,sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)) total_awal,
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)) total_final,
                sum(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0)) selisih,
                sum(abs(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0))) dampak_nominal
                FROM vw_pengembalian_monthly_current{where} GROUP BY 1 ORDER BY 7 DESC,2 DESC""",params)
            names=[x[0] for x in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            for item in rows:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
            return {"status":"ok","rows":rows}
        finally: conn.close()

    def bnba_trace(self, nip, period_id=None):
        nip=_nip(nip)
        if not nip: raise ValueError("NIP wajib diisi")
        where="m.nip=?"; params=[nip]
        if period_id:
            where+=" AND m.periode_id=?"; params.append(period_id)
        conn=self.connect()
        try:
            cursor=conn.execute(f"""SELECT m.row_id,m.periode_id,m.nip,m.nama,m.jabatan,m.tahun,m.bulan,
                m.bulan_nama,coalesce(m.pengembalian_awal_exact,cast(m.pengembalian_awal AS DOUBLE)) pengembalian_awal,
                coalesce(m.final_pengembalian_exact,cast(m.final_pengembalian AS DOUBLE)) final_pengembalian,
                coalesce(m.selisih_pengembalian_exact,cast(m.selisih_pengembalian AS DOUBLE)) selisih_pengembalian,
                m.status_perubahan,m.cek_selisih,m.alasan_perubahan,m.keterangan_penyesuaian,
                m.status_send_bpk,m.tahun_sp2d_awal,m.tahun_sp2d_update,m.kdasatker_awal,
                m.kdasatker_update,m.status_tendik,m.zona_waktu,m.provinsi,m.kabupaten_kota,
                m.validation_status,m.validation_flags,m.source_row_number,i.source_filename,
                i.source_sheet,i.imported_at FROM vw_pengembalian_monthly_current m
                JOIN pengembalian_monthly_import i ON i.import_id=m.import_id WHERE {where}
                ORDER BY m.tahun,m.bulan""",params)
            names=[x[0] for x in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            if not rows: raise ValueError("BNBA tidak ditemukan")
            for item in rows:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
                    elif isinstance(value,datetime): item[key]=str(value)
            obligations_cursor=conn.execute("""SELECT periode_id,tahun_sp2d,final_tukin,final_total,status_sp2d,
                tahun_sp2d_awal,tahun_sp2d_update,sumber_sp2d,flags_json,updated_at
                FROM pengembalian_master_obligation WHERE nip=? ORDER BY tahun_sp2d""",[nip])
            obligation_names=[x[0] for x in obligations_cursor.description]
            obligations=[dict(zip(obligation_names,row)) for row in obligations_cursor.fetchall()]
            for item in obligations:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
                    elif isinstance(value,datetime): item[key]=str(value)
            return {"status":"ok","nip":nip,"nama":rows[0]["nama"],"rows":rows,
                    "obligations":obligations,"summary":{"months":len(rows),
                    "total_awal":sum(item["pengembalian_awal"] or 0 for item in rows),
                    "total_final":sum(item["final_pengembalian"] or 0 for item in rows),
                    "total_selisih":sum(item["selisih_pengembalian"] or 0 for item in rows)}}
        finally: conn.close()

    def imports(self,limit=100):
        conn=self.connect()
        try:
            cursor=conn.execute("""SELECT import_id,periode_id,source_filename,source_file_hash,source_sheet,status,is_active,
                total_rows,valid_rows,warning_rows,invalid_rows,duplicate_rows,total_awal,total_final,
                imported_by,imported_at,activated_at,superseded_at FROM pengembalian_monthly_import
                ORDER BY imported_at DESC LIMIT ?""",[int(limit)])
            names=[x[0] for x in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            for item in rows:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
                    elif isinstance(value,datetime): item[key]=str(value)
            return {"status":"ok","rows":rows}
        finally: conn.close()

    def activate(self,import_id,activated_by="system"):
        conn=self.connect()
        try:
            found=conn.execute("SELECT periode_id FROM pengembalian_monthly_import WHERE import_id=?",[import_id]).fetchone()
            if not found: raise ValueError("Import tidak ditemukan")
            conn.execute("BEGIN"); period_id=found[0]
            conn.execute("UPDATE pengembalian_monthly_import SET status='SUPERSEDED',is_active=FALSE,superseded_at=now() WHERE periode_id=? AND is_active=TRUE",[period_id])
            conn.execute("UPDATE pengembalian_monthly_import SET status='ACTIVE',is_active=TRUE,activated_at=now(),superseded_at=NULL WHERE import_id=?",[import_id])
            self._sync_obligations(conn,period_id); conn.execute("COMMIT")
            return {"status":"ok","import_id":import_id,"periode_id":period_id,"activated_by":activated_by}
        except Exception:
            try:
                conn.execute("ROLLBACK")
            except Exception:
                pass
            raise
        finally: conn.close()

    def delete_import(self, import_id, deleted_by="system"):
        """Delete one superseded import and all rows belonging to that version."""
        normalized_id = str(import_id or "").strip()
        if not normalized_id:
            raise ValueError("ID import wajib diisi")
        conn = self.connect()
        try:
            found = conn.execute(
                "SELECT periode_id, status, is_active FROM pengembalian_monthly_import WHERE import_id=?",
                [normalized_id],
            ).fetchone()
            if not found:
                raise ValueError("Import tidak ditemukan")
            period_id, status, is_active = found
            if bool(is_active) or str(status).upper() == "ACTIVE":
                raise ValueError("Import aktif tidak dapat dihapus. Aktifkan versi lain terlebih dahulu.")
            result = delete_where_without_secondary_indexes(conn, [
                ("pengembalian_monthly_source_row", "import_id=?", [normalized_id]),
                ("pengembalian_monthly_master", "import_id=?", [normalized_id]),
                ("pengembalian_monthly_import", "import_id=?", [normalized_id]),
            ])
            return {
                "status": "ok",
                "import_id": normalized_id,
                "periode_id": period_id,
                "deleted_by": deleted_by,
                "deleted_rows": result["deleted_rows"],
                "tables": result["deleted_tables"],
            }
        finally:
            conn.close()

    def reset_source_database(self, reset_by="system"):
        """Remove all monthly source versions without deleting Progress ledger rows.

        Obligations materialized from the monthly source are removed as well so
        an empty source database cannot continue presenting stale Final values.
        A subsequent official-source import recreates those obligations.
        """
        conn = self.connect()
        try:
            progress_rows = 0
            if conn.execute("""SELECT count(*) FROM information_schema.tables
                WHERE table_schema=current_schema()
                  AND table_name='pengembalian_progress_transaction'""").fetchone()[0]:
                progress_rows = int(conn.execute(
                    "SELECT count(*) FROM pengembalian_progress_transaction"
                ).fetchone()[0])
            result = delete_where_without_secondary_indexes(conn, [
                ("pengembalian_master_obligation",
                 "sumber_sp2d=? OR obligation_id LIKE 'OBL-MONTHLY-%'",
                 [OFFICIAL_SOURCE_CODE]),
                ("pengembalian_monthly_source_row", "1=1", []),
                ("pengembalian_monthly_master", "1=1", []),
                ("pengembalian_monthly_import", "1=1", []),
            ])
            return {
                **result,
                "reset_by": reset_by,
                "progress_preserved": True,
                "progress_rows": progress_rows,
            }
        finally:
            conn.close()

    def reconciliation(self,period_id=None):
        conn=self.connect()
        try:
            where=" WHERE periode_id=?" if period_id else ""; params=[period_id] if period_id else []
            cursor=conn.execute(f"""SELECT coalesce(status_send_bpk,'BELUM SEND BPK') status_send_bpk,
                count(*) jumlah_data,count(DISTINCT nip) jumlah_sdm,sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)) nominal_awal,
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)) nominal_final,
                sum(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0)) selisih,
                count_if(status_perubahan='NAIK') naik,count_if(status_perubahan='TURUN') turun,
                count_if(status_perubahan='TIDAK BERUBAH') tidak_berubah
                FROM vw_pengembalian_monthly_current{where} GROUP BY 1 ORDER BY 1""",params)
            names=[x[0] for x in cursor.description]
            rows=[dict(zip(names,row)) for row in cursor.fetchall()]
            for item in rows:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
            month_cursor=conn.execute(f"""SELECT coalesce(status_send_bpk,'BELUM MASUK BPK') status_send_bpk,
                bulan,bulan_nama,count(*) jumlah_data,count(DISTINCT nip) jumlah_sdm,
                sum(coalesce(pengembalian_awal_exact,cast(pengembalian_awal AS DOUBLE),0)) nominal_awal,
                sum(coalesce(final_pengembalian_exact,cast(final_pengembalian AS DOUBLE),0)) nominal_final,
                sum(coalesce(selisih_pengembalian_exact,cast(selisih_pengembalian AS DOUBLE),0)) selisih FROM vw_pengembalian_monthly_current{where}
                GROUP BY 1,2,3 ORDER BY 1,2""",params)
            month_names=[x[0] for x in month_cursor.description]
            by_month=[dict(zip(month_names,row)) for row in month_cursor.fetchall()]
            for item in by_month:
                for key,value in list(item.items()):
                    if isinstance(value,Decimal): item[key]=float(value)
            return {"status":"ok","rows":rows,"by_month":by_month}
        finally:conn.close()

    def export_xlsx(self,filters=None):
        page=self.list_rows(filters,limit=1000000,offset=0)
        wb=Workbook(write_only=True); ws=wb.create_sheet("DATA FINAL")
        headers=["NIP","Nama","Jabatan","Tahun","Bulan","Tahap","Predikat Kinerja","Tukin",
                 "Pengembalian Awal","Final Pengembalian","Selisih","Status Perubahan","Tahun SP2D Awal",
                 "Tahun SP2D Update","Status Send BPK","Status Tendik","Zona Waktu","Provinsi","Kabupaten/Kota",
                 "Cek Selisih","Alasan Perubahan","Keterangan Penyesuaian","Validasi"]
        ws.append(headers)
        keys=["nip","nama","jabatan","tahun","bulan_nama","tahap","predikat_kinerja","tukin",
              "pengembalian_awal","final_pengembalian","selisih_pengembalian","status_perubahan","tahun_sp2d_awal",
              "tahun_sp2d_update","status_send_bpk","status_tendik","zona_waktu","provinsi","kabupaten_kota",
               "cek_selisih","alasan_perubahan","keterangan_penyesuaian","validation_status"]
        for row in page["rows"]: ws.append([row.get(key) for key in keys])
        summary=wb.create_sheet("SUMMARY"); dash=self.dashboard(filters.get("periode_id") if filters else None)
        summary.append(["METRIK","NILAI"])
        for key in ("total_rows","total_sdm","total_awal","total_final","selisih","naik","turun","tidak_berubah","warning_rows"):
            summary.append([key.replace("_"," ").upper(),dash[key]])
        stream=BytesIO(); wb.save(stream); return stream.getvalue()

    def clear(self):
        # The API routes this operation through ArchiveService.clear_database so
        # monthly data and every dependent archive table are cleared together.
        from backend.pengembalian_archive import ArchiveService
        return ArchiveService(self.db_path).clear_database()

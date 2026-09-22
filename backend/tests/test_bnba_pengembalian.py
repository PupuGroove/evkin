import openpyxl

from backend.pengembalian_archive import ArchiveService


def test_bnba_export_contains_current_state_and_traceability(tmp_path):
    service = ArchiveService(str(tmp_path / "bnba.duckdb"))
    source = {
        "NIP": "100",
        "Nama": "Budi",
        "Tahun SP2D": 2025,
        "Nomor SP2D": "SP2D-1",
        "Tahun": 2025,
        "Bulan": 1,
        "Periode": "2025-01",
        "Komponen": "TUKIN",
        "Total Tukin": 1000,
        "Potongan SKP": 10,
        "ESDM": 20,
        "E-Absensi": 30,
        "Tendik": 5,
    }
    service.import_rows([source], "DATA_AWAL", {
        "filename": "BNBA PENGEMBALIAN ALL SDM.xlsx",
        "send_bpk_status": "SEND BPK BATCH 1",
    })
    service.import_rows([dict(source, **{"Total Pengembalian": 25})],
                        "DATA_PENYESUAIAN")

    content = service.export_xlsx(source="DATA_FINAL")
    output = tmp_path / "bnba.xlsx"
    output.write_bytes(content)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    try:
        assert "BNBA PENGEMBALIAN ALL SDM" in workbook.sheetnames
        headers = list(next(workbook["BNBA PENGEMBALIAN ALL SDM"].values))
        assert {
            "NIP", "Tahun SP2D", "Total Pengembalian Awal",
            "Total Pengembalian Final", "Batch BPK", "Comparison Key",
        } <= set(headers)
        assert {"REKAP", "REKAP PROVINSI", "REKAP BATCH"} <= set(workbook.sheetnames)
        recap_headers = list(next(workbook["REKAP"].values))
        assert {"Jumlah SDM", "Total Pengembalian Final", "Total Selisih"} <= set(recap_headers)
    finally:
        workbook.close()

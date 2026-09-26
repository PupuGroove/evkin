"""Regression coverage for the canonical Absensi + SKP -> Tukin -> Final flow."""
from decimal import Decimal
import io

import duckdb
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.final_tukin import FinalTukinService, calculate_final_tukin, create_router
from backend.shared_master import SharedMasterService
from backend.skp_assessment import (
    SkpAssessmentService,
    canonical_skp_reduction_rate,
    classify_skp_report_status,
)
from backend.tukin_annual import _report_status as annual_report_status
from backend.tukin_annual import TukinAnnualService
from backend.tukin_combined import (
    build,
    create_calculation_snapshot,
    create_router as create_combined_router,
    export_calculation_snapshot,
    integration_readiness,
)


PERIOD = "2026-09"
NIP_READY = "198701012026011001"
NIP_NO_ATTENDANCE = "198701012026011002"
NIP_NO_SKP = "198701012026011003"


def setup_integration(tmp_path):
    path = str(tmp_path / "canonical.duckdb")
    master = SharedMasterService(path)
    master.ingest_rows(
        [
            {"NIP": NIP_READY, "Nama": "READY", "Tukin": 10_000_000, "Berhak Tukin": "YA"},
            {"NIP": NIP_NO_ATTENDANCE, "Nama": "NO ABS", "Tukin": 8_000_000, "Berhak Tukin": "YA"},
            {"NIP": NIP_NO_SKP, "Nama": "NO SKP", "Tukin": 6_000_000, "Berhak Tukin": "YA"},
        ],
        "2026-01-01", "master.xlsx", "AUDITOR", True,
    )
    skp = SkpAssessmentService(path)
    with skp.connect() as conn:
        conn.execute(
            """CREATE TABLE final_absensi (
                nip VARCHAR, nama VARCHAR, bulan VARCHAR, potPct DECIMAL(9,4),
                import_batch_id VARCHAR, finalized_at TIMESTAMP
            )"""
        )
        conn.execute(
            """INSERT INTO final_absensi VALUES
               (?, 'READY', ?, 20, 'ABS-BATCH-1', now()),
               (?, 'NO SKP', ?, 5, 'ABS-BATCH-1', now())""",
            [NIP_READY, PERIOD, NIP_NO_SKP, PERIOD],
        )
        conn.execute(
            """INSERT INTO skp_assessment_monthly
               (assessment_period,nip,nama,final_predicate,report_status,review_status,
                import_batch_id,rule_version,approved_by,approved_at)
               VALUES
               (?,?,'READY','Kurang','SUBMITTED','APPROVED','SKP-BATCH-1','RULE-V1','APPROVER',now()),
               (?,?,'NO ABS','Baik','SUBMITTED','APPROVED','SKP-BATCH-1','RULE-V1','APPROVER',now())""",
            [PERIOD, NIP_READY, PERIOD, NIP_NO_ATTENDANCE],
        )
    return path


def complete_all_sources(path):
    """Fill the two deliberate readiness gaps before tests that create Final Tukin."""
    with duckdb.connect(path) as conn:
        conn.execute(
            "INSERT INTO final_absensi VALUES (?, 'NO ABS', ?, 0, 'ABS-BATCH-2', now())",
            [NIP_NO_ATTENDANCE, PERIOD],
        )
        conn.execute(
            """INSERT INTO skp_assessment_monthly
               (assessment_period,nip,nama,final_predicate,report_status,review_status,
                import_batch_id,rule_version,approved_by,approved_at)
               VALUES (?,?,'NO SKP','Baik','SUBMITTED','APPROVED',
                       'SKP-BATCH-2','RULE-V1','APPROVER',now())""",
            [PERIOD, NIP_NO_SKP],
        )


def test_final_attendance_and_approved_skp_are_joined_by_nip_and_period(tmp_path):
    path = setup_integration(tmp_path)
    rows = build(path, PERIOD, canonical=True)
    ready = next(row for row in rows if row["nip"] == NIP_READY)
    assert ready["attendance_participant"] is True
    assert ready["attendance_batch_id"] == "ABS-BATCH-1"
    assert ready["skp_batch_id"] == "SKP-BATCH-1"
    assert ready["skp_source_version"] == "RULE-V1"
    assert ready["attendance_rate"] == 20
    assert ready["skp_rate"] == 10
    assert ready["status"] == "LENGKAP"

    # Same NIP in another period must never be selected.
    with duckdb.connect(path) as conn:
        conn.execute(
            "INSERT INTO final_absensi VALUES (?, 'READY', '2026-08', 100, 'WRONG-PERIOD', now())",
            [NIP_READY],
        )
    ready = next(row for row in build(path, PERIOD, canonical=True) if row["nip"] == NIP_READY)
    assert ready["attendance_rate"] == 20
    assert ready["attendance_batch_id"] == "ABS-BATCH-1"


def test_readiness_detects_missing_sources_and_excludes_mismatch(tmp_path):
    path = setup_integration(tmp_path)
    readiness = integration_readiness(path, PERIOD)
    assert readiness["scope"] == "workflow"
    assert readiness["target_count"] == 2
    assert readiness["master_sdm"] == 2
    assert readiness["attendance_final"] == 2
    assert readiness["skp_approved"] == 1
    assert readiness["ready"] == 2
    assert readiness["mismatch"] == 0
    assert readiness["skp_missing_fallback"] == 1
    assert readiness["warning_count"] == 1
    assert readiness["can_calculate"] is True
    assert readiness["can_finalize"] is False
    assert readiness["blocker_counts"]["skp_not_approved"] == 0
    all_scope = integration_readiness(path, PERIOD, "all")
    assert all_scope["target_count"] == 3
    assert all_scope["mismatch"] == 1
    rows = build(path, PERIOD, canonical=True)
    assert next(row for row in rows if row["nip"] == NIP_NO_ATTENDANCE)["final_tukin"] is None
    fallback = next(row for row in rows if row["nip"] == NIP_NO_SKP)
    assert fallback["final_tukin"] is not None
    assert fallback["skp_source_status"] == "MISSING_FROM_FINAL_IMPORT"


def test_workflow_population_ignores_master_only_employees(tmp_path):
    path = setup_integration(tmp_path)
    master = SharedMasterService(path)
    master.ingest_rows([
        {"NIP": f"999{i:015d}", "Nama": f"MASTER ONLY {i}", "Tukin": 5_000_000,
         "Berhak Tukin": "YA"}
        for i in range(97)
    ], "2026-01-01", "master-only.xlsx", "AUDITOR", True)
    complete_all_sources(path)
    readiness = integration_readiness(path, PERIOD, "workflow")
    assert readiness["master_universe"] == 100
    assert readiness["target_count"] == 3
    assert readiness["complete_count"] == 3
    assert readiness["blocked_count"] == 0
    assert readiness["can_calculate"] is True
    assert readiness["source_ready"] is True


def test_partial_and_empty_workflow_calculation_readiness(tmp_path):
    path = setup_integration(tmp_path)
    partial = integration_readiness(path, PERIOD, "workflow")
    assert (partial["target_count"], partial["complete_count"], partial["blocked_count"]) == (2, 2, 0)
    assert partial["can_calculate"] is True
    with duckdb.connect(path) as conn:
        conn.execute(
            "UPDATE skp_assessment_monthly SET review_status='PENDING_REVIEW' WHERE nip=? AND assessment_period=?",
            [NIP_READY, PERIOD],
        )
        conn.execute(
            """INSERT INTO skp_assessment_monthly
               (assessment_period,nip,nama,final_predicate,report_status,validation_status,review_status)
               VALUES (?,?,'NO SKP','Baik','SUBMITTED','VALID','PENDING_REVIEW')""",
            [PERIOD, NIP_NO_SKP],
        )
    empty = integration_readiness(path, PERIOD, "workflow")
    assert empty["complete_count"] == 0
    assert empty["can_calculate"] is False


def test_scope_isolation_and_freshness_ignore_non_workflow_master(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    assert integration_readiness(path, PERIOD, "workflow")["target_count"] == 3
    assert integration_readiness(path, PERIOD, "attendance")["target_count"] == 3
    assert integration_readiness(path, PERIOD, "all")["target_count"] == 3
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR", scope="workflow")
    SharedMasterService(path).ingest_rows(
        [{"NIP": "999999999999999999", "Nama": "OUTSIDE", "Tukin": 1_000_000,
          "Berhak Tukin": "YA"}],
        "2026-01-01", "outside.xlsx", "AUDITOR", True,
    )
    current = integration_readiness(path, PERIOD, "workflow")
    assert current["calculation"]["calculation_id"] == calculation["calculation_id"]
    assert current["calculation"]["status"] == "CALCULATED"
    assert integration_readiness(path, PERIOD, "all")["target_count"] == 4
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE final_absensi SET potPct=21 WHERE nip=? AND bulan=?", [NIP_READY, PERIOD])
    assert integration_readiness(path, PERIOD, "workflow")["calculation"]["status"] == "STALE"


def test_combined_api_uses_explicit_scope_for_readiness_and_calculation(tmp_path):
    path = setup_integration(tmp_path)
    app = FastAPI()
    app.include_router(create_combined_router(lambda: path))
    client = TestClient(app)
    readiness = client.get(
        "/api/tukin/combined/readiness",
        params={"assessment_period": PERIOD, "scope": "workflow"},
    )
    assert readiness.status_code == 200
    payload = readiness.json()
    assert payload["scope"] == "workflow"
    assert payload["target_count"] == 2
    calculated = client.post("/api/tukin/combined/calculate", json={
        "assessment_period": PERIOD,
        "scope": "workflow",
        "calculated_by": "OPERATOR",
        "expected": payload["source_fingerprint"],
    })
    assert calculated.status_code == 200
    assert calculated.json()["scope"] == "workflow"
    history = client.get(
        "/api/tukin/combined/calculations", params={"assessment_period": PERIOD}
    ).json()["rows"]
    assert history[0]["scope"] == "workflow"
    invalid = client.get(
        "/api/tukin/combined/readiness",
        params={"assessment_period": PERIOD, "scope": "arbitrary"},
    )
    assert invalid.status_code == 400


def test_canonical_skp_classification_is_shared():
    raw = "CUTI TUGAS BELAJAR TIDAK MEMBUAT SKP"
    assert classify_skp_report_status("", raw) == "EXEMPT_OTHER"
    assert annual_report_status("", raw)[0] == "EXEMPT_OTHER"
    assert canonical_skp_reduction_rate("Kurang", "SUBMITTED") == 10
    assert canonical_skp_reduction_rate("", "NOT_SUBMITTED_VERIFIED", 2) == 70


def test_calculation_snapshot_is_versioned_and_preserves_formula(tmp_path):
    path = setup_integration(tmp_path)
    first = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    second = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    assert (first["version"], second["version"]) == (1, 2)
    assert first["calculation_id"] != second["calculation_id"]
    with duckdb.connect(path) as conn:
        batches = conn.execute(
            "SELECT version,count(*) FROM tukin_calculation_batch GROUP BY version ORDER BY version"
        ).fetchall()
        results = conn.execute(
            """SELECT calculation_id,attendance_deduction,skp_deduction,total_deduction,tukin_paid,
                      attendance_batch_id,skp_batch_id
               FROM tukin_calculation_result WHERE nip=? ORDER BY calculated_at""", [NIP_READY]
        ).fetchall()
    assert batches == [(1, 1), (2, 1)]
    assert len(results) == 2
    assert tuple(results[0][1:5]) == (600_000, 700_000, 1_300_000, 8_700_000)
    assert tuple(results[0][5:]) == ("ABS-BATCH-1", "SKP-BATCH-1")
    exported = export_calculation_snapshot(path, first["calculation_id"])
    assert exported.getvalue()[:2] == b"PK"


def test_final_references_snapshot_and_source_change_does_not_mutate_it(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    service = FinalTukinService(path)
    final = service.approve_calculation(
        calculation["calculation_id"], "APPROVER", "Dokumen verifikasi 01"
    )
    assert service.export_final_decision(final["final_decision_id"]).getvalue()[:2] == b"PK"
    with duckdb.connect(path) as conn:
        before = conn.execute(
            "SELECT tukin_paid,snapshot FROM tukin_calculation_result WHERE calculation_id=? AND nip=?",
            [calculation["calculation_id"], NIP_READY],
        ).fetchone()
        conn.execute("UPDATE final_absensi SET potPct=30 WHERE nip=? AND bulan=?", [NIP_READY, PERIOD])
        decision = conn.execute(
            "SELECT calculation_id,calculation_version,status FROM tukin_final_decision WHERE id=?",
            [final["final_decision_id"]],
        ).fetchone()
        after = conn.execute(
            "SELECT tukin_paid,snapshot FROM tukin_calculation_result WHERE calculation_id=? AND nip=?",
            [calculation["calculation_id"], NIP_READY],
        ).fetchone()
    assert decision == (calculation["calculation_id"], 1, "FINAL")
    assert after == before
    assert integration_readiness(path, PERIOD)["calculation"]["status"] == "STALE"


def test_paid_final_cannot_be_downgraded(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    service = FinalTukinService(path)
    final = service.approve_calculation(calculation["calculation_id"], "APPROVER", "Siap bayar")
    assert service.mark_paid(final["final_decision_id"], "BENDAHARA", "SP2D-001")["status"] == "PAID"
    try:
        service.mark_paid(final["final_decision_id"], "BENDAHARA", "SP2D-002")
    except ValueError as exc:
        assert "sudah dibayar" in str(exc)
    else:
        raise AssertionError("PAID decision accepted an ordinary repeated transition")


def test_finalization_remains_strict_for_identity_mismatch_and_freshness(tmp_path):
    path = setup_integration(tmp_path)
    service = FinalTukinService(path)
    with duckdb.connect(path) as conn:
        conn.execute(
            """INSERT INTO skp_assessment_monthly
               (assessment_period,nip,nama,final_predicate,report_status,validation_status,review_status)
               VALUES (?,?,'NO SKP','Baik','SUBMITTED','VALID','PENDING_REVIEW')""",
            [PERIOD, NIP_NO_SKP],
        )
    partial = create_calculation_snapshot(path, PERIOD, "CALCULATOR", scope="workflow")
    with pytest.raises(ValueError, match="masih perlu diperbaiki"):
        service.approve_calculation(partial["calculation_id"], "APPROVER", "Belum lengkap")
    with pytest.raises(ValueError, match="Approver dan alasan"):
        service.approve_calculation(partial["calculation_id"], "", "")

    with duckdb.connect(path) as conn:
        conn.execute("DELETE FROM skp_assessment_monthly WHERE nip=? AND assessment_period=?", [NIP_NO_SKP, PERIOD])
    complete_all_sources(path)
    current = create_calculation_snapshot(path, PERIOD, "CALCULATOR", scope="workflow")
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE final_absensi SET potPct=22 WHERE nip=? AND bulan=?", [NIP_READY, PERIOD])
    with pytest.raises(ValueError, match="berubah"):
        service.approve_calculation(current["calculation_id"], "APPROVER", "Harus hitung ulang")


def test_legacy_final_calculator_uses_same_canonical_amounts():
    row = calculate_final_tukin(
        NIP_READY, PERIOD, "ABSENSI_SKP", {"total": 20}, {"tukin_gross": 10_000_000},
        {"nip": NIP_READY, "periode": PERIOD, "predikat_skp": "KURANG",
         "potongan_skp_percent": 10, "status_skp": "APPROVED", "source_type": "ASSESSMENT"},
    )
    assert row["potongan_absen_rp"] == 600_000
    assert row["potongan_skp_rp"] == 700_000
    assert row["tukin_dibayar"] == 8_700_000
    assert Decimal(str(row["tukin_bruto"])) == Decimal("10000000")


def test_locked_legacy_period_rejects_financial_condition(tmp_path):
    path = str(tmp_path / "locked.duckdb")
    service = TukinAnnualService(path)
    service.import_rows({
        "source_type": "FINAL", "assessment_period": "2026-09",
        "rows": [{"NIP": NIP_READY, "Tukin": 10_000_000}],
    })
    service.lock_period("2026-09", "PAID")
    with pytest.raises(ValueError, match="terkunci|dibayar"):
        service.add_condition({
            "assessment_period": "2026-09", "nip": NIP_READY,
            "condition_code": "HUKDIS", "effect_type": "UNPAID", "reason": "Uji",
        })


def test_governed_unpaid_row_can_close_period(tmp_path):
    path = str(tmp_path / "unpaid.duckdb")
    service = TukinAnnualService(path)
    service.import_rows({
        "source_type": "FINAL", "assessment_period": "2026-09",
        "rows": [{"NIP": NIP_READY, "Tukin": 10_000_000,
                  "STATUS PEMBAYAN": "tidak dibayarkan"}],
    })
    # Set the governed terminal outcome explicitly; UNPAID is complete, not HOLD.
    with service.connect() as conn:
        conn.execute(
            "UPDATE tukin_final_monthly SET payment_status='UNPAID',paid_amount=0 WHERE assessment_period=?",
            [PERIOD],
        )
    service.lock_period(PERIOD, "APPROVED", "APPROVER", "Sudah diverifikasi")
    result = service.lock_period(PERIOD, "LOCKED", "APPROVER", "Tutup periode")
    assert result["period_status"] == "LOCKED"


def test_final_snapshot_ignores_later_absence_and_skp_changes(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    service = FinalTukinService(path)
    final = service.approve_calculation(calculation["calculation_id"], "APPROVER", "Lock v1")
    before = next(row for row in service.final_decision(final["final_decision_id"])["rows"] if row["nip"] == NIP_READY)
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE final_absensi SET potPct=99 WHERE nip=? AND bulan=?", [NIP_READY, PERIOD])
        conn.execute(
            "UPDATE skp_assessment_monthly SET final_predicate='SANGAT KURANG' WHERE nip=? AND assessment_period=?",
            [NIP_READY, PERIOD],
        )
    after = next(row for row in service.final_decision(final["final_decision_id"])["rows"] if row["nip"] == NIP_READY)
    assert after["tukin_dibayar"] == before["tukin_dibayar"] == 8_700_000
    assert after["potongan_absensi"] == before["potongan_absensi"] == 600_000
    assert after["potongan_skp"] == before["potongan_skp"] == 700_000
    assert service.final_decision(final["final_decision_id"])["batch"]["source_status"] == "STALE"


def test_new_calculation_version_can_create_new_final_without_mutating_v1(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    service = FinalTukinService(path)
    v1 = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    final_v1 = service.approve_calculation(v1["calculation_id"], "APPROVER", "Final v1")
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE final_absensi SET potPct=30 WHERE nip=? AND bulan=?", [NIP_READY, PERIOD])
    v2 = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    final_v2 = service.approve_calculation(v2["calculation_id"], "APPROVER", "Final koreksi v2")
    old = service.final_decision(final_v1["final_decision_id"])
    new = service.final_decision(final_v2["final_decision_id"])
    old_row = next(row for row in old["rows"] if row["nip"] == NIP_READY)
    new_row = next(row for row in new["rows"] if row["nip"] == NIP_READY)
    assert old["decision"]["status"] == "FINAL"
    assert old["decision"]["is_current_decision"] is False
    assert old_row["tukin_dibayar"] == 8_700_000
    assert new["decision"]["calculation_version"] == 2
    assert new_row["tukin_dibayar"] == 8_400_000
    with pytest.raises(ValueError, match="digantikan"):
        service.mark_paid(final_v1["final_decision_id"], "BENDAHARA", "SP2D-OLD")


def test_legacy_final_routes_are_canonical_adapters(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    app = FastAPI()
    app.include_router(create_router(lambda: path))
    client = TestClient(app)
    preview = client.post("/api/final-tukin/preview", json={"periode": PERIOD}).json()
    ready = next(row for row in preview["rows"] if row["nip"] == NIP_READY)
    assert preview["deprecated"] is True
    assert preview["preview_signature"] == calculation["calculation_id"]
    assert ready["tukin_dibayar"] == 8_700_000
    response = client.post("/api/final-tukin/finalize", json={
        "preview_signature": preview["preview_signature"], "created_by": "APPROVER", "reason": "Legacy adapter",
    })
    assert response.status_code == 200
    final = response.json()
    legacy_result = client.get("/api/final-tukin/results/" + calculation["calculation_id"]).json()
    assert next(row for row in legacy_result["rows"] if row["nip"] == NIP_READY)["tukin_dibayar"] == 8_700_000
    export = client.get("/api/final-tukin/export/" + calculation["calculation_id"])
    assert export.status_code == 200 and export.content[:2] == b"PK"
    assert client.post("/api/final-tukin/archive/" + calculation["calculation_id"], json={"confirm": True}).status_code == 410
    assert final["calculation_id"] == calculation["calculation_id"]


def test_final_export_contains_decision_reference_and_snapshot_nominal(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    service = FinalTukinService(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    final = service.approve_calculation(calculation["calculation_id"], "APPROVER", "Export audit")
    workbook = load_workbook(io.BytesIO(service.export_final_decision(final["final_decision_id"]).getvalue()), data_only=True)
    metadata = {row[0]: row[1] for row in workbook["FINAL DECISION"].iter_rows(min_row=2, values_only=True)}
    headers = [cell.value for cell in workbook["FINAL TUKIN"][1]]
    values = dict(zip(headers, [cell.value for cell in workbook["FINAL TUKIN"][2]]))
    assert metadata["id"] == final["final_decision_id"]
    assert metadata["calculation_id"] == calculation["calculation_id"]
    assert metadata["calculation_version"] == 1
    assert values["tukin_dibayar"] == 8_700_000


def test_tukin_final_monthly_never_overrides_canonical_final(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    service = FinalTukinService(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    final = service.approve_calculation(calculation["calculation_id"], "APPROVER", "Canonical only")
    with duckdb.connect(path) as conn:
        conn.execute("CREATE TABLE tukin_final_monthly(assessment_period VARCHAR,nip VARCHAR,paid_amount BIGINT)")
        conn.execute("INSERT INTO tukin_final_monthly VALUES (?,?,1)", [PERIOD, NIP_READY])
    row = next(row for row in service.final_decision(final["final_decision_id"])["rows"] if row["nip"] == NIP_READY)
    assert row["tukin_dibayar"] == 8_700_000


def test_legacy_annual_export_and_payment_route_to_canonical_decision(tmp_path):
    path = setup_integration(tmp_path)
    complete_all_sources(path)
    final_service = FinalTukinService(path)
    calculation = create_calculation_snapshot(path, PERIOD, "CALCULATOR")
    final = final_service.approve_calculation(calculation["calculation_id"], "APPROVER", "Canonical payment")
    annual = TukinAnnualService(path)
    annual.ensure_period(PERIOD, "2026-10", "LIVE")
    workbook = load_workbook(io.BytesIO(annual.export_final(PERIOD).getvalue()), data_only=True)
    assert workbook["FINAL DECISION"]["B2"].value == final["final_decision_id"]
    with pytest.raises(ValueError, match="Final Decision canonical"):
        annual.lock_period(PERIOD, "PAID", "BENDAHARA", "Jalur legacy harus ditolak")

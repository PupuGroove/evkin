from io import BytesIO
from pathlib import Path
from decimal import Decimal
import duckdb
import pytest
from openpyxl import Workbook, load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient
from backend.final_tukin import FinalTukinService, calculate_final_tukin, migrate, create_router
from backend.shared_master import SharedMasterService

NIP='198701012026011001'
PERIOD='2026-08'

def calc(mode='ABSENSI_ONLY', skp=None, gross=3915950, rate=1):
    return calculate_final_tukin(NIP,PERIOD,mode,{'total':rate},{'tukin_gross':gross},skp)

def skp(rate=5):
    return dict(nip=NIP,periode=PERIOD,predikat_skp='BUTUH PERBAIKAN',potongan_skp_percent=rate,
                status_skp='IMPORTED_FINAL',source_type='TEMPLATE_IMPORT')

def test_case1_attendance_only():
    row=calc()
    assert row['status']=='READY'
    assert row['komponen_absen_30']==1174785
    assert row['potongan_absen_rp']==11748
    assert row['tukin_dibayar']==3904202
    for field in ['predikat_skp','komponen_skp_70','potongan_skp_percent','potongan_skp_rp']:
        assert row[field] is None

def test_case2_both_components():
    row=calc('ABSENSI_SKP',skp())
    assert row['komponen_skp_70']==2741165
    assert row['potongan_skp_rp']==137058
    assert row['tukin_dibayar']==3767144

def test_case3_optional_missing_skp():
    assert calc()['status']=='READY'

def test_case4_required_missing_skp():
    row=calc('ABSENSI_SKP')
    assert row['status']=='READY' and row['reason']==''
    assert row['skp_source_status']=='MISSING_FROM_FINAL_IMPORT'
    assert row['skp_calculation_mode']=='FALLBACK'
    assert row['potongan_skp_percent']=='35.0'
    assert row['potongan_skp_rp']==959408

def test_case5_duplicate_skp():
    row=calc('ABSENSI_SKP',{**skp(),'_duplicate':True})
    assert row['status']=='HOLD' and row['reason']=='DUPLICATE_SKP'

def test_case6_missing_gross_no_import_fallback():
    row=calc('ABSENSI_SKP',{**skp(),'tukin_gross':999999},gross=None)
    assert row['status']=='HOLD' and 'MISSING_TUKIN' in row['reason']

@pytest.mark.parametrize('gross',[1,2,5,15,3915950,Decimal('3915950.55')])
@pytest.mark.parametrize('rate',[0,1,Decimal('3.3333'),100])
def test_case8_integer_invariant(gross,rate):
    row=calc('ABSENSI_SKP',skp(rate),gross,rate)
    assert row['tukin_bruto']==row['tukin_dibayar']+row['potongan_absen_rp']+row['potongan_skp_rp']
    assert all(isinstance(row[k],int) for k in ['tukin_bruto','tukin_dibayar','potongan_absen_rp','potongan_skp_rp'])
    assert row['tukin_dibayar']>=0

@pytest.mark.parametrize('rate',['NaN','Infinity',-1,101,None])
def test_invalid_percentage_holds(rate):
    assert calc(rate=rate)['status']=='HOLD'
    assert calc('ABSENSI_SKP',skp(rate))['status']=='HOLD'

@pytest.fixture
def service(tmp_path):
    path=str(tmp_path/'test.duckdb')
    master=SharedMasterService(path)
    master.ingest_rows([{'NIP':NIP,'Nama':'Canonical','Tukin':3915950,'Berhak Tukin':'YA'}],
                       '2026-01-01','master.xlsx','Tester',True)
    with duckdb.connect(path) as c:
        c.execute('CREATE TABLE rekap_sdm(nip VARCHAR,bulan VARCHAR,total DECIMAL(9,4))')
        c.execute('INSERT INTO rekap_sdm VALUES(?,?,1)',[NIP,PERIOD])
        c.execute('CREATE TABLE workflow_flags(k VARCHAR,verified INTEGER)')
        c.execute('INSERT INTO workflow_flags VALUES(?,1)',['semifinal:'+PERIOD])
    return FinalTukinService(path)

def workbook(rows,headers=None,formatted=False):
    wb=Workbook();ws=wb.active;ws.title='Final Tukin'
    ws.append(['Referensi']);ws.append(headers or ['NIP PPPK','Nama','Predikat Kinerja Periodik','Potongan SKP (%)'])
    for row in rows:ws.append(row)
    if formatted:ws['D3'].number_format='0%'
    stream=BytesIO();wb.save(stream);return stream.getvalue()

def imported(service,rate=5):
    p=service.preview_import(workbook([[NIP,'Different name','BUTUH PERBAIKAN',rate]]),'skp.xlsx',PERIOD,'TEMPLATE_IMPORT')
    result=service.import_preview(p,p['preview_signature'])
    return result['import_batch_id']

def save(service,mode='ABSENSI_ONLY',batch=None):
    args=dict(periode=PERIOD,mode=mode)
    if mode=='ABSENSI_SKP':args.update(source_type='TEMPLATE_IMPORT',skp_batch_id=batch)
    p=service.preview(**args)
    return service.finalize(**args,preview_signature=p['preview_signature'])

def test_case7_outdated_snapshot_versions_and_restart(service):
    batch=imported(service)
    only=save(service)
    both=save(service,'ABSENSI_SKP',batch)
    with duckdb.connect(service.db_path) as c:
        c.execute('UPDATE skp_results SET potongan_skp_percent=10 WHERE import_batch_id=?',[batch])
    reopened=FinalTukinService(service.db_path)
    result=reopened.results(both['calculation_batch_id'])
    row=result['rows'][0]
    assert row['status']=='OUTDATED' and row['status_label']=='PERLU FINALISASI ULANG'
    assert row['potongan_skp_rp']==137058
    assert reopened.results(only['calculation_batch_id'])['rows'][0]['status']=='FINALIZED'
    new=save(reopened,'ABSENSI_SKP',batch)
    assert new['version']==2 and len(reopened.batches())==3
    assert reopened.results(new['calculation_batch_id'])['rows'][0]['potongan_skp_rp']==274117

def test_safe_migration_and_null_database(service):
    migrate(service.db_path);migrate(service.db_path)
    saved=save(service)
    with duckdb.connect(service.db_path) as c:
        assert c.execute('SELECT count(*) FROM sdm_master').fetchone()[0]==1
        assert c.execute('SELECT count(*) FROM master_skp_rules').fetchone()[0]==5
        assert c.execute('SELECT potongan_skp_percent,potongan_skp_rp,predikat_skp FROM final_tukin_results').fetchone()==(None,None,None)
    service.archive(saved['calculation_batch_id'])
    assert service.batches()==[]
    with duckdb.connect(service.db_path) as c:
        assert c.execute('SELECT count(*) FROM final_tukin_results').fetchone()[0]==1
        assert c.execute('SELECT count(*) FROM rekap_sdm').fetchone()[0]==1

def test_import_preview_no_commit_duplicates_and_repeat(service):
    data=workbook([[NIP,'N','Baik',0],[NIP,'N','Kurang',10]])
    p=service.preview_import(data,'duplicates.xlsx',PERIOD,'EXTERNAL_IMPORT')
    assert p['stats']['duplikat']==2
    with pytest.raises(ValueError):service.import_preview(p,p['preview_signature'])
    with duckdb.connect(service.db_path) as c:
        assert c.execute('SELECT count(*) FROM skp_results').fetchone()[0]==0
    p=service.preview_import(workbook([[NIP,'N','Baik',0]]),'same.xlsx',PERIOD,'EXTERNAL_IMPORT')
    service.import_preview(p,p['preview_signature'])
    with pytest.raises(ValueError,match='sudah diimport'):service.import_preview(p,p['preview_signature'])

def test_import_header_alias_percent_and_rules(service):
    for value,formatted,expected in [(0.05,True,'5.00'),('7,5%',False,'7.5'),(None,False,'5.0000')]:
        p=service.preview_import(workbook([[NIP,'Name','Butuh Perbaikan',value]],formatted=formatted),'x.xlsx',PERIOD,'EXTERNAL_IMPORT')
        assert p['header_row']==1 and not p['stats']['invalid']
        assert Decimal(p['rows'][0]['potongan_skp_percent'])==Decimal(expected)
        assert p['rows'][0]['skp_percentage_source']==('RULE_MAPPING' if value is None else 'IMPORTED')
    data=workbook([[NIP,'Baik']],headers=['ID khusus','Nilai khusus'])
    p=service.preview_import(data,'manual.xlsx',PERIOD,'EXTERNAL_IMPORT',mapping={'nip':0,'predikat_skp':1},header_row=1)
    assert p['stats']['invalid']==0

def test_valid_only_never_picks_a_duplicate_winner(service):
    other='198701012026011002'
    data=workbook([[NIP,'N','Baik',0],[NIP,'N','Kurang',10],[other,'O','Baik',0]])
    p=service.preview_import(data,'conflicts.xlsx',PERIOD,'EXTERNAL_IMPORT')
    result=service.import_preview(p,p['preview_signature'],valid_only=True)
    assert result['imported']==1 and result['excluded']==2
    with duckdb.connect(service.db_path) as c:
        assert c.execute('SELECT nip FROM skp_results').fetchall()==[(other,)]
        assert c.execute('SELECT excluded_count FROM skp_import_batches').fetchone()[0]==2


def test_tn_final_format_imports_unavailable_as_hold(service):
    headers=['NIP','Nama','Jabatan','Rating Hasil Kinerja','Rating Perilaku Kerja',
             'Predikat Kinerja Periodik','Penilai']
    data=workbook([[NIP,'ANI','PLO','Sesuai Ekspektasi','Sesuai Ekspektasi','-','PENILAI']],headers=headers)
    preview=service.preview_import(data,'TN AGUSTUS 2026 08092026.xlsx',PERIOD,'EXTERNAL_IMPORT')
    assert preview['mapping']['nip']==0 and preview['mapping']['predikat_skp']==5
    assert preview['stats']['invalid']==0 and preview['stats']['skp_hold']==1
    saved=service.import_preview(preview,preview['preview_signature'])
    final_preview=service.preview(PERIOD,'ABSENSI_SKP','EXTERNAL_IMPORT',saved['import_batch_id'])
    row=final_preview['rows'][0]
    assert row['status']=='HOLD' and row['reason']=='SKP_UNAVAILABLE'
    assert row['potongan_skp_percent'] is None and row['potongan_skp_rp'] is None


def test_tn_final_format_duplicate_rows_are_all_excluded(service):
    other='198701012026011002'
    headers=['NIP','Nama','Jabatan','Rating Hasil Kinerja','Rating Perilaku Kerja',
             'Predikat Kinerja Periodik','Penilai']
    data=workbook([[NIP,'ANI','PLO','Sesuai','Sesuai','Baik','P'],
                   [NIP,'ANI','PLO','Sesuai','Sesuai','Kurang','P'],
                   [other,'BUDI','PLO','Sesuai','Sesuai','Baik','P']],headers=headers)
    preview=service.preview_import(data,'TN.xlsx',PERIOD,'EXTERNAL_IMPORT')
    assert preview['stats']['duplikat']==2
    saved=service.import_preview(preview,preview['preview_signature'],valid_only=True)
    assert saved['imported']==1 and saved['excluded']==2
    with duckdb.connect(service.db_path) as c:
        assert c.execute('SELECT nip FROM skp_results').fetchall()==[(other,)]

def test_real_template_headers(service):
    data=(Path(__file__).resolve().parents[2]/'assets/templates/Template_Penilaian_SKP.xlsx').read_bytes()
    p=service.preview_import(data,'Template_Penilaian_SKP.xlsx',PERIOD,'TEMPLATE_IMPORT')
    assert p['sheet']=='DATA' and p['mapping']['nip']==0 and p['mapping']['predikat_skp']==4

def test_reference_percentage_beats_nominal_and_indonesian_month(service):
    headers=['NIP','Tahun','Bulan','Predikat Kinerja Periodik','Potongan SKP (%)','Potongan SKP','Tukin']
    data=workbook([[NIP,2026,'Agustus','-', '35,00%',959408,999999]],headers=headers)
    p=service.preview_import(data,'reference.xlsx',PERIOD,'EXTERNAL_IMPORT')
    assert p['mapping']['potongan_skp_percent']==4
    assert p['rows'][0]['periode']==PERIOD and p['rows'][0]['potongan_skp_percent']=='35.00'
    assert p['stats']['invalid']==0

def test_assessment_approval_and_persisted_normalized_link(service):
    with duckdb.connect(service.db_path) as c:
        c.execute('''CREATE TABLE skp_assessment_monthly(nip VARCHAR,assessment_period VARCHAR,
            final_predicate VARCHAR,report_status VARCHAR,review_status VARCHAR)''')
        c.execute("INSERT INTO skp_assessment_monthly VALUES(?,?,'BAIK','SUBMITTED','PENDING_REVIEW')",[NIP,PERIOD])
    args=dict(periode=PERIOD,mode='ABSENSI_SKP',source_type='ASSESSMENT')
    assert service.preview(**args)['rows'][0]['reason']=='SKP_NOT_FINAL'
    with duckdb.connect(service.db_path) as c:c.execute("UPDATE skp_assessment_monthly SET review_status='APPROVED'")
    p=service.preview(**args)
    saved=service.finalize(**args,preview_signature=p['preview_signature'])
    with duckdb.connect(service.db_path) as c:
        assert c.execute('''SELECT count(*) FROM final_tukin_results f JOIN skp_results s ON s.id=f.skp_result_id
                            WHERE f.calculation_batch_id=?''',[saved['calculation_batch_id']]).fetchone()[0]==1
    assert service.results(saved['calculation_batch_id'])['rows'][0]['status']=='FINALIZED'
    with duckdb.connect(service.db_path) as c:c.execute("UPDATE skp_assessment_monthly SET review_status='RETURNED'")
    assert service.results(saved['calculation_batch_id'])['rows'][0]['status']=='OUTDATED'

def test_changed_sources_reject_final_and_mark_outdated(service):
    p=service.preview(PERIOD)
    with duckdb.connect(service.db_path) as c:c.execute('UPDATE rekap_sdm SET total=2')
    with pytest.raises(ValueError,match='berubah'):service.finalize(PERIOD,'ABSENSI_ONLY',p['preview_signature'])
    final=save(service)
    with duckdb.connect(service.db_path) as c:c.execute('UPDATE sdm_tukin_profile_history SET tukin_gross=4000000')
    r=service.results(final['calculation_batch_id'])['rows'][0]
    assert r['status']=='OUTDATED' and r['tukin_bruto']==3915950

def test_unverified_rekap_and_other_period(service):
    with duckdb.connect(service.db_path) as c:c.execute('UPDATE workflow_flags SET verified=0')
    assert service.preview(PERIOD)['rows'][0]['reason']=='ATTENDANCE_NOT_VERIFIED'
    assert service.preview('2026-09')['rows']==[]

def test_legacy_current_master_fallback_is_explicit(service):
    with duckdb.connect(service.db_path) as c:
        c.execute("UPDATE sdm_tukin_profile_history SET valid_from='2026-09-01'")
    row=service.preview(PERIOD)['rows'][0]
    assert row['status']=='READY' and row['tukin_bruto']==3915950
    assert row['tukin_bruto_source']=='MASTER SDM TERKINI'
    assert 'Histori bruto periode belum tersedia' in row['warning']

def test_export_reconciles_and_preserves_nip(service):
    saved=save(service)
    wb=load_workbook(service.export(saved['calculation_batch_id']),data_only=True)
    assert wb.sheetnames==['BNBA FINAL','REKAP','DATA BERMASALAH','METADATA PERHITUNGAN']
    rows=list(wb['BNBA FINAL'].values);r=dict(zip(rows[0],rows[1]))
    assert r['nip']==NIP and r['potongan_skp_rp'] is None
    assert isinstance(r['potongan_absen_percent'],(int,float))
    assert r['tukin_bruto']==r['potongan_absen_rp']+r['tukin_dibayar']

def test_api_legacy_preview_requires_canonical_snapshot(service):
    app=FastAPI();app.include_router(create_router(lambda:service.db_path));client=TestClient(app)
    response=client.post('/api/final-tukin/preview',json={'periode':PERIOD})
    assert response.status_code==400
    assert 'calculation snapshot canonical' in response.json()['detail']
    assert client.get('/api/final-tukin/batches').json()['rows']==[]

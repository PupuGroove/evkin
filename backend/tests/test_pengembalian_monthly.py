from io import BytesIO
import importlib.util
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from backend.pengembalian_monthly import MonthlyMasterService


ROOT = Path(__file__).resolve().parents[1]


def _migration(filename, db_path):
    path = ROOT / "migrations" / filename
    spec = importlib.util.spec_from_file_location(filename.replace(".py", ""), path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.main(db_path)


def _database(tmp_path):
    path = tmp_path / "monthly.duckdb"
    for filename in (
        "023_create_pengembalian_archive_compare.py",
        "028_create_pengembalian_progress.py",
        "029_create_pengembalian_archive_v2.py",
        "030_create_pengembalian_progress_v2.py",
        "031_optimize_pengembalian_progress_import.py",
        "033_create_multi_period_master.py",
    ):
        _migration(filename, path)
    return path


def _workbook(finals=(100, 200)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([
        "Penilai", "NIP", "Nama", "Jabatan", "Tahap", "Predikat Kinerja Periodik",
        "Tahun", "Bulan", "Potongan SKP (%)", "Potongan Absensi eSDM (%)",
        "Potongan Absensi e-Absensi (%)", "Tukin", "Potongan SKP",
        "Potongan Kehadiran eSDM", "Potongan Kehadiran e-Absensi",
        "Tukin Dibayar (eSDM)", "Tukin Dibayar (e-Absensi)", "FINAL KEHADIRAN",
        "FINAL PENGEMBALIAN", "POTONGAN KEHADIRAN AWAL", "PENGEMBALIAN AWAL",
        "STATUS SP2D UPDATE", "KDASATKER UPDATE", "STATUS SP2D", "KDASATKER",
        "STATUS SEND BPK", "STATUS TENDIK", "ZONA WAKTU", "PROPINSI", "KABUPATEN",
        "CUTI", "TANGGAL MULAI CUTI", "TANGGAL SELESAI CUTI", "TENDIK HITUNG ABSEN",
        "PRESENTASE", "ABSEN FINAL TENDIK", "PENGEMBALIAN FINAL TENDIK",
        "STATUS PENGEMBALIAN FINAL", "STATUS PENGEMBALIAN AWAL", "cek selisih", "S", "CEK BUP",
    ])
    for index, (month, final) in enumerate(zip(("Oktober", "November"), finals), 1):
        presence = final - 10
        ws.append([
            "PENILAI", "198001012025211001", "PEGAWAI", "PENATA", "Periode I", "Baik",
            2025, month, "0,00%", "1,00%", "2,00%", 1000, 10, 10, 20, 980, 970,
            presence, final, presence + 5, final + 5, 2025, "P1", 2025, "P1",
            "TAHAP 1 SEND BPK", "bukan tendik", "WIB", "JAWA BARAT", "KAB. BANDUNG",
            "TIDAK CUTI", None, None, "#REF!" if index == 1 else 0,
            "#REF!" if index == 1 else 0, "#REF!" if index == 1 else 0,
            "#REF!" if index == 1 else 0, "PENGEMBALIAN", "PENGEMBALIAN", 0, 0, "20 tahun",
        ])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _workbook_with_rekap_first():
    wb = load_workbook(BytesIO(_workbook()))
    ws = wb.create_sheet("rekap 3 bulan", 0)
    ws.append([None, None, "STATUS SP2D", "Values"])
    ws.append(["STATUS SEND BPK", "Bulan", "Count", "FINAL PENGEMBALIAN"])
    ws.append(["TAHAP 1 SEND BPK", "Oktober", 1, 100])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_monthly_import_is_canonical_and_syncs_obligation(tmp_path):
    db = _database(tmp_path)
    service = MonthlyMasterService(db)
    preview = service.preview_file(_workbook(), "master.xlsx")
    assert preview["summary"]["total_rows"] == 2
    assert preview["summary"]["warning_rows"] == 1
    assert preview["summary"]["total_final"] == 300

    result = service.import_file(_workbook(), "master.xlsx", "PENGEMBALIAN-2025", "tester")
    assert result["summary"]["invalid_rows"] == 0
    conn = duckdb.connect(str(db), read_only=True)
    try:
        assert conn.execute("select count(*) from vw_pengembalian_monthly_current").fetchone()[0] == 2
        obligation = conn.execute("""select final_tukin,final_total,sumber_sp2d
            from pengembalian_master_obligation""").fetchone()
        assert float(obligation[0]) == 300
        assert float(obligation[1]) == 300
        assert obligation[2] == "MASTER_UPDATE_21062026"
    finally:
        conn.close()


def test_new_version_supersedes_and_old_version_can_be_reactivated(tmp_path):
    db = _database(tmp_path)
    service = MonthlyMasterService(db)
    first = service.import_file(_workbook(), "master-v1.xlsx", "PENGEMBALIAN-2025", "tester")
    second = service.import_file(_workbook((50, 75)), "master-v2.xlsx", "PENGEMBALIAN-2025", "tester")
    assert service.dashboard()["total_final"] == 125
    service.activate(first["import_id"], "reviewer")
    assert service.dashboard()["total_final"] == 300
    imports = service.imports()["rows"]
    assert sum(1 for row in imports if row["is_active"]) == 1
    assert next(row for row in imports if row["is_active"])["import_id"] == first["import_id"]


def test_superseded_version_can_be_deleted_but_active_version_cannot(tmp_path):
    db = _database(tmp_path)
    service = MonthlyMasterService(db)
    first = service.import_file(_workbook(), "master-v1.xlsx", "PENGEMBALIAN-2025", "tester")
    second = service.import_file(_workbook((50, 75)), "master-v2.xlsx", "PENGEMBALIAN-2025", "tester")

    with pytest.raises(ValueError, match="aktif"):
        service.delete_import(second["import_id"])

    result = service.delete_import(first["import_id"])
    assert result["deleted_rows"] > 0
    assert all(row["import_id"] != first["import_id"] for row in service.imports()["rows"])


def test_rekap_sheet_is_skipped_and_bnba_reason_is_traceable(tmp_path):
    db = _database(tmp_path)
    service = MonthlyMasterService(db)
    content = _workbook_with_rekap_first()
    preview = service.preview_file(content, "21062026 MASTER UPDATE- (1).xlsx")
    assert preview["sheet"] == "Sheet1"
    assert preview["summary"]["total_rows"] == 2
    service.import_file(content, "21062026 MASTER UPDATE- (1).xlsx", "PENGEMBALIAN-2025", "tester")
    dashboard = service.dashboard()
    assert dashboard["by_status_month"][0]["status_send_bpk"] == "TAHAP 1 SEND BPK"
    assert dashboard["by_reason"][0]["alasan_perubahan"] == "PERUBAHAN TANPA KETERANGAN"
    assert service.reasons()["rows"][0]["rows"] == 2
    trace = service.bnba_trace("198001012025211001")
    assert trace["summary"]["months"] == 2
    assert trace["rows"][0]["source_filename"] == "21062026 MASTER UPDATE- (1).xlsx"


def test_reset_source_database_preserves_progress_ledger(tmp_path):
    db = _database(tmp_path)
    service = MonthlyMasterService(db)
    service.import_file(_workbook(), "master.xlsx", "PENGEMBALIAN-2025", "tester")
    conn = duckdb.connect(str(db))
    try:
        conn.execute("""INSERT INTO pengembalian_progress_import_batch
            (import_batch_id,status,total_rows) VALUES ('PROGRESS-KEEP','COMMITTED',1)""")
    finally:
        conn.close()

    result = service.reset_source_database("tester")
    assert result["progress_preserved"] is True
    assert result["progress_rows"] == 0

    conn = duckdb.connect(str(db), read_only=True)
    try:
        assert conn.execute("SELECT count(*) FROM pengembalian_monthly_import").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM pengembalian_monthly_master").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM pengembalian_monthly_source_row").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM pengembalian_master_obligation").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM pengembalian_progress_import_batch").fetchone()[0] == 1
    finally:
        conn.close()

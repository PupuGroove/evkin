from io import BytesIO
import duckdb
import pandas as pd
import pytest

from backend.pengembalian_archive import ArchiveService
from backend.pengembalian_progress import ProgressService, _expand_master_progress_rows, normalize_bank


def service(tmp_path):
    path = tmp_path / "progress-v2.duckdb"
    ArchiveService(str(path)).repository.ensure_schema()
    progress = ProgressService(str(path))
    conn = duckdb.connect(str(path))
    conn.executemany("""INSERT INTO pengembalian_obligation
        (obligation_id,nip,tahun_sp2d,nama_sdm,provinsi,final_tukin,final_uang_makan,final_total)
        VALUES (?,?,?,?,?,?,?,?)""", [
        ["O-A-2025","A",2025,"SDM A","Jawa Barat",1_000_000,0,1_000_000],
        ["O-B-2025","B",2025,"SDM B","Jawa Barat",1_000_000,500_000,1_500_000],
        ["O-B-2026","B",2026,"SDM B","Jawa Barat",500_000,0,500_000],
    ])
    conn.close()
    return progress


def row(nip="A", year=2025, tukin=0, um=0, status="BERHASIL", period="Juli 2026", date="18/07/2026"):
    return {"NIP": nip, "Nama SDM": "Test", "Tahun SP2D": year, "Provinsi": "Jawa Barat",
            "Nama Bank": "Bank BRI", "Periode Debit": period, "Tanggal Debit": date,
            "Status Debit": status, "Debit Tukin": tukin, "Debit Uang Makan": um}


def find(result, nip, year):
    return next(x for x in result["rows"] if x["nip"] == nip and x["tahun_sp2d"] == year)


def test_bank_normalization():
    assert normalize_bank("BRI") == normalize_bank("Bank BRI") == normalize_bank("PT Bank Rakyat Indonesia")


def test_wide_master_progress_is_expanded_into_stage_events():
    rows = [{
        "NIP": "001", "Nama sk": "SDM", "Provinsi": "Jawa Barat",
        "Tahun Sp2d UPDATE": 2025, "NAMA BANK": "BANK MANDIRI",
        "FINAL TOTAL PENGEMBALIAN TUKIN": 1_000,
        "SEND BANK 1  (DARI TOTAL TUKIN DIBAGI 3 BULAN DEBITAN BAGI YANG >500RB BAGI <500RB LANGSUNG DIBULAN JULI)": 1_000,
        "HASIL DEBIT JULI SP2D 2025": 700, "NTPN": "NTPN-1",
        "HASIL DEBIT SP2D 2026 ": 0, "NTPN_1": "NTPN-2",
        "GAGAL": 200, "KETERANGAN GAGAL": "Saldo kurang",
        "SEND YANG GAGAL TAHAP 2": 0, "HASIL DEBIT TAHAP 2 2025": 0,
        "HASIL DEBIT TAHAP 2 SP2D 2026": 0, "GAGAL TAHAP 2": 0,
        "SEND BANK TAHAP 3": 0, "HASIL DEBIT AGUSTUS TAHAP 3  SP2D 2025": 0,
        "HASIL DEBIT TAHAP 3 SP2D 2026": 0,
        "TOTAL PROGRESS TUKIN SP2D TAHUH 2025'": 700,
    }]
    expanded, adapter = _expand_master_progress_rows(rows)
    assert adapter == {"format": "MASTER_PROGRESS_WIDE", "source_rows": 1, "transaction_rows": 3}
    assert [(r["Status Debit"], r["Debit Tukin"]) for r in expanded] == [
        ("BERHASIL", 700), ("GAGAL", 200), ("PENDING", 100)
    ]
    assert expanded[0]["Tahun SP2D"] == 2025
    assert expanded[0]["NTPN"] == "NTPN-1"
    assert all(r["Source Row Number"] == 2 for r in expanded)


def test_tukin_only_and_failed_do_not_reduce_receivable(tmp_path):
    svc = service(tmp_path)
    result = svc.import_rows([row(tukin=500_000), row(tukin=500_000, status="GAGAL", date="19/07/2026")])
    assert result["inserted"] == 2
    item = find(svc.progress(limit=20), "A", 2025)
    assert item["sudah_debit_tukin"] == 500_000
    assert item["sisa_tukin"] == 500_000
    assert item["nominal_debit_gagal"] == 500_000
    assert item["total_sisa_piutang"] == 500_000
    assert item["status_progress"] == "PARSIAL"


def test_component_specific_and_cumulative_months(tmp_path):
    svc = service(tmp_path)
    svc.import_rows([
        row("B",2025,300_000,0,period="Juli 2026"),
        row("B",2025,200_000,100_000,period="Agustus 2026",date="18/08/2026"),
        row("B",2025,0,100_000,period="September 2026",date="18/09/2026"),
    ])
    item=find(svc.progress(limit=20),"B",2025)
    assert (item["sudah_debit_tukin"],item["sudah_debit_uang_makan"]) == (500_000,200_000)
    assert (item["sisa_tukin"],item["sisa_uang_makan"],item["total_sisa_piutang"]) == (500_000,300_000,800_000)
    periods=svc.recaps()["periode"]
    assert [x["akumulasi_berhasil"] for x in periods] == [300_000,600_000,700_000]


def test_duplicate_is_invalid_and_not_double_counted(tmp_path):
    svc=service(tmp_path); source=row(tukin=500_000)
    assert svc.import_rows([source])["inserted"] == 1
    duplicate=svc.preview([source])
    assert duplicate["statistics"]["duplicate"] == 1
    with pytest.raises(ValueError, match="FILE_ALREADY_IMPORTED"):
        svc.import_rows([source])
    assert find(svc.progress(limit=20),"A",2025)["sudah_debit_tukin"] == 500_000


def test_same_nip_different_sp2d_is_separate(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,100_000),row("B",2026,200_000,date="19/07/2026")])
    result=svc.progress(limit=20)
    assert find(result,"B",2025)["sudah_debit_tukin"] == 100_000
    assert find(result,"B",2026)["sudah_debit_tukin"] == 200_000


def test_lunas_and_overpayment(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row(tukin=1_000_000)])
    assert find(svc.progress(limit=20),"A",2025)["status_progress"] == "LUNAS"
    svc.import_rows([row(tukin=100_000,date="19/07/2026")])
    item=find(svc.progress(limit=20),"A",2025)
    assert item["status_progress"] == "OVERPAYMENT"
    assert item["kelebihan_bayar"] == 100_000


def test_tukin_overpayment_does_not_reduce_uang_makan_receivable(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,1_100_000,0)])
    item=find(svc.progress(limit=20),"B",2025)
    assert item["kelebihan_bayar"] == 100_000
    assert item["sisa_tukin"] == 0
    assert item["sisa_uang_makan"] == 500_000
    assert item["total_sisa_piutang"] == 500_000
    assert item["status_progress"] == "OVERPAYMENT"


def test_uang_makan_overpayment_does_not_reduce_tukin_receivable(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,0,600_000)])
    item=find(svc.progress(limit=20),"B",2025)
    assert item["kelebihan_bayar"] == 100_000
    assert item["sisa_tukin"] == 1_000_000
    assert item["sisa_uang_makan"] == 0
    assert item["total_sisa_piutang"] == 1_000_000
    assert item["status_progress"] == "OVERPAYMENT"


def test_both_components_overpaid_are_summed(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,1_100_000,700_000)])
    item=find(svc.progress(limit=20),"B",2025)
    assert item["kelebihan_bayar"] == 300_000
    assert item["sisa_tukin"] == 0
    assert item["sisa_uang_makan"] == 0
    assert item["total_sisa_piutang"] == 0


def test_both_components_still_have_receivables(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,600_000,200_000)])
    item=find(svc.progress(limit=20),"B",2025)
    assert item["kelebihan_bayar"] == 0
    assert item["sisa_tukin"] == 400_000
    assert item["sisa_uang_makan"] == 300_000
    assert item["total_sisa_piutang"] == 700_000
    assert item["status_progress"] == "PARSIAL"


def test_exact_payment_has_no_remaining_or_overpayment(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,1_000_000,500_000)])
    item=find(svc.progress(limit=20),"B",2025)
    assert item["kelebihan_bayar"] == 0
    assert item["sisa_tukin"] == 0
    assert item["sisa_uang_makan"] == 0
    assert item["total_sisa_piutang"] == 0
    assert item["status_progress"] == "LUNAS"


def test_reversal_recomputes_component_remaining_without_cross_offset(tmp_path):
    svc=service(tmp_path)
    first=svc.preview([row("B",2025,1_100_000,0)],{"filename":"paid.xlsx","file_hash":"paid-over"})
    svc.commit_batch(first["import_batch_id"])
    original_id=svc.history({"import_batch_id":first["import_batch_id"]})["rows"][0]["transaction_id"]
    reversal=row("B",2025,200_000,0,status="REVERSAL",period="Agustus 2026",date="18/08/2026")
    reversal["Original Transaction ID"]=original_id
    second=svc.preview([reversal],{"filename":"reversal.xlsx","file_hash":"reverse-over"})
    svc.commit_batch(second["import_batch_id"])

    item=find(svc.progress(limit=20),"B",2025)
    assert item["sudah_debit_tukin"] == 900_000
    assert item["kelebihan_bayar"] == 0
    assert item["sisa_tukin"] == 100_000
    assert item["sisa_uang_makan"] == 500_000
    assert item["total_sisa_piutang"] == 600_000


def test_pagination_summary_and_export(tmp_path):
    svc=service(tmp_path)
    result=svc.progress(limit=1,offset=0)
    assert len(result["rows"]) == 1 and result["total"] == 3
    assert result["summary"]["total_kewajiban"] == 3_000_000
    assert result["summary"]["total_sdm"] == 3
    workbook=pd.ExcelFile(BytesIO(svc.export_xlsx({"tahun_sp2d":2025})))
    expected={"SUMMARY","PROGRESS_SDM","HISTORY_DEBIT","REKAP_PROVINSI","REKAP_BANK","REKAP_SP2D","REKAP_PERIODE","DEBIT_GAGAL","SISA_PIUTANG","DATA_LUNAS","NEEDS_REVIEW","MULTI_SP2D"}
    assert expected.issubset(set(workbook.sheet_names))


def test_export_history_and_period_recaps_follow_progress_filter(tmp_path):
    svc=service(tmp_path)
    svc.import_rows([row("B",2025,100_000,period="Juli 2026"),
                     row("B",2026,200_000,period="Agustus 2026",date="18/08/2026")])
    content=svc.export_xlsx({"tahun_sp2d":2025})
    history=pd.read_excel(BytesIO(content),sheet_name="HISTORY_DEBIT")
    periods=pd.read_excel(BytesIO(content),sheet_name="REKAP_PERIODE")
    assert set(history["tahun_sp2d"].astype(int)) == {2025}
    assert set(periods["periode_debit"]) == {"Juli 2026"}


def test_keyset_and_tenor_filter_scale_to_more_than_50000_rows(tmp_path):
    svc = service(tmp_path)
    conn = duckdb.connect(svc.db_path)
    conn.execute("""
        INSERT INTO pengembalian_obligation
            (obligation_id,nip,tahun_sp2d,nama_sdm,final_tukin,final_total)
        SELECT 'SCALE-' || i::VARCHAR, lpad(i::VARCHAR, 6, '0'), 2025,
               'Scale ' || i::VARCHAR, 100, 100
        FROM range(1, 50002) AS values(i)
    """)
    conn.execute("""
        INSERT INTO pengembalian_progress_bulanan
            (progress_id,nip,nama,tenor_bulan)
        VALUES ('TENOR-1','050000','Scale 50000',12),
               ('TENOR-2','050001','Scale 50001',12)
    """)
    conn.close()

    first = svc.progress({"tenor_bulan": 12}, limit=1)
    assert first["total"] == 2
    assert first["summary"]["total_sdm"] == 2
    assert first["has_more"] is True
    second = svc.progress({"tenor_bulan": 12}, limit=1, cursor=first["next_cursor"])
    assert [row["nip"] for row in second["rows"]] == ["050001"]
    assert second["has_more"] is False


def test_template_has_import_and_guide(tmp_path):
    svc=service(tmp_path)
    workbook=pd.ExcelFile(BytesIO(svc.template_xlsx()))
    assert workbook.sheet_names == ["IMPORT_PROGRESS","PETUNJUK"]


def test_preview_uses_staging_without_final_transaction(tmp_path):
    svc=service(tmp_path)
    preview=svc.preview([row(tukin=100_000)],{"filename":"juli.xlsx","file_hash":"abc"})
    conn=duckdb.connect(svc.db_path)
    assert conn.execute("SELECT status FROM pengembalian_progress_import_batch WHERE import_batch_id=?",[preview["import_batch_id"]]).fetchone()[0] == "READY"
    assert conn.execute("SELECT count(*) FROM pengembalian_progress_staging WHERE import_batch_id=?",[preview["import_batch_id"]]).fetchone()[0] == 1
    assert conn.execute("SELECT count(*) FROM pengembalian_progress_transaction").fetchone()[0] == 0
    conn.close()
    assert preview["impact"]["proyeksi_sisa"] == 2_900_000


def test_commit_reversal_and_soft_rollback(tmp_path):
    svc=service(tmp_path)
    first=svc.preview([row(tukin=500_000)],{"filename":"juli.xlsx","file_hash":"juli"})
    committed=svc.commit_batch(first["import_batch_id"])
    assert committed["inserted"] == 1
    original_id=svc.history({"import_batch_id":first["import_batch_id"]})["rows"][0]["transaction_id"]
    reversal=row(tukin=500_000,status="REVERSAL",period="Agustus 2026",date="18/08/2026")
    reversal["Original Transaction ID"]=original_id
    second=svc.preview([reversal],{"filename":"agustus.xlsx","file_hash":"agustus"})
    svc.commit_batch(second["import_batch_id"])
    assert find(svc.progress(limit=20),"A",2025)["sudah_debit_tukin"] == 0
    rolled=svc.rollback_batch(second["import_batch_id"])
    assert rolled["status_batch"] == "REVERTED"
    assert find(svc.progress(limit=20),"A",2025)["sudah_debit_tukin"] == 500_000
    history=svc.history({"import_batch_id":second["import_batch_id"]})["rows"]
    assert history[0]["is_active"] is True
    reversal_history=svc.history({"import_batch_id":rolled["revert_batch_id"]})["rows"]
    assert reversal_history[0]["correction_type"] == "REVERT_BATCH"


def test_reversal_requires_original_and_cannot_exceed_it(tmp_path):
    svc=service(tmp_path)
    first=svc.preview([row(tukin=200_000)],{"filename":"paid.xlsx","file_hash":"paid"})
    svc.commit_batch(first["import_batch_id"])
    original_id=svc.history({"import_batch_id":first["import_batch_id"]})["rows"][0]["transaction_id"]
    missing=svc.preview([row(tukin=100_000,status="REVERSAL",date="20/07/2026")],{"file_hash":"missing"})
    assert "REVERSAL_ORIGINAL_REQUIRED" in missing["rows"][0]["errors"]
    excessive=row(tukin=300_000,status="REVERSAL",date="21/07/2026")
    excessive["Original Transaction ID"]=original_id
    checked=svc.preview([excessive],{"file_hash":"excessive"})
    assert "REVERSAL_EXCEEDS_ORIGINAL" in checked["rows"][0]["errors"]


def test_ambiguous_multi_period_obligation_requires_period_id(tmp_path):
    svc=service(tmp_path)
    conn=duckdb.connect(svc.db_path)
    conn.execute("""INSERT INTO pengembalian_master_obligation
        (obligation_id,periode_id,nip,tahun_sp2d,nama_sdm,final_tukin,final_total)
        VALUES ('MP-1','PERIOD-A','M',2025,'Multi',100,100),
               ('MP-2','PERIOD-B','M',2025,'Multi',200,200)""")
    conn.close()
    ambiguous=svc.preview([row("M",2025,100)],{"file_hash":"ambiguous"})
    assert ambiguous["rows"][0]["classification"] == "NEEDS_REVIEW"
    assert ambiguous["statistics"]["ambiguous_obligation"] == 1
    selected=row("M",2025,100,date="19/07/2026")
    selected["Periode Master"]="PERIOD-A"
    matched=svc.preview([selected],{"file_hash":"selected"})
    assert matched["rows"][0]["obligation_id"] == "MP-1"


def test_commit_can_exclude_warning_rows(tmp_path):
    svc=service(tmp_path)
    warning=row(tukin=100_000)
    warning["Tanggal Debit"]=""
    preview=svc.preview([warning],{"file_hash":"warning"})
    assert preview["rows"][0]["classification"] == "WARNING"
    committed=svc.commit_batch(preview["import_batch_id"],include_warnings=False)
    assert committed["inserted"] == 0


def test_total_only_requires_component_allocation_and_validation_download(tmp_path):
    svc=service(tmp_path)
    source=row(tukin="",um="")
    source["Nominal Debit Total"]=100_000
    preview=svc.preview([source],{"filename":"total.xlsx","file_hash":"total"})
    item=preview["rows"][0]
    assert item["classification"] == "NEEDS_REVIEW"
    assert "COMPONENT_ALLOCATION_REQUIRED" in item["errors"]
    result=svc.commit_batch(preview["import_batch_id"])
    assert result["inserted"] == 0
    workbook=pd.ExcelFile(BytesIO(svc.validation_result_xlsx(preview["import_batch_id"])))
    assert workbook.sheet_names == ["VALID","WARNING","INVALID","DUPLICATE","NEEDS_REVIEW"]


def test_file_hash_duplicate_is_blocked(tmp_path):
    svc=service(tmp_path)
    first=svc.preview([row(tukin=100_000)],{"filename":"a.xlsx","file_hash":"same"})
    svc.commit_batch(first["import_batch_id"])
    second=svc.preview([row(tukin=200_000,date="19/07/2026")],{"filename":"copy.xlsx","file_hash":"same"})
    assert second["file_duplicate"]["warning"] == "FILE_ALREADY_IMPORTED"
    conn=duckdb.connect(svc.db_path)
    assert conn.execute("SELECT status FROM pengembalian_progress_import_batch WHERE import_batch_id=?",[second["import_batch_id"]]).fetchone()[0] == "FAILED"
    conn.close()
    with pytest.raises(ValueError,match="FAILED"):
        svc.commit_batch(second["import_batch_id"])


def test_unlimited_stage_ntpn_and_bank_history(tmp_path):
    svc=service(tmp_path)
    rows=[]
    for stage,bank,ntpn,amount in [(5,"MANDIRI","AAA",300_000),(5,"BRI","BBB",200_000),(15,"BSI","CCC",100_000)]:
        rows.append({"NIP":"A","Nama":"SDM A","Provinsi":"Jawa Barat","Bank":bank,
            "Nomor Rekening":f"R-{bank}","Tahun SP2D":2025,"Komponen":"TUKIN","Tahap":stage,
            "Nominal Diajukan":amount,"Nominal Berhasil":amount,"Status":"BERHASIL",
            "NTPN":ntpn,"Tanggal Proses":f"{stage:02d}/09/2026"})
    result=svc.import_rows(rows,{"periode_id":"LEGACY","import_mode":"PROGRESS_STAGE_IMPORT","file_hash":"multi-bank"})
    assert result["inserted"] == 3
    history=svc.history({"nip":"A"},limit=10)["rows"]
    assert {item["nama_bank_normalized"] for item in history} == {"MANDIRI","BRI","BSI"}
    assert {item["tahap"] for item in history} == {5,15}
    assert {item["ntpn"] for item in history} == {"AAA","BBB","CCC"}
    wide=svc.wide_progress(limit=10)
    assert wide["stages"] == [5,15]
    assert wide["rows"][0]["stages"]["TUKIN"]["5"]["2025"]["ntpn_count"] == 2
    conn=duckdb.connect(svc.db_path)
    ledger=conn.execute("SELECT periode_master_id,nip,tahun_sp2d,komponen,tahap,bank,ntpn,transaction_fingerprint FROM pengembalian_transaction_ledger ORDER BY tahap,bank").fetchall()
    assert len(ledger) == 3 and all(item[-1] for item in ledger)
    obligation=conn.execute("SELECT komponen,nominal_kewajiban FROM pengembalian_obligation_component WHERE nip='A' ORDER BY komponen").fetchall()
    conn.close()
    assert obligation == [("TUKIN",1_000_000),("UANG_MAKAN",0)]


def test_baseline_parser_uses_only_authoritative_sheets_and_pengembalian_columns(tmp_path):
    from openpyxl import Workbook
    path=tmp_path/"baseline.xlsx"
    workbook=Workbook();workbook.remove(workbook.active)
    for name,bank in [("BNBA PENGEMBALIAN MANDIRI","BANK MANDIRI"),("BNBA PENGEMBALIAN BRI","BANK RAKYAT INDONESIA"),("BNBA PENGEMBALIAN BSI","BANK SYARIAH INDONESIA")]:
        sheet=workbook.create_sheet(name)
        headers=["NIP","Nama sk","Provinsi","NAMA BANK","NOMER REKENING","Pengembalian Tukin Tahap XV","Sisa Piutang Tukin Tahap XV","Pengembalian Uang Makan Tahap V"]
        sheet.append(headers);sheet.append([None,None,None,None,None,2025,2025,2025]);sheet.append([None,None,None,None,None,"NTPN-T","BAD","NTPN-U"]);sheet.append([]);sheet.append([None]*8);sheet.append(["A","SDM A","Jawa Barat",bank,"123",100,999,50])
    all_sheet=workbook.create_sheet("BNBA PENGEMBALIAN ALL SDM");all_sheet.append(["NIP","Pengembalian Tukin Tahap XV"]);all_sheet.append([None,2025]);all_sheet.append([None,"DOUBLE"]);all_sheet.append([]);all_sheet.append([]);all_sheet.append(["A",100])
    workbook.create_sheet("REKAPITULASI")
    workbook.save(path)
    svc=service(tmp_path)
    rows,adapter=svc.parse_baseline_workbook(path)
    assert len(rows) == 6
    assert adapter["transaction_rows"] == 6
    assert {item["Tahap"] for item in rows} == {5,15}
    assert all(item["Source Sheet"] != "BNBA PENGEMBALIAN ALL SDM" for item in rows)
    assert sum(item["Nominal Berhasil"] for item in rows) == 450


def test_migration_backfills_stage_and_ntpn_for_committed_legacy_progress(tmp_path):
    svc=service(tmp_path)
    conn=duckdb.connect(svc.db_path)
    conn.execute("""INSERT INTO pengembalian_progress_transaction
        (transaction_id,fingerprint,import_batch_id,obligation_id,nip,tahun_sp2d,
         periode_debit,komponen_debit,nominal_debit_tukin,nominal_debit_total,status_debit,raw_data)
        VALUES ('LEGACY-TX','LEGACY-FP','LEGACY-BATCH','O-A-2025','A',2025,
                'Tahap 12 - September','TUKIN',100,100,'BERHASIL','{"NTPN":"NTPN-LEGACY"}')""")
    conn.close()
    import importlib.util
    from pathlib import Path
    migration_path=Path(__file__).resolve().parents[1]/"migrations"/"046_create_progress_stage_ledger.py"
    spec=importlib.util.spec_from_file_location("migration_046_test",migration_path)
    migration=importlib.util.module_from_spec(spec);spec.loader.exec_module(migration)
    migration.main(svc.db_path)
    conn=duckdb.connect(svc.db_path)
    value=conn.execute("SELECT tahap,ntpn,nominal_diajukan,nominal_berhasil FROM pengembalian_progress_transaction WHERE transaction_id='LEGACY-TX'").fetchone()
    conn.close()
    assert value == (12,"NTPN-LEGACY",100,100)


def _dashboard_metadata(svc):
    conn = duckdb.connect(svc.db_path)
    rows = [
        ("F-A", "A", 2025, "Jawa Barat", "Bandung", "BATCH 1", "BANK BRI"),
        ("F-B25", "B", 2025, "Jawa Barat", "Bogor", "BATCH 2", "BANK MANDIRI"),
        ("F-B26", "B", 2026, "Jawa Barat", "Bogor", "BATCH 2", "BANK MANDIRI"),
    ]
    conn.executemany("""INSERT INTO pengembalian_final
        (record_id,comparison_key,nip,tahun_sp2d,provinsi,kabupaten_kota,batch,row_json)
        VALUES (?,?,?,?,?,?,?,?)""", [
        (record_id, record_id, nip, year, province, city, batch,
         '{"nama_bank":"' + bank + '"}')
        for record_id, nip, year, province, city, batch, bank in rows
    ])
    conn.close()


def _stage_row(nip, year, stage, amount, component="TUKIN", bank="BANK BRI", status="BERHASIL", date="18/07/2026"):
    return {"NIP": nip, "Nama": f"SDM {nip}", "Provinsi": "Jawa Barat", "Bank": bank,
            "Tahun SP2D": year, "Komponen": component, "Tahap": stage,
            "Nominal Diajukan": amount, "Nominal Berhasil": amount if status == "BERHASIL" else 0,
            "Status": status, "Tanggal Proses": date}


def test_operational_dashboard_filters_are_and_cumulative_and_read_only(tmp_path):
    svc = service(tmp_path)
    _dashboard_metadata(svc)
    svc.import_rows([
        _stage_row("A", 2025, 1, 100_000),
        _stage_row("A", 2025, 2, 200_000, bank="BANK MANDIRI", date="19/07/2026"),
        _stage_row("B", 2025, 2, 300_000, component="UANG_MAKAN", bank="BANK BRI", date="20/07/2026"),
        _stage_row("B", 2026, 3, 100_000, bank="BANK BRI", date="21/07/2026"),
    ], {"file_hash": "dashboard-filters"})

    baseline = svc.operational_dashboard()["kpi"]
    assert baseline["total_kewajiban"] == 3_000_000
    assert svc.operational_dashboard({"tahun_sp2d": 2025})["kpi"]["total_kewajiban"] == 2_500_000
    # Master/payroll bank is BRI for A even though its Tahap 2 ledger says Mandiri.
    bank = svc.operational_dashboard({"bank": "BRI"})["kpi"]
    assert (bank["total_kewajiban"], bank["total_realisasi_netto"], bank["total_sisa"], bank["sebagian"]) == (
        1_000_000, 300_000, 700_000, 1,
    )
    mandiri = svc.operational_dashboard({"bank": "mandiri"})["kpi"]
    assert (mandiri["total_kewajiban"], mandiri["total_realisasi_netto"], mandiri["total_sisa"], mandiri["sebagian"]) == (
        2_000_000, 400_000, 1_600_000, 2,
    )
    assert svc.operational_dashboard({"bank": "BANK-TIDAK-ADA"})["kpi"] == {
        "total_kewajiban": 0, "total_realisasi_netto": 0, "total_sisa": 0,
        "total_kelebihan_bayar": 0, "lunas": 0, "sebagian": 0, "belum": 0, "overpaid": 0,
    }
    assert svc.operational_dashboard({"provinsi": "Jawa Barat", "kabupaten_kota": "Bogor"})["kpi"]["total_kewajiban"] == 2_000_000
    assert svc.operational_dashboard({"batch_bpk": "BATCH 2"})["kpi"]["total_kewajiban"] == 2_000_000
    assert svc.operational_dashboard({"komponen": "UANG_MAKAN"})["kpi"]["total_kewajiban"] == 500_000
    # Tahap selects A but preserves A's entire cumulative ledger: 100k + 200k.
    stage = svc.operational_dashboard({"tahap": 2, "stage_mode": "single"})["kpi"]
    assert (stage["total_kewajiban"], stage["total_realisasi_netto"]) == (1_000_000, 300_000)
    stage_um = svc.operational_dashboard({"tahap": 2, "stage_mode": "single", "komponen": "UANG_MAKAN"})["kpi"]
    assert (stage_um["total_kewajiban"], stage_um["total_realisasi_netto"]) == (500_000, 300_000)
    combined = svc.operational_dashboard({"tahun_sp2d": 2025, "bank": "BRI", "kabupaten_kota": "Bandung", "tahap": 2})["kpi"]
    assert combined["total_kewajiban"] == 1_000_000
    assert svc.operational_dashboard({"provinsi": "Tidak Ada"})["kpi"] == {
        "total_kewajiban": 0, "total_realisasi_netto": 0, "total_sisa": 0,
        "total_kelebihan_bayar": 0, "lunas": 0, "sebagian": 0, "belum": 0, "overpaid": 0,
    }


def test_operational_dashboard_options_follow_current_metadata_without_duplicates(tmp_path):
    svc = service(tmp_path)
    _dashboard_metadata(svc)
    conn = duckdb.connect(svc.db_path)
    conn.executemany("""INSERT INTO pengembalian_final
        (record_id,comparison_key,nip,tahun_sp2d,provinsi,kabupaten_kota,batch,row_json)
        VALUES (?,?,?,?,?,?,?,?)""", [
        ("F-A-SPACE", "F-A-SPACE", "A", 2025, "Jawa Barat", "Bandung", "BATCH 1", '{"nama_bank":" BRI "}'),
        ("F-A-CASE", "F-A-CASE", "A", 2025, "Jawa Barat", "Bandung", "BATCH 1", '{"nama_bank":"bri"}'),
        ("F-A-NULL", "F-A-NULL", "A", 2025, "Jawa Barat", "Bandung", "BATCH 1", '{"nama_bank":null}'),
        ("F-A-BLANK", "F-A-BLANK", "A", 2025, "Jawa Barat", "Bandung", "BATCH 1", '{"nama_bank":""}'),
    ])
    conn.close()
    svc.import_rows([
        _stage_row("A", 2025, 1, 100_000),
        _stage_row("A", 2025, 2, 200_000, date="19/07/2026"),
        _stage_row("B", 2025, 2, 300_000, component="UANG_MAKAN", date="20/07/2026"),
    ], {"file_hash": "dashboard-options"})
    options = svc.filter_options()["options"]
    assert options["tahun_sp2d"] == [2025, 2026]
    assert options["bank"] == ["BRI", "MANDIRI"]
    assert options["kabupaten_kota"] == ["Bandung", "Bogor"]
    assert options["batch_bpk"] == ["BATCH 1", "BATCH 2"]
    assert options["tahap"] == [1, 2]
    assert svc.filter_options({"provinsi": "Jawa Barat", "kabupaten_kota": "Bandung"})["options"]["kabupaten_kota"] == ["Bandung"]

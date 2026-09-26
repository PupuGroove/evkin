import duckdb
from openpyxl import load_workbook

from backend.final_tukin import FinalTukinService, calculate_final_tukin
from backend.shared_master import SharedMasterService
from backend.skp_assessment import SkpAssessmentService, skp_ready_for_tukin
from backend.tukin_combined import (
    calculate_tukin_amounts,
    create_calculation_snapshot,
    export_calculation_snapshot,
    filter_skp_status,
    integration_readiness,
    combine,
)


PERIOD = "2026-09"
NIPS = [f"19870101202601100{i}" for i in range(1, 5)]


def setup_four(tmp_path):
    path = str(tmp_path / "direct-final.duckdb")
    SharedMasterService(path).ingest_rows([
        {"NIP": nip, "Nama": f"PESERTA {index}", "Tukin": 10_000_000,
         "Berhak Tukin": "YA"}
        for index, nip in enumerate(NIPS, 1)
    ], "2026-01-01", "master.xlsx", "ADMIN", True)
    service = SkpAssessmentService(path)
    with service.connect() as conn:
        conn.execute("""CREATE TABLE final_absensi (
            nip VARCHAR, nama VARCHAR, bulan VARCHAR, potPct DECIMAL(9,4),
            import_batch_id VARCHAR, finalized_at TIMESTAMP)""")
        conn.executemany("INSERT INTO final_absensi VALUES (?,?,?,10,'ABS-FINAL',now())", [
            [nip, f"PESERTA {index}", PERIOD] for index, nip in enumerate(NIPS, 1)
        ])
    return path, service


def import_final(service, nips=NIPS):
    return service.import_rows({
        "assessment_period": PERIOD, "import_mode": "FINAL", "confirmed": True,
        "final_confirmed": True, "imported_by": "PETUGAS FINAL",
        "source_file": "hasil-final.xlsx",
        "rows": [{"NIP": nip, "Nama": f"PESERTA {index}", "Predikat": "Baik"}
                 for index, nip in enumerate(nips, 1)],
    })


def test_canonical_gate_keeps_internal_strict_and_accepts_valid_direct_final():
    approved = {"review_status": "APPROVED", "validation_status": "VALID",
                "final_predicate": "Baik", "report_status": "SUBMITTED"}
    pending = {**approved, "review_status": "PENDING_REVIEW"}
    direct = {**approved, "review_status": "FINAL_IMPORTED"}
    invalid_direct = {**direct, "validation_status": "INVALID"}
    assert skp_ready_for_tukin(approved)["ready"] is True
    pending_result = skp_ready_for_tukin(pending)
    assert pending_result["ready"] is False
    assert pending_result["source_status"] == "INTERNAL_NOT_FINAL"
    assert pending_result["ready_mode"] == "BLOCKED"
    assert pending_result["reason"] == "SKP internal belum APPROVED"
    assert skp_ready_for_tukin(direct)["ready"] is True
    result = skp_ready_for_tukin(invalid_direct)
    assert result["ready"] is False and result["reason_code"] == "FINAL_IMPORT_INVALID"


def test_four_valid_direct_final_rows_are_ready_and_calculable(tmp_path):
    path, service = setup_four(tmp_path)
    result = import_final(service)
    assert result["source_type"] == "DIRECT_FINAL_IMPORT"
    readiness = integration_readiness(path, PERIOD)
    assert readiness["target_count"] == 4
    assert readiness["complete_count"] == 4
    assert readiness["blocked_count"] == 0
    assert readiness["skp_ready"] == 4
    assert readiness["can_calculate"] is True
    assert readiness["blocker_counts"]["skp_internal_not_approved"] == 0
    assert readiness["blocker_counts"]["skp_final_import_invalid"] == 0
    calculated = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    assert calculated["status"] == "CALCULATED" and calculated["ready"] == 4
    after_calculation = integration_readiness(path, PERIOD)
    assert after_calculation["can_finalize"] is True
    calculation_export = export_calculation_snapshot(path, calculated["calculation_id"])
    assert len(calculation_export.getvalue()) > 1_000
    final = FinalTukinService(path).approve_calculation(
        calculated["calculation_id"], "APPROVER", "Dokumen final uji"
    )
    assert final["status"] == "FINAL"
    final_export = FinalTukinService(path).export_final_decision(final["final_decision_id"])
    assert len(final_export.getvalue()) > 1_000


def test_mixed_source_only_blocks_internal_pending(tmp_path):
    path, service = setup_four(tmp_path)
    import_final(service, NIPS[2:])
    with service.connect() as conn:
        conn.executemany("""INSERT INTO skp_assessment_monthly
            (assessment_period,nip,nama,final_predicate,report_status,validation_status,
             review_status,final_reduction_rate,source_file)
            VALUES (?,?,?,'Baik','SUBMITTED','VALID',?,?, 'internal.xlsx')""", [
            [PERIOD, NIPS[0], "PESERTA 1", "APPROVED", 0],
            [PERIOD, NIPS[1], "PESERTA 2", "PENDING_REVIEW", 0],
        ])
    readiness = integration_readiness(path, PERIOD)
    assert (readiness["target_count"], readiness["complete_count"], readiness["blocked_count"]) == (4, 3, 1)
    assert readiness["blocker_counts"]["skp_internal_not_approved"] == 1
    assert readiness["blocker_counts"]["skp_final_import_invalid"] == 0
    assert "SKP internal belum APPROVED" in readiness["blockers"][0]["issues"]


def test_invalid_direct_final_is_blocked_for_validation_not_approval(tmp_path):
    path, service = setup_four(tmp_path)
    import_final(service)
    with duckdb.connect(path) as conn:
        conn.execute("""UPDATE skp_assessment_monthly SET validation_status='INVALID'
                        WHERE assessment_period=? AND nip=?""", [PERIOD, NIPS[0]])
    readiness = integration_readiness(path, PERIOD)
    assert (readiness["complete_count"], readiness["blocked_count"]) == (3, 1)
    assert readiness["blocker_counts"]["skp_final_import_invalid"] == 1
    assert readiness["blocker_counts"]["skp_internal_not_approved"] == 0
    assert "hasil import final SKP tidak valid" in readiness["blockers"][0]["issues"]


def test_direct_final_formula_parity_and_snapshot_freshness(tmp_path):
    internal = calculate_tukin_amounts(10_000_000, 10, 5)
    direct = calculate_tukin_amounts(10_000_000, 10, 5)
    assert internal == direct
    path, service = setup_four(tmp_path)
    import_final(service)
    calculation = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    assert integration_readiness(path, PERIOD)["calculation"]["status"] == "CALCULATED"
    with duckdb.connect(path) as conn:
        conn.execute("""UPDATE skp_assessment_monthly
                        SET final_predicate='Kurang',final_reduction_rate=10,updated_at=now()
                        WHERE assessment_period=? AND nip=?""", [PERIOD, NIPS[0]])
    stale = integration_readiness(path, PERIOD)
    assert stale["calculation"]["calculation_id"] == calculation["calculation_id"]
    assert stale["calculation"]["status"] == "STALE"


def test_legacy_final_tukin_consumer_uses_same_direct_final_semantics(tmp_path):
    path, service = setup_four(tmp_path)
    import_final(service)
    master, _, sources = FinalTukinService(path).sources(
        PERIOD, "ABSENSI_SKP", "ASSESSMENT"
    )
    source = next(item for item in sources if item["nip"] == NIPS[0])
    assert source["source_type"] == "DIRECT_FINAL_IMPORT"
    result = calculate_final_tukin(
        NIPS[0], PERIOD, "ABSENSI_SKP", {"total": 10}, master[NIPS[0]], source
    )
    assert result["status"] == "READY"


def test_actual_scale_1563_with_four_genuinely_missing_is_fully_calculable(tmp_path):
    path = str(tmp_path / "scale-1563.duckdb")
    nips = [f"19{index:016d}" for index in range(1, 1564)]
    SharedMasterService(path).ingest_rows([
        {"NIP": nip, "Nama": f"PESERTA {index}", "Tukin": 3_915_950,
         "Berhak Tukin": "YA"}
        for index, nip in enumerate(nips, 1)
    ], "2026-01-01", "master-1563.xlsx", "ADMIN", True)
    service = SkpAssessmentService(path)
    with service.connect() as conn:
        conn.execute("""CREATE TABLE final_absensi (
            nip VARCHAR, nama VARCHAR, bulan VARCHAR, potPct DECIMAL(9,4),
            import_batch_id VARCHAR, finalized_at TIMESTAMP)""")
        conn.executemany("INSERT INTO final_absensi VALUES (?,?,?,0,'ABS-1563',now())", [
            [nip, f"PESERTA {index}", PERIOD] for index, nip in enumerate(nips, 1)
        ])
        conn.executemany("""INSERT INTO skp_assessment_monthly
            (assessment_period,nip,nama,final_predicate,report_status,validation_status,
             review_status,final_reduction_rate,source_file)
            VALUES (?,?,?,'Baik','SUBMITTED','VALID','FINAL_IMPORTED',0,'final-1559.xlsx')""", [
            [PERIOD, nip, f"PESERTA {index}"] for index, nip in enumerate(nips[:1559], 1)
        ])
    readiness = integration_readiness(path, PERIOD)
    assert readiness["target_count"] == 1563
    assert readiness["skp_normal_ready"] == 1559
    assert readiness["skp_missing_fallback"] == 4
    assert readiness["complete_count"] == 1563
    assert readiness["blocked_count"] == 0
    assert readiness["warning_count"] == 4
    assert readiness["can_calculate"] is True


def test_mixed_semantic_population_has_deterministic_precedence():
    nips = [f"20{index:016d}" for index in range(1, 13)]
    master = {nip: {"nama": nip, "tukin_gross": 1_000_000, "eligible": True} for nip in nips}
    attendance = [{"nip": nip, "potPct": 0} for nip in nips]
    assessments = []
    for nip in nips[:5]:
        assessments.append({"nip": nip, "final_predicate": "Baik", "report_status": "SUBMITTED",
                            "validation_status": "VALID", "review_status": "FINAL_IMPORTED",
                            "final_reduction_rate": 0, "_direct_final_import": True})
    for nip in nips[5:8]:
        assessments.append({"nip": nip, "final_predicate": "Baik", "report_status": "SUBMITTED",
                            "validation_status": "VALID", "review_status": "APPROVED",
                            "final_reduction_rate": 0})
    assessments.append({"nip": nips[10], "final_predicate": "Baik", "report_status": "SUBMITTED",
                        "validation_status": "VALID", "review_status": "PENDING_REVIEW",
                        "final_reduction_rate": 0})
    assessments.append({"nip": nips[11], "final_predicate": "Baik", "report_status": "SUBMITTED",
                        "validation_status": "INVALID", "review_status": "FINAL_IMPORTED",
                        "final_reduction_rate": 0, "_direct_final_import": True})
    rows = combine(master, attendance, assessments, [], PERIOD)
    statuses = [row["skp_source_status"] for row in rows]
    assert statuses.count("FINAL_IMPORTED") == 5
    assert statuses.count("INTERNAL_APPROVED") == 3
    assert statuses.count("MISSING_FROM_FINAL_IMPORT") == 2
    assert statuses.count("INTERNAL_NOT_FINAL") == 1
    assert statuses.count("INVALID_FINAL_IMPORT") == 1
    assert sum(row["status"] == "LENGKAP" for row in rows) == 10
    assert len(filter_skp_status(rows, "normal")) == 8
    assert len(filter_skp_status(rows, "missing")) == 2
    assert len(filter_skp_status(rows, "problem")) == 2


def test_missing_formula_uses_35_percent_of_performance_component_only():
    amounts = calculate_tukin_amounts(3_915_950, 0, 35)
    assert amounts["skp_component"] == 2_741_165
    assert amounts["skp_deduction"] == 959_408
    assert amounts["attendance_component"] == 1_174_785
    assert amounts["attendance_deduction"] == 0
    assert amounts["calculated_entitlement"] == 2_956_542


def test_missing_snapshot_exports_flags_and_becomes_stale_when_final_arrives(tmp_path):
    path, service = setup_four(tmp_path)
    import_final(service, NIPS[:3])
    before = integration_readiness(path, PERIOD)
    assert before["skp_missing_fallback"] == 1
    calculation = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    workbook = load_workbook(export_calculation_snapshot(path, calculation["calculation_id"]), read_only=True)
    sheet = workbook["Hasil Perhitungan"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    source_index = headers.index("Status Sumber SKP")
    mode_index = headers.index("Mode Perhitungan SKP")
    policy_index = headers.index("Kebijakan Fallback SKP")
    exported = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
    fallback = next(row for row in exported if row[source_index] == "MISSING_FROM_FINAL_IMPORT")
    assert fallback[mode_index] == "FALLBACK"
    assert fallback[policy_index] == "MISSING_SKP_35_PERFORMANCE"
    import_final(service, [NIPS[3]])
    stale = integration_readiness(path, PERIOD)
    assert stale["calculation"]["status"] == "STALE"
    assert stale["skp_missing_fallback"] == 0 and stale["skp_normal_ready"] == 4
    recalculated = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    assert recalculated["ready"] == 4


def test_final_source_disappearing_after_calculation_also_marks_snapshot_stale(tmp_path):
    path, service = setup_four(tmp_path)
    import_final(service)
    calculation = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    with duckdb.connect(path) as conn:
        conn.execute("DELETE FROM skp_assessment_monthly WHERE assessment_period=? AND nip=?", [PERIOD, NIPS[0]])
    readiness = integration_readiness(path, PERIOD)
    assert readiness["calculation"]["calculation_id"] == calculation["calculation_id"]
    assert readiness["calculation"]["status"] == "STALE"
    assert readiness["skp_missing_fallback"] == 1


def test_fallback_participants_can_be_finalized_and_exported(tmp_path):
    path, _ = setup_four(tmp_path)
    readiness = integration_readiness(path, PERIOD)
    assert readiness["complete_count"] == 4 and readiness["warning_count"] == 4
    calculation = create_calculation_snapshot(path, PERIOD, "KALKULATOR")
    service = FinalTukinService(path)
    decision = service.approve_calculation(calculation["calculation_id"], "APPROVER", "Fallback sah")
    workbook = load_workbook(service.export_final_decision(decision["final_decision_id"]), read_only=True)
    headers = [cell.value for cell in next(workbook["FINAL TUKIN"].iter_rows(min_row=1, max_row=1))]
    assert "Status Sumber SKP" in headers
    assert "Mode Perhitungan SKP" in headers
    assert "Kebijakan Fallback SKP" in headers

import json

import duckdb
import pytest
from openpyxl import load_workbook

from backend.tukin_reconciliation import (
    TukinReconciliationService,
    _money,
    reconcile_records,
)


PERIOD = "2026-08"
NIP = "198701012026011001"


def canonical(**overrides):
    row = {
        "nip": NIP, "nama": "UAT", "unit": "UNIT", "periode": PERIOD,
        "bruto": 10_000_000, "potongan_absensi": 600_000,
        "potongan_skp": 700_000, "adjustment": 0, "tukin_final": 8_700_000,
        "stale_status": "CURRENT", "source_invalid": False, "broken_lineage": False,
        "lineage_issues": [],
    }
    row.update(overrides)
    return row


def legacy(**overrides):
    row = {
        "nip": NIP, "nama": "UAT", "unit": "UNIT", "periode": PERIOD,
        "bruto_legacy": 10_000_000, "potongan_absensi_legacy": 600_000,
        "potongan_skp_legacy": 700_000, "adjustment_legacy": 0,
        "tukin_final_legacy": 8_700_000, "legacy_reference": "OFFICIAL-01",
    }
    row.update(overrides)
    return row


def setup_db(tmp_path):
    path = str(tmp_path / "shadow.duckdb")
    snapshot = {
        "nama": "UAT", "unit": "UNIT", "attendance_component": 3_000_000,
        "skp_component": 7_000_000, "attendance_rate": 20,
        "predicate_skp": "Kurang",
    }
    with duckdb.connect(path) as conn:
        conn.execute("CREATE TABLE master_data(nip VARCHAR,nama VARCHAR,unit VARCHAR,tukin_gross BIGINT,master_version INTEGER)")
        conn.execute("INSERT INTO master_data VALUES (?,?,?,10000000,1)", [NIP, "UAT", "UNIT"])
        conn.execute("CREATE TABLE final_absensi(nip VARCHAR,bulan VARCHAR,potPct DECIMAL(9,4))")
        conn.execute("INSERT INTO final_absensi VALUES (?,?,20)", [NIP, PERIOD])
        conn.execute("""CREATE TABLE skp_assessment_monthly(
            nip VARCHAR,assessment_period VARCHAR,review_status VARCHAR,rule_version VARCHAR,final_predicate VARCHAR)""")
        conn.execute("INSERT INTO skp_assessment_monthly VALUES (?,?,'APPROVED','RULE-V1','KURANG')", [NIP, PERIOD])
        conn.execute("""CREATE TABLE tukin_calculation_batch(
            id VARCHAR,assessment_period VARCHAR,version INTEGER,status VARCHAR,
            source_fingerprint VARCHAR,calculated_at TIMESTAMP)""")
        conn.execute("INSERT INTO tukin_calculation_batch VALUES ('CALC-1',?,1,'CALCULATED','FP-1',now())", [PERIOD])
        conn.execute("""CREATE TABLE tukin_calculation_result(
            calculation_id VARCHAR,assessment_period VARCHAR,nip VARCHAR,status VARCHAR,
            tukin_gross BIGINT,attendance_deduction BIGINT,skp_deduction BIGINT,
            adjustment BIGINT,total_deduction BIGINT,tukin_paid BIGINT,master_version INTEGER,
            attendance_source_id VARCHAR,attendance_batch_id VARCHAR,skp_source_id VARCHAR,
            skp_source_version VARCHAR,skp_batch_id VARCHAR,source_fingerprint VARCHAR,
            calculated_at TIMESTAMP,snapshot VARCHAR)""")
        conn.execute("""INSERT INTO tukin_calculation_result VALUES
            ('CALC-1',?,?,'READY',10000000,600000,700000,0,1300000,8700000,1,
             'ABS-1','ABS-B1','SKP-1','RULE-V1','SKP-B1','ROW-FP',now(),?)""",
            [PERIOD, NIP, json.dumps(snapshot)])
        conn.execute("""CREATE TABLE tukin_final_decision(
            id VARCHAR,assessment_period VARCHAR,calculation_id VARCHAR,status VARCHAR)""")
        conn.execute("INSERT INTO tukin_final_decision VALUES ('FINAL-1',?,'CALC-1','FINAL')", [PERIOD])
        conn.execute("""CREATE TABLE tukin_final_monthly(
            assessment_period VARCHAR,nip VARCHAR,nama VARCHAR,unit VARCHAR,tukin_gross DECIMAL(18,2),
            attendance_deduction DECIMAL(18,2),skp_deduction DECIMAL(18,2),manual_adjustment DECIMAL(18,2),
            paid_amount DECIMAL(18,2),payment_status VARCHAR,import_batch_id VARCHAR,source_file VARCHAR,source_row INTEGER)""")
        conn.execute("""INSERT INTO tukin_final_monthly VALUES
            (?,?, 'UAT','UNIT',10000000,600000,700000,0,8700000,'FINAL','LEG-B1','official.xlsx',2)""", [PERIOD, NIP])
    return path


def test_exact_match_and_delta_one_rupiah_are_distinct():
    matched = reconcile_records([canonical()], [legacy()])[0]
    changed = reconcile_records([canonical()], [legacy(tukin_final_legacy=8_699_999)])[0]
    assert matched["category"] == "MATCH"
    assert changed["category"] == "UNEXPLAINED_DIFFERENCE"
    assert changed["delta_final"] == 1


def test_missing_canonical_and_missing_legacy_are_detected():
    assert reconcile_records([], [legacy()])[0]["category"] == "MISSING_CANONICAL"
    assert reconcile_records([canonical()], [])[0]["category"] == "MISSING_LEGACY"


def test_scope_excludes_master_and_legacy_outside_processing_population():
    in_scope = legacy()
    out_scope = legacy(nip="198701012026011099")
    population = {
        "processing_nips": {NIP},
        "final_assessment_nips": {NIP},
        "final_attendance_nips": {NIP},
        "master_nips": {NIP},
    }
    rows = reconcile_records([], [in_scope, out_scope], population=population)
    by_nip = {row["nip"]: row["category"] for row in rows}
    assert by_nip[NIP] == "MISSING_CANONICAL"
    assert by_nip[out_scope["nip"]] == "OUT_OF_SCOPE"


def test_duplicate_is_detected_at_nip_period_grain():
    result = reconcile_records([canonical(), canonical()], [legacy()])
    assert result[0]["category"] == "DUPLICATE"
    assert result[0]["duplicate_canonical"] is True


def test_verified_explanation_does_not_become_unexplained():
    result = reconcile_records(
        [canonical()], [legacy(tukin_final_legacy=8_600_000)],
        {NIP: {"verified": True, "reason": "Aturan historis", "evidence": "SK-01"}},
    )[0]
    assert result["category"] == "EXPLAINED_DIFFERENCE"
    assert result["explanation_evidence"] == "SK-01"


def test_broken_lineage_and_stale_source_are_detected(tmp_path):
    path = setup_db(tmp_path)
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE tukin_calculation_result SET attendance_source_id='' WHERE calculation_id='CALC-1'")
    broken = TukinReconciliationService(path).reconcile(PERIOD, "tukin_final_monthly")
    assert broken["details"][0]["category"] == "SOURCE_INVALID"
    assert broken["acceptance"]["broken_calculation_lineage"] == 1
    with duckdb.connect(path) as conn:
        conn.execute("UPDATE tukin_calculation_result SET attendance_source_id='ABS-1'")
        conn.execute("UPDATE final_absensi SET potPct=21")
    stale = TukinReconciliationService(path).reconcile(PERIOD, "tukin_final_monthly")
    assert stale["details"][0]["category"] == "STALE"
    assert stale["acceptance"]["stale_calculation_approved"] == 1


def test_totals_equal_detail_aggregation_and_money_is_integer(tmp_path):
    result = TukinReconciliationService(setup_db(tmp_path)).reconcile(PERIOD, "tukin_final_monthly")
    detail = result["details"]
    assert result["totals"]["final"]["canonical"] == sum(row["tukin_final"] or 0 for row in detail)
    assert result["totals"]["final"]["legacy"] == sum(row["tukin_final_legacy"] or 0 for row in detail)
    assert isinstance(result["totals"]["final"]["delta"], int)
    assert _money("10000000.00") == 10_000_000
    with pytest.raises(ValueError, match="bilangan bulat"):
        _money(1.5)


def test_shadow_reconciliation_does_not_write_or_pay(tmp_path):
    path = setup_db(tmp_path)
    with duckdb.connect(path) as conn:
        before = conn.execute("SELECT status,count(*) FROM tukin_final_decision GROUP BY status").fetchall()
        tables_before = conn.execute("SHOW TABLES").fetchall()
    service = TukinReconciliationService(path)
    result = service.reconcile(PERIOD, "tukin_final_monthly")
    workbook = load_workbook(service.export(result), read_only=True, data_only=True)
    with duckdb.connect(path) as conn:
        after = conn.execute("SELECT status,count(*) FROM tukin_final_decision GROUP BY status").fetchall()
        tables_after = conn.execute("SHOW TABLES").fetchall()
    assert result["mode"] == "SHADOW_READ_ONLY"
    assert before == after == [("FINAL", 1)]
    assert tables_before == tables_after
    assert workbook.sheetnames == [
        "SUMMARY", "MATCH", "EXPLAINED DIFFERENCE", "UNEXPLAINED DIFFERENCE",
        "MISSING CANONICAL", "MISSING LEGACY", "MISSING ATTENDANCE",
        "MISSING ASSESSMENT", "MISSING MASTER", "SOURCE INVALID", "STALE",
        "DUPLICATE", "OUT OF SCOPE", "SOURCE LINEAGE",
    ]
